"""Offline namespace controls, not proof of new LibreOffice qualification."""
import base64
import hashlib
import io
import json
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

from skillopt.coevolution_v5.core import seal
from skillopt.continual_eval import sheet_recalc as v5
from skillopt.continual_eval import sheet_recalc_functions as v8
from skillopt.continual_eval import sheet_recalc_v7 as v7
from tests.test_continual_eval_sheet_recalc_v7 import _cache


def book(path, formula="=_xlfn.IFNA(A1,5)", *, cache="5", array=False, observer=False, name=False):
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.active["A1"] = 5
    wb.active["B1"] = ArrayFormula(ref="B1", text=formula) if array else formula
    if observer:
        wb.active["C1"] = "=FORMULATEXT(B1)"
    if name:
        wb.defined_names.add(DefinedName("SomeRange", attr_text="S!A1"))
    wb.save(path)
    wb.close()
    _cache(path, cache)
    return path


def unpack(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        return {name: z.read(name) for name in z.namelist()}


def native_fixture(self, path):
    """Synthetic engine receipt, deliberately not a real Calc qualification."""
    _, proof = v8.screening_view(path)
    manifest = v5._numeric_manifest(path)
    return seal({"input_sha256": v5.sha(path), "identity": self.identity,
                 "status": "available", "reason": "fixture", "output_sha256": v5.sha(path),
                 "output_base64": base64.b64encode(path.read_bytes()).decode(), "function_alias_view": proof,
                 "numeric_manifest_sha256": hashlib.sha256(json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()).hexdigest(),
                 "numeric_literals_verified": len(manifest["cells"]),
                 "cleanup_confirmed": True, "container_execution_attempted": True,
                 "duration_seconds": 0.25, "model_api_calls": 0})


@pytest.mark.parametrize("formula,expected", [
    ("=_xlfn.IFNA(A1,5)", "=IFNA(A1,5)"),
    ("=_xlfn.AGGREGATE(15,6,A1:A5,1)", "=AGGREGATE(15,6,A1:A5,1)"),
    ('=_xlfn.IFNA(A1,"_xlfn.IFNA(A1,99)")', '=IFNA(A1,"_xlfn.IFNA(A1,99)")'),
    ('="_xlfn.IFNA(A1,5)"', '="_xlfn.IFNA(A1,5)"'),
    ("='_xlfn.IFNA'!A1", "='_xlfn.IFNA'!A1"),
    ("=_xlfn.IFNA_RANGE", "=_xlfn.IFNA_RANGE"),
    ("=_xlfn.IFNA2(A1,5)", "=_xlfn.IFNA2(A1,5)"),
    ("=_xlfn.UNIQUE(A1:A5)", "=_xlfn.UNIQUE(A1:A5)"),
    ("=IFNA(A1,5)", "=IFNA(A1,5)"),
    ("=ifna(A1,5)", "=IFNA(A1,5)"),
    ("=aggregate(15,6,A1:A5,1)", "=AGGREGATE(15,6,A1:A5,1)"),
    ('=ifna(A1,"ifna(A1,99)")', '=IFNA(A1,"ifna(A1,99)")'),
    ("=IfNa(A1,5)", "=IfNa(A1,5)"),
    ("=sum(A1:A5)", "=sum(A1:A5)"),
    ("=ifna2(A1,5)", "=ifna2(A1,5)"),
    ("=_xlfn.ifna(A1,5)", "=_xlfn.ifna(A1,5)"),
])
def test_only_two_exact_function_tokens_are_aliases(formula, expected):
    assert v8.canonical_formula(formula) == expected


@pytest.mark.parametrize("array", [False, True])
def test_execution_and_comparison_views_only_change_formula_text(tmp_path, array):
    source = book(tmp_path / "source.xlsx", array=array)
    original_bytes = source.read_bytes()
    raw, proof = v8.screening_view(source)
    assert proof["original_input_sha256"] == v5.sha(source)
    assert proof["screening_view_sha256"] == hashlib.sha256(raw).hexdigest()
    assert proof["execution_input_sha256"] == v5.sha(source)
    before, after = unpack(original_bytes), unpack(raw)
    assert before.keys() == after.keys()
    for name in before:
        assert (before[name].replace(b"_xlfn.IFNA(", b"IFNA(") if name == "xl/worksheets/sheet1.xml"
                else before[name]) == after[name]
    executed = tmp_path / "execution-view.xlsx"
    executed.write_bytes(raw)
    assert v5.screen(executed) is None
    restored, output_proof = v8.comparison_view(source, executed)
    assert unpack(restored) == before
    assert output_proof["restored_formula_count"] == 1
    assert source.read_bytes() == original_bytes


@pytest.mark.parametrize("formula", ['=_xlfn.IFNA(A1,5)', '=_xlfn.AGGREGATE(15,6,A1:A5,1)'])
def test_export_already_has_original_prefix_needs_no_rewrite(tmp_path, formula):
    source = book(tmp_path / "source.xlsx", formula)
    output, proof = v8.comparison_view(source, source)
    assert output == source.read_bytes() and proof["restored_formula_count"] == 0


@pytest.mark.parametrize("array", [False, True])
@pytest.mark.parametrize("source_formula,export_formula", [
    ("=_xlfn.IFNA(A1,5)", "=ifna(A1,5)"),
    ("=_xlfn.AGGREGATE(15,6,A1:A5,1)", "=aggregate(15,6,A1:A5,1)"),
])
def test_known_lowercase_engine_tokens_restore_without_cache_edits(tmp_path, array, source_formula, export_formula):
    source = book(tmp_path / "source.xlsx", source_formula, array=array)
    target = book(tmp_path / "engine.xlsx", export_formula, cache="99", array=array)
    original, exported = source.read_bytes(), target.read_bytes()
    restored, proof = v8.comparison_view(source, target)
    before, after = unpack(exported), unpack(restored)
    for name in before:
        assert (before[name].replace(export_formula[1:].encode(), source_formula[1:].encode())
                if name == "xl/worksheets/sheet1.xml" else before[name]) == after[name]
    values = openpyxl.load_workbook(io.BytesIO(restored), data_only=True)
    assert values.active["B1"].value == 99
    values.close()
    assert proof["restored_formula_count"] == 1
    assert source.read_bytes() == original and target.read_bytes() == exported


@pytest.mark.parametrize("export_formula", [
    '=ifna(A1,"LEFT")',  # quoted arguments remain case-sensitive
    '=ifna(a1,"left")',  # no broad reference case normalization
    '=ifna(A1,"left")+0',
    '=IfNa(A1,"left")',  # only the observed exact spelling is qualified
])
def test_lowercase_allowlist_does_not_hide_any_other_formula_change(tmp_path, export_formula):
    source = book(tmp_path / "source.xlsx", '=_xlfn.IFNA(A1,"left")')
    target = book(tmp_path / "engine.xlsx", export_formula)
    with pytest.raises(ValueError, match="Alias formula"):
        v8.comparison_view(source, target)


@pytest.mark.parametrize("change", ["argument", "different_function", "array_range", "removed"])
def test_changed_result_formula_is_not_restored_into_original(tmp_path, change):
    source = book(tmp_path / "source.xlsx", array=change == "array_range")
    target = book(tmp_path / "after.xlsx", formula="=IFNA(A1,5)", array=change == "array_range")
    wb = openpyxl.load_workbook(target)
    if change == "argument":
        wb.active["B1"] = "=IFNA(A1,6)"
    elif change == "different_function":
        wb.active["B1"] = "=IFERROR(A1,5)"
    elif change == "array_range":
        wb.active["B1"] = ArrayFormula(ref="B1:B2", text="=IFNA(A1,5)")
    else:
        wb.active["B1"] = 5
    wb.save(target)
    wb.close()
    with pytest.raises(ValueError, match="Alias formula"):
        v8.comparison_view(source, target)


@pytest.mark.parametrize("kind", ["observer", "named_expression", "unrelated_extension", "mixed_extension", "quoted_alias"])
def test_unsafe_or_unqualified_scope_submits_no_container(tmp_path, monkeypatch, kind):
    formula = {"unrelated_extension": "=_xlfn.UNIQUE(A1:A5)",
               "mixed_extension": "=_xlfn.IFNA(_xlfn.UNIQUE(A1:A5),5)",
               "quoted_alias": '="_xlfn.IFNA(A1,5)"'}.get(kind, "=_xlfn.IFNA(A1,5)")
    source = book(tmp_path / "source.xlsx", formula, observer=kind == "observer", name=kind == "named_expression")
    calls = []
    monkeypatch.setattr(v7.Recalculator, "run", lambda *args: calls.append(1))
    result = v8.Recalculator("sha256:" + "a" * 64).run(source)
    assert result["status"] == "unknown" and not result["container_execution_attempted"]
    assert result["reason"] == "unsupported_extended_function" and not calls
    assert result["cleanup_confirmed"]


def test_execution_receipt_and_input_views_are_explicit_not_mislabeled(tmp_path, monkeypatch):
    source = book(tmp_path / "source.xlsx")
    original = source.read_bytes()
    seen = []
    def fixture(self, path):
        assert v5.screen(path) == "unsupported_extended_function"
        seen.append(path.read_bytes())
        return native_fixture(self, path)
    monkeypatch.setattr(v8.Recalculator, "_run_native", fixture)
    result = v8.Recalculator("sha256:" + "a" * 64).run(source)
    assert result["status"] == "available" and len(seen) == 1
    assert result["input_sha256"] == v5.sha(source)
    assert result["execution_receipt"]["input_sha256"] == result["execution_input_sha256"]
    assert result["execution_input_sha256"] == result["input_sha256"]
    assert result["function_alias_view"]["screening_view_sha256"] != result["input_sha256"]
    assert original == seen[0]
    assert "numeric_manifest_sha256" not in result
    assert result["execution_receipt"]["numeric_literals_verified"] == 1
    assert source.read_bytes() == original
    assert unpack(base64.b64decode(result["output_base64"])) == unpack(original)


def test_wrong_numeric_cache_stays_wrong_after_namespace_restoration(tmp_path):
    source = book(tmp_path / "source.xlsx", cache="5")
    wrong = book(tmp_path / "wrong.xlsx", "=IFNA(A1,5)", cache="99")
    raw, _ = v8.comparison_view(source, wrong)
    values = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    assert values.active["B1"].value == 99
    values.close()


def test_wrong_cache_remains_a_failed_paired_comparison(tmp_path, monkeypatch):
    reference = book(tmp_path / "reference.xlsx", cache="5")
    prediction = book(tmp_path / "prediction.xlsx", cache="99")
    monkeypatch.setattr(v8.Recalculator, "_run_native", native_fixture)
    result = v5.evaluate_pair(prediction, reference, "B1", v8.Recalculator("sha256:" + "a" * 64), tmp_path / "receipts")
    assert result["status"] == "fail" and len(result["receipts"]) == 2


def test_xml_escaped_argument_content_and_cache_are_preserved(tmp_path):
    formula = '=_xlfn.IFNA(IF(A1<9,"A&B","quoted ""text"""),"fallback")'
    source = book(tmp_path / "source.xlsx", formula)
    raw, _ = v8.screening_view(source)
    target = tmp_path / "execution-view.xlsx"
    target.write_bytes(raw)
    before, after = unpack(source.read_bytes()), unpack(raw)
    for name in before:
        assert (before[name].replace(b"_xlfn.IFNA(", b"IFNA(")
                if name == "xl/worksheets/sheet1.xml" else before[name]) == after[name]
    restored, _ = v8.comparison_view(source, target)
    assert unpack(restored) == before


def test_cleanup_failure_and_closed_unknown_receipts_are_preserved(tmp_path, monkeypatch):
    source = book(tmp_path / "source.xlsx")
    def fixture(self, path):
        return seal({"input_sha256": v5.sha(path), "identity": self.identity,
                     "status": "unknown", "reason": "container_cleanup_unconfirmed",
                     "cleanup_confirmed": False, "container_execution_attempted": True,
                     "duration_seconds": 0.25, "model_api_calls": 0})
    monkeypatch.setattr(v8.Recalculator, "_run_native", fixture)
    result = v8.Recalculator("sha256:" + "a" * 64).run(source)
    assert result["status"] == "unknown" and result["container_execution_attempted"]
    assert not result["cleanup_confirmed"] and result["duration_seconds"] == .25
    assert result["execution_receipt"]["reason"] == result["reason"]


@pytest.mark.parametrize("cleanup_code", [0, 1])
def test_native_driver_mounts_original_not_screening_bytes_and_always_cleans(tmp_path, monkeypatch, cleanup_code):
    source = book(tmp_path / "source.xlsx")
    original = source.read_bytes()
    image = "sha256:" + "a" * 64
    calls = []
    monkeypatch.setattr(v8.platform, "system", lambda: "Linux")
    monkeypatch.setattr(v8.shutil, "which", lambda name: "/test/docker")
    def command(argv, *args):
        calls.append(argv)
        if argv[1] == "image":
            return v5._Command(0, json.dumps({"Id": image, "Os": "linux", "Config": {}}).encode())
        if argv[1] == "rm":
            return v5._Command(cleanup_code)
        assert argv[1] == "run"
        for flag in ["--network=none", "--read-only", "--user=65534:65534", "--cap-drop=ALL",
                     "--security-opt=no-new-privileges", "--memory=1024m", "--pids-limit=128"]:
            assert flag in argv
        mounted = next(x for x in argv if x.startswith("type=bind,src=") and ",dst=/input,readonly" in x)
        input_dir = Path(mounted.split("src=", 1)[1].split(",dst=", 1)[0])
        assert (input_dir / "book.xlsx").read_bytes() == original
        assert json.loads((input_dir / "numeric_literals.json").read_bytes()) == v5._numeric_manifest(source)
        return v5._Command(1)  # No actual container; exercise cleanup after execution failure.
    monkeypatch.setattr(v5, "_bounded_command", command)
    result = v8.Recalculator(image).run(source)
    assert len(calls) == 3 and calls[-1][:3] == ["docker", "rm", "-f"]
    assert result["status"] == "unknown" and result["container_execution_attempted"]
    assert result["cleanup_confirmed"] is (cleanup_code == 0)
    assert source.read_bytes() == original
    assert result["execution_receipt"]["input_sha256"] == v5.sha(source)


@pytest.mark.parametrize("valid_proof", [True, False])
def test_native_driver_binds_worker_proof_to_original_input(tmp_path, monkeypatch, valid_proof):
    from tests.test_continual_eval_sheet_recalc import numeric_proof
    source = book(tmp_path / "source.xlsx")
    image = "sha256:" + "a" * 64
    monkeypatch.setattr(v8.platform, "system", lambda: "Linux")
    monkeypatch.setattr(v8.shutil, "which", lambda name: "/test/docker")
    def command(argv, *args):
        if argv[1] == "image":
            return v5._Command(0, json.dumps({"Id": image, "Os": "linux", "Config": {}}).encode())
        if argv[1] == "rm":
            return v5._Command(0)
        mount = next(x for x in argv if x.startswith("type=bind,src=") and ",dst=/input,readonly" in x)
        input_dir = Path(mount.split("src=", 1)[1].split(",dst=", 1)[0])
        manifest = json.loads((input_dir / "numeric_literals.json").read_bytes())
        assert manifest["input_sha256"] == v5.sha(source)
        response = {"status": "available", "output_base64": base64.b64encode(source.read_bytes()).decode(),
                    "output_sha256": v5.sha(source), "calculate_all": True, "macros": "never_execute",
                    "links": "no_update", "libreoffice_version": "fixture-not-live"}
        if valid_proof:
            response.update(numeric_proof(manifest))
        return v5._Command(0, json.dumps(response).encode())
    monkeypatch.setattr(v5, "_bounded_command", command)
    result = v8.Recalculator(image).run(source)
    assert result["cleanup_confirmed"]
    assert result["execution_input_sha256"] == v5.sha(source)
    if valid_proof:
        assert result["status"] == "available"
        assert result["execution_receipt"]["raw_export_sha256"] == v5.sha(source)
    else:
        assert result["status"] == "unknown" and result["reason"] == "invalid_recalculation_receipt"


def test_unsupported_native_result_cannot_be_accepted_as_case_spelling_fix(tmp_path, monkeypatch):
    source = book(tmp_path / "source.xlsx")
    after = book(tmp_path / "export.xlsx", "=ifna(A1,5)")
    _cache(after, "#NAME?", kind="e")
    def fixture(self, path):
        row = native_fixture(self, path)
        row.update(output_base64=base64.b64encode(after.read_bytes()).decode(), output_sha256=v5.sha(after))
        return seal({k: val for k, val in row.items() if k != "record_hash"})
    monkeypatch.setattr(v8.Recalculator, "_run_native", fixture)
    result = v8.Recalculator("sha256:" + "a" * 64).run(source)
    assert result["status"] == "unknown" and result["reason"] == "qualified_native_function_view_unverified"


def test_old_supported_path_uses_v7_unchanged(tmp_path, monkeypatch):
    source = book(tmp_path / "source.xlsx", "=SUM(A1:A5)")
    calls = []
    def fixture(self, path):
        calls.append(path)
        return {"unchanged_fixture": True}
    monkeypatch.setattr(v7.Recalculator, "run", fixture)
    assert v8.Recalculator("sha256:" + "a" * 64).run(source) == {"unchanged_fixture": True}
    assert calls == [source]
