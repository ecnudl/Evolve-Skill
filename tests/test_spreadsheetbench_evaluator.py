"""Evidence guards around the official cell-value SpreadsheetBench metric.

All workbooks here are authored fixtures, not natural model efficacy data.
No spreadsheet code, formula engine, external service, or agent is executed.
"""
from __future__ import annotations

import datetime
from unittest.mock import Mock

import openpyxl
import pytest

from skillopt.envs.spreadsheetbench import evaluator


def _book(values=None, sheet="Sheet"):
    workbook = openpyxl.Workbook()
    workbook.active.title = sheet
    for coordinate, value in (values or {}).items():
        workbook.active[coordinate] = value
    return workbook


def _save(tmp_path, name, values=None, sheet="Sheet"):
    path = tmp_path / name
    workbook = _book(values, sheet)
    workbook.save(path)
    workbook.close()
    return str(path)


def _evaluate(gold, pred, target):
    return evaluator.evaluate(pred, gold, "test-fixture", target)


@pytest.mark.parametrize("value,other,expected", [
    (7, 7, True), (7, 999, False), (7, "7.00", True),
    (None, None, True), (None, "", True), ("hello", "HELLO", False),
    (1.234, 1.231, True), (1.234, 1.24, False), (True, 1, True),
    (datetime.datetime(2020, 1, 1), 43831, True),
])
def test_original_cell_value_semantics_remain(value, other, expected):
    assert evaluator._compare_cell_value(value, other) is expected


def test_real_workbooks_correct_incorrect_and_absolute_references(tmp_path):
    gold = _save(tmp_path, "gold.xlsx", {"A1": 7, "B1": 9})
    pred = _save(tmp_path, "pred.xlsx", {"A1": 7, "B1": 999})
    assert evaluator.compare_workbooks(gold, pred, "$a$1") == (True, "")
    assert _evaluate(gold, pred, "A1")["status"] == "passed"
    result = _evaluate(gold, pred, "A1:B1")
    assert result["ok"] is False
    assert result["status"] == "failed"
    assert "value@Sheet!B1" in result["reason"]


@pytest.mark.parametrize("target", [
    "", " ", None, 42, ",", "A1,", ",A1", "A1,,B1", "A1, ,B1",
    "B2:A1", "B1:A2", "A2:B1", "A0", "A-1", "A01", "1A", "A", "1",
    "A1:B2:C3", "A1:", ":B2", "XFE1", "AAAA1", "A1048577",
    "A999999999999999999999999", "A1;B1", "!A1", "Sheet!", "Sheet!!A1",
    "'Sheet!A1", "Sheet!A1junk", "A1 B2", "Sheet!A1,'Other'!",
])
def test_invalid_ranges_never_pass_or_claim_semantic_failure(tmp_path, target):
    gold = _save(tmp_path, "gold.xlsx", {"A1": 7})
    pred = _save(tmp_path, "pred.xlsx", {"A1": 999})
    result = _evaluate(gold, pred, target)
    assert result["ok"] is False
    assert result["status"] == "invalid_spec"
    assert result["reason"].startswith("invalid_spec:")


def test_all_union_members_are_validated_before_a_value_mismatch(tmp_path):
    gold = _save(tmp_path, "gold.xlsx", {"A1": 7})
    pred = _save(tmp_path, "pred.xlsx", {"A1": 999})
    assert _evaluate(gold, pred, "A1,B2:A1")["status"] == "invalid_spec"
    assert _evaluate(gold, pred, "A1,Missing!A1")["status"] == "invalid_spec"


def test_multi_range_checks_do_not_skip_a_bad_value(tmp_path):
    gold = _save(tmp_path, "gold.xlsx", {"A1": 7, "B3": 8})
    pred = _save(tmp_path, "pred.xlsx", {"A1": 7, "B3": 0})
    assert _evaluate(gold, pred, "A1, B3")["status"] == "failed"
    assert evaluator.compare_workbooks(gold, gold, "A1, B3") == (True, "")


@pytest.mark.parametrize("sheet,target", [
    ("Data Sheet", "'Data Sheet'!A1"),
    ("Revenue, Net", "'Revenue, Net'!A1"),
    ("Budget's", "'Budget''s'!A1"),
    ("Hello!", "'Hello!'!A1"),
    ("Sheet", '"Sheet"!"A1"'),
])
def test_quoted_sheet_names_and_ranges(tmp_path, sheet, target):
    gold = _save(tmp_path, "gold.xlsx", {"A1": 7}, sheet)
    pred = _save(tmp_path, "pred.xlsx", {"A1": 7}, sheet)
    assert evaluator.compare_workbooks(gold, pred, target) == (True, "")


def test_missing_files_and_sheets_are_not_semantic_value_failures(tmp_path):
    gold = _save(tmp_path, "gold.xlsx", {"A1": 7}, "Gold")
    pred = _save(tmp_path, "pred.xlsx", {"A1": 7}, "Other")
    missing = str(tmp_path / "missing.xlsx")
    assert _evaluate(gold, missing, "A1")["status"] == "missing_output"
    assert _evaluate(missing, pred, "A1")["status"] == "unknown"
    assert _evaluate(gold, pred, "A1")["status"] == "missing_output"
    assert _evaluate(gold, pred, "Unknown!A1")["status"] == "invalid_spec"
    assert _evaluate(missing, missing, "")["status"] == "invalid_spec"


def test_actual_empty_cells_can_still_pass(tmp_path):
    gold = _save(tmp_path, "gold.xlsx")
    pred = _save(tmp_path, "pred.xlsx")
    assert _evaluate(gold, pred, "A1:B2")["status"] == "passed"


@pytest.mark.parametrize("gold_value,pred_value,source", [
    ("=1+1", "=2", "reference"),
    (None, "=1+1", "prediction"),
    ("=1+1", None, "reference"),
    (2, "=1+1", "prediction"),
    (None, '=IF(TRUE,"",2)', "prediction"),
])
def test_missing_formula_results_cannot_become_blank_equality(
    tmp_path, gold_value, pred_value, source
):
    # openpyxl writes formulas but does not calculate cached results.
    gold = _save(tmp_path, "gold.xlsx", {"A1": gold_value})
    pred = _save(tmp_path, "pred.xlsx", {"A1": pred_value})
    result = _evaluate(gold, pred, "A1")
    assert result["ok"] is False
    assert result["status"] == "unknown"
    assert f"formula result missing in {source}" in result["reason"]
    assert "recalculation required (not performed)" in result["reason"]


def test_missing_formula_evidence_and_confirmed_mismatch_are_both_retained(tmp_path):
    gold = _save(tmp_path, "gold.xlsx", {"A1": 7, "B1": 2})
    pred = _save(tmp_path, "pred.xlsx", {"A1": 999, "B1": "=1+1"})
    result = _evaluate(gold, pred, "A1,B1")
    assert result["status"] == "failed"
    assert "value@Sheet!A1" in result["reason"]
    assert "partial_unavailable_cells=1" in result["reason"]
    assert "formula result missing" in result["reason"]
    assert _evaluate(gold, pred, "B1,A1")["status"] == "failed"


def test_unchecked_formula_is_not_an_extra_task_obligation(tmp_path):
    gold = _save(tmp_path, "gold.xlsx", {"A1": 7})
    pred = _save(tmp_path, "pred.xlsx", {"A1": 7, "B1": "=1+1"})
    assert _evaluate(gold, pred, "A1")["status"] == "passed"


@pytest.mark.parametrize("pred_cache,expected_status", [(2, "passed"), (99, "failed")])
def test_alternative_formulas_compare_cached_values_not_formula_text(
    tmp_path, monkeypatch, pred_cache, expected_status
):
    gold = _save(tmp_path, "gold.xlsx")
    pred = _save(tmp_path, "pred.xlsx")
    views = [
        _book({"A1": 2}), _book({"A1": pred_cache}),
        _book({"A1": "=1+1"}), _book({"A1": "=SUM(1,1)"}),
    ]
    for book in views:
        book.close = Mock(wraps=book.close)
    loader = Mock(side_effect=views)
    monkeypatch.setattr(evaluator.openpyxl, "load_workbook", loader)
    assert _evaluate(gold, pred, "A1")["status"] == expected_status
    assert [call.kwargs["data_only"] for call in loader.call_args_list] == [
        True, True, False, False,
    ]
    for book in views:
        book.close.assert_called_once()


def test_load_failure_closes_already_opened_workbooks(tmp_path, monkeypatch):
    gold = _save(tmp_path, "gold.xlsx")
    pred = _save(tmp_path, "pred.xlsx")
    opened = _book()
    opened.close = Mock(wraps=opened.close)
    loader = Mock(side_effect=[opened, OSError("fixture read failure")])
    monkeypatch.setattr(evaluator.openpyxl, "load_workbook", loader)
    result = _evaluate(gold, pred, "A1")
    assert result["status"] == "unknown"
    assert "workbook load error" in result["reason"]
    opened.close.assert_called_once()


def test_cell_generator_validates_and_keeps_original_column_major_order():
    assert evaluator._generate_cell_names("$a$1:b2") == ["A1", "A2", "B1", "B2"]
    assert evaluator._generate_cell_names("XFD1048576") == ["XFD1048576"]
    for invalid in ("B2:A1", "", "XFE1", "A1048577"):
        with pytest.raises(ValueError):
            evaluator._generate_cell_names(invalid)


def test_direct_cell_helper_no_longer_succeeds_vacuously():
    gold, pred = _book({"A1": 7}), _book({"A1": 999})
    try:
        ok, reason = evaluator._cell_level_compare(gold, pred, "Sheet", "B2:A1")
        assert ok is False
        assert reason.startswith("invalid_spec:")
    finally:
        gold.close()
        pred.close()


def test_repeat_comparison_is_deterministic_and_does_not_modify_workbooks(tmp_path):
    gold = _save(tmp_path, "gold.xlsx", {"A1": 7})
    pred = _save(tmp_path, "pred.xlsx", {"A1": 999})
    from pathlib import Path
    before = (Path(gold).read_bytes(), Path(pred).read_bytes())
    first = _evaluate(gold, pred, "A1")
    assert _evaluate(gold, pred, "A1") == first
    assert (Path(gold).read_bytes(), Path(pred).read_bytes()) == before
