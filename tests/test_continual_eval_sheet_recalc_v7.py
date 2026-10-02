"""Engineering controls only; no model data or new engine qualification."""
import base64
import hashlib
import io
import json
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pytest
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from skillopt.coevolution_v5.core import seal
from skillopt.continual_eval import sheet_recalc as v5
from skillopt.continual_eval import sheet_recalc_v6 as v6
from skillopt.continual_eval import sheet_recalc_v7 as v7
from skillopt.continual_eval import spreadsheet_compat as compat


def _cache(path, text, kind="n", coordinate="B1"):
    source = io.BytesIO(path.read_bytes())
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(path, "w") as target:
        for info in archive.infolist():
            content = archive.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                tree = ET.fromstring(content)
                cell = next(c for c in tree.iter(f"{{{compat.MAIN_NS}}}c") if c.get("r") == coordinate)
                cell.set("t", kind)
                cell.find(f"{{{compat.MAIN_NS}}}v").text = text
                content = ET.tostring(tree)
            target.writestr(info, content)


def _book(path, *, exported=False, formula="=A1+1", cache="46001",
          source_format="General", formula_format=None, epoch=None):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet["A1"] = 46000
    if exported:
        sheet["A1"].number_format = "m/d/yyyy"
    else:
        sheet["A1"].number_format = "0"
        sheet["A1"]._style.numFmtId = 30
    sheet["B1"] = formula
    sheet["B1"].number_format = formula_format or ("mm/dd/yy" if exported else source_format)
    # These genuinely date-styled cells must keep their read semantics.
    sheet["C1"] = 46003
    sheet["C1"].number_format = "yyyy-mm-dd"
    sheet["D1"] = "=DATE(2025,12,12)"
    sheet["D1"].number_format = "yyyy-mm-dd"
    if epoch:
        book.epoch = epoch
    book.save(path)
    book.close()
    if exported:
        _cache(path, cache)
        _cache(path, "46003", coordinate="D1")
    return path


def _view(tmp_path, before, after):
    raw, proof = v7.comparison_view(before, after)
    target = tmp_path / "view.xlsx"
    target.write_bytes(raw)
    return target, proof


def _fake(before, after, identity):
    manifest = v5._numeric_manifest(before)
    encoded = json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
    return seal({"status": "available", "reason": "fixture", "identity": identity,
        "input_sha256": v5.sha(before), "output_sha256": v5.sha(after),
        "output_base64": base64.b64encode(after.read_bytes()).decode(),
        "numeric_manifest_sha256": hashlib.sha256(encoded).hexdigest(),
        "numeric_literals_verified": len(manifest["cells"]),
        "cleanup_confirmed": True, "model_api_calls": 0})


@pytest.mark.parametrize("formula,cache,expected", [("=A1+1", "46001", 46001), ("=A1+2", "46002", 46002)])
def test_q20_shaped_formula_autoformat_and_wrong_result(tmp_path, formula, cache, expected):
    before = _book(tmp_path / "before.xlsx", formula=formula)
    after = _book(tmp_path / "after.xlsx", exported=True, formula=formula, cache=cache)
    originals = before.read_bytes(), after.read_bytes()
    view, proof = _view(tmp_path, before, after)
    assert v5._recalc_valid(before, view) is None
    with zipfile.ZipFile(after) as a, zipfile.ZipFile(view) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name == "xl/worksheets/sheet1.xml":
                assert (a.read(name).replace(b'r="A1" s="1"', b'r="A1" s="0"')
                        .replace(b'r="B1" s="2"', b'r="B1" s="0"')) == b.read(name)
            else:
                assert a.read(name) == b.read(name)
    book = openpyxl.load_workbook(view, data_only=True)
    assert book["S"]["A1"].value == 46000
    assert book["S"]["B1"].value == expected
    assert book["S"]["C1"].is_date
    assert book["S"]["D1"].is_date
    book.close()
    assert len(proof["adapted_cells"]) == len(proof["formula_adapted_cells"]) == 1
    assert not proof["formula_result_correctness_proven"]
    assert originals == (before.read_bytes(), after.read_bytes())
    assert (expected == 46001) == (formula == "=A1+1")  # Wrong output is not repaired.


def test_v6_and_v7_do_not_share_scope(tmp_path):
    before = _book(tmp_path / "before.xlsx")
    after = _book(tmp_path / "after.xlsx", exported=True)
    raw, old_proof = v6.comparison_view(before, after)
    old = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    assert old["S"]["B1"].is_date
    old.close()
    assert old_proof["version"] != v7.VIEW_VERSION
    image = "sha256:" + "a" * 64
    assert v7.Recalculator(image).identity != v6.Recalculator(image).identity


@pytest.mark.parametrize("source_format", ["yyyy-mm-dd", "mm/dd/yy", "0", "0.00"])
def test_original_date_formula_and_non_general_style_not_adapted(tmp_path, source_format):
    before = _book(tmp_path / "before.xlsx", source_format=source_format)
    after = _book(tmp_path / "after.xlsx", exported=True)
    view, proof = _view(tmp_path, before, after)
    assert not proof["formula_adapted_cells"]
    book = openpyxl.load_workbook(view, data_only=True)
    assert book["S"]["B1"].is_date
    book.close()


def test_reserved_unknown_formula_format_is_not_general_fallback(tmp_path):
    before = _book(tmp_path / "before.xlsx")
    book = openpyxl.load_workbook(before)
    book["S"]["B1"].number_format = "0"
    book["S"]["B1"]._style.numFmtId = 30
    book.save(before)
    book.close()
    after = _book(tmp_path / "after.xlsx", exported=True)
    _, proof = _view(tmp_path, before, after)
    assert not proof["formula_adapted_cells"]


@pytest.mark.parametrize("formula", ["=A1+2", "=DATE(2025,12,10)", "=46001"])
def test_changed_source_formula_is_not_authorized(tmp_path, formula):
    before = _book(tmp_path / "before.xlsx")
    after = _book(tmp_path / "after.xlsx", exported=True, formula=formula)
    with pytest.raises(ValueError, match="General formula changed"):
        v7.comparison_view(before, after)


@pytest.mark.parametrize("kind,text", [("str", "46001"), ("b", "1"), ("e", "#VALUE!"),
                                      ("n", None), ("n", "NaN"), ("n", "1e999")])
def test_missing_nonnumeric_or_nonfinite_formula_cache_not_authorized(tmp_path, kind, text):
    before = _book(tmp_path / "before.xlsx")
    after = _book(tmp_path / "after.xlsx", exported=True)
    _cache(after, text, kind)
    with pytest.raises((ValueError, OverflowError), match="numeric cache|NaN|infinity|int"):
        v7.comparison_view(before, after)


@pytest.mark.parametrize("formula", ['=CELL("format",A1)', '=TEXT(A1,"yyyy")', '=FORMULATEXT(A1)'])
def test_format_observer_blocks_literal_and_formula_views(tmp_path, formula):
    before = _book(tmp_path / "before.xlsx", formula=formula)
    after = _book(tmp_path / "after.xlsx", exported=True, formula=formula)
    view, proof = _view(tmp_path, before, after)
    assert view.read_bytes() == after.read_bytes()
    assert not proof["adapted_cells"] and not proof["formula_adapted_cells"]
    assert v5._recalc_valid(before, view) == "recalculation_changed_literal_cell"


def test_epoch_change_not_authorized(tmp_path):
    before = _book(tmp_path / "before.xlsx")
    after = _book(tmp_path / "after.xlsx", exported=True, epoch=CALENDAR_MAC_1904)
    view, proof = _view(tmp_path, before, after)
    assert view.read_bytes() == after.read_bytes()
    assert not proof["formula_adapted_cells"]


def test_full_run_retains_engine_export_and_numeric_proof_limits(tmp_path, monkeypatch):
    before = _book(tmp_path / "before.xlsx")
    after = _book(tmp_path / "after.xlsx", exported=True)
    monkeypatch.setattr(v5.Recalculator, "run", lambda self, path: _fake(path, after, self.identity))
    result = v7.Recalculator("sha256:" + "a" * 64).run(before)
    assert result["status"] == "available"
    assert result["engine_export_sha256"] == v5.sha(after)
    assert result["engine_export_base64"] == base64.b64encode(after.read_bytes()).decode()
    assert result["comparison_view"]["version"] == v7.VIEW_VERSION
    assert not result["comparison_view"]["formula_result_correctness_proven"]
    assert result["output_sha256"] != result["engine_export_sha256"]
    assert result["model_api_calls"] == 0


def test_run_changed_formula_returns_unknown(tmp_path, monkeypatch):
    before = _book(tmp_path / "before.xlsx")
    after = _book(tmp_path / "after.xlsx", exported=True, formula="=A1+2")
    monkeypatch.setattr(v5.Recalculator, "run", lambda self, path: _fake(path, after, self.identity))
    result = v7.Recalculator("sha256:" + "a" * 64).run(before)
    assert result["status"] == "unknown"
    assert result["reason"] == "general_formula_read_view_unverified"


def test_wrong_numeric_cache_is_not_repaired_into_expected(tmp_path, monkeypatch):
    reference = _book(tmp_path / "reference.xlsx")
    _cache(reference, "46001")  # Original public reference cache, not injected into the view.
    prediction = _book(tmp_path / "prediction.xlsx")
    correct_export = _book(tmp_path / "correct-export.xlsx", exported=True)
    wrong_export = _book(tmp_path / "wrong-export.xlsx", exported=True, cache="46002")
    exports = {v5.sha(reference): correct_export, v5.sha(prediction): wrong_export}
    monkeypatch.setattr(v5.Recalculator, "run",
                        lambda self, path: _fake(path, exports[v5.sha(path)], self.identity))
    engine = v7.Recalculator("sha256:" + "a" * 64)
    result = v5.evaluate_pair(prediction, reference, "B1", engine, tmp_path / "receipts")
    assert result["status"] == "fail"
    assert len(result["receipts"]) == 2
    wrong = engine.run(prediction)
    values = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wrong["output_base64"])), data_only=True)
    assert values["S"]["B1"].value == 46002
    values.close()


@pytest.mark.parametrize("damage", ["duplicate_cell", "shared_formula", "duplicate_formula"])
def test_source_formula_xml_binding_is_required(tmp_path, damage):
    before = _book(tmp_path / "before.xlsx")
    after = _book(tmp_path / "after.xlsx", exported=True)
    source = io.BytesIO(before.read_bytes())
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(before, "w") as target:
        for item in archive.infolist():
            content = archive.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                if damage == "duplicate_cell":
                    content = content.replace(b"</row>", b'<c r="B1"><f>A1+1</f><v/></c></row>')
                elif damage == "duplicate_formula":
                    content = content.replace(b"<f>A1+1</f>", b"<f>A1+1</f><f>A1+1</f>")
                else:
                    content = content.replace(b"<f>A1+1</f>", b'<f t="shared" si="0" ref="B1:B2">A1+1</f>')
            target.writestr(item, content)
    with pytest.raises(ValueError, match="Duplicate|XML binding"):
        v7.comparison_view(before, after)


def test_no_docker_stays_unknown(tmp_path, monkeypatch):
    monkeypatch.setattr(v5.platform, "system", lambda: "Darwin")
    result = v7.Recalculator("sha256:" + "a" * 64).run(_book(tmp_path / "in.xlsx"))
    assert result["status"] == "unknown" and result["reason"] == "linux_docker_unavailable"
