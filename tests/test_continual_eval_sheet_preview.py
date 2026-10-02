"""Stable public workbook preview fixtures; no model or container execution."""
import builtins
import json
import re
from copy import deepcopy
from datetime import date, datetime, time, timedelta

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from skillopt.continual_eval import backends as b
from skillopt.continual_eval.core import freeze_plan, load_plan
from skillopt.continual_eval.fixtures import fixture_config

PROFILE = {"spreadsheet_preview_version": b.SHEET_PREVIEW_VERSION}


def book(tmp_path, *, formulas=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Public input"
    ws["A1"], ws["A2"], ws["A3"] = 3, "Text", None
    if formulas:
        ws["B1"] = ArrayFormula(ref="B1:B2", text="=A1:A2*2")
        ws["C1"] = DataTableFormula(ref="C1:D2", dt2D=True, r1="A1", r2="A2")
    path = tmp_path/"input.xlsx"
    wb.save(path)
    wb.close()
    return path


def test_read_only_array_and_data_table_formulas_have_stable_public_fields(tmp_path):
    path = book(tmp_path, formulas=True)
    before = path.read_bytes()
    first = b._sheet_preview(path, version=b.SHEET_PREVIEW_VERSION)
    second = b._sheet_preview(path, version=b.SHEET_PREVIEW_VERSION)
    assert json.dumps(first, ensure_ascii=False) == json.dumps(second, ensure_ascii=False)
    assert path.read_bytes() == before
    assert "object at 0x" not in json.dumps(first)
    array = first[0]["preview"][0][1]
    assert array == {"kind": "array_formula", "text": "=A1:A2*2", "ref": "B1:B2",
                     "truncated_fields": [], "unsupported_fields": [], "missing_fields": []}
    table = first[0]["preview"][0][2]
    assert table == {"kind": "data_table_formula", "ref": "C1:D2", "r1": "A1", "r2": "A2",
                     "ca": False, "dt2D": True, "dtr": False, "del1": False, "del2": False,
                     "truncated_fields": [], "unsupported_fields": [], "missing_fields": []}
    legacy = b._sheet_preview(path)
    assert re.fullmatch(r"<openpyxl\.worksheet\.formula\.ArrayFormula object at 0x[0-9a-fA-F]+>",
                        legacy[0]["preview"][0][1])
    assert re.fullmatch(r"<openpyxl\.worksheet\.formula\.DataTableFormula object at 0x[0-9a-fA-F]+>",
                        legacy[0]["preview"][0][2])


@pytest.mark.parametrize("value", [None, "", "Plain text", "界" * 201, "=SUM(A1:A2)",
    42, -3, 0.125, float("inf"), float("nan"), True, False, date(2026, 10, 1),
    datetime(2026, 10, 1, 7, 30), time(7, 30), timedelta(days=2, seconds=3)])
def test_supported_scalars_keep_exact_legacy_display(value):
    assert b._sheet_public_value(value) == (str(value)[:200] if value is not None else None)


class UntrustedValue:
    def __str__(self):
        raise AssertionError("Custom str must never execute")

    def __repr__(self):
        raise AssertionError("Custom repr must never execute")


def test_unknown_object_never_invokes_custom_stringification():
    assert b._sheet_public_value(UntrustedValue()) == {
        "kind": "unsupported_cell", "reason": "unsupported_public_value_type"}


def test_metaclass_equality_cannot_impersonate_a_builtin_scalar():
    class EqualToEverything(type):
        def __eq__(cls, other):
            raise AssertionError("Type whitelist must not invoke custom equality")

    class CustomScalar(UntrustedValue, metaclass=EqualToEverything):
        pass

    assert b._sheet_public_value(CustomScalar())["kind"] == "unsupported_cell"


def test_subclass_cannot_smuggle_custom_formula_fields_or_representation():
    class CustomFormula(ArrayFormula):
        def __str__(self):
            raise AssertionError("Subclass str must never execute")
    value = CustomFormula(ref="PRIVATE_REF", text="PRIVATE_TEXT")
    assert b._sheet_public_value(value)["kind"] == "unsupported_cell"
    assert "PRIVATE" not in json.dumps(b._sheet_public_value(value))


@pytest.mark.parametrize("length,expected", [(200, []), (201, ["text", "ref"])])
def test_array_text_and_reference_budgets_are_explicit(length, expected):
    value = ArrayFormula(ref="R" * length, text="界" * length)
    value.private_oracle = "PRIVATE_ORACLE"
    result = b._sheet_public_value(value)
    assert result["text"] == "界" * min(length, 200)
    assert result["ref"] == "R" * min(length, 200)
    assert result["truncated_fields"] == expected
    assert "PRIVATE" not in json.dumps(result)


def test_missing_and_unsupported_formula_fields_are_not_fabricated_or_stringified():
    array = b._sheet_public_value(ArrayFormula(ref=UntrustedValue(), text=None))
    assert array["ref"] is None and array["text"] is None
    assert array["unsupported_fields"] == ["ref"] and array["missing_fields"] == ["text"]
    table = b._sheet_public_value(DataTableFormula(ref="A" * 201, r1="B" * 201,
        r2=UntrustedValue(), ca=UntrustedValue(), dt2D=1))
    assert table["truncated_fields"] == ["ref", "r1"]
    assert table["unsupported_fields"] == ["r2", "ca", "dt2D"]
    assert table["r2"] is None and table["ca"] is None and table["dt2D"] is None


@pytest.mark.parametrize("value,expected", [("1", True), ("0", False), ("true", True), ("false", False)])
def test_data_table_accepts_only_explicit_xml_boolean_lexical_values(value, expected):
    result = b._sheet_public_value(DataTableFormula(ref="A1", dt2D=value))
    assert result["dt2D"] is expected and result["unsupported_fields"] == []


@pytest.mark.parametrize("value", ["yes", "TRUE", "false trailing", 1, 0])
def test_invalid_data_table_boolean_does_not_gain_truth_from_coercion(value):
    result = b._sheet_public_value(DataTableFormula(ref="A1", dt2D=value))
    assert result["dt2D"] is None and result["unsupported_fields"] == ["dt2D"]


def test_default_preview_structure_and_scalar_prompt_bytes_are_unchanged(tmp_path, monkeypatch):
    path = book(tmp_path)
    assert b._sheet_preview(path) == b._sheet_preview(path, version=b.SHEET_PREVIEW_VERSION)
    observed, calls = [], []
    monkeypatch.setattr(b, "_native", lambda req, runtime: calls.append(req) or {"status": "fixture"})

    def call(system, user):
        observed.append((system, user))
        return {"ok": True, "response": "```python\npass\n```", "usage": {}, "finish_reason": "stop"}

    public = {"instruction": "Keep existing data", "input_files": [str(path)], "answer_position": "A1"}
    old = b.solve("spreadsheetbench", public, "", call)
    new = b.solve("spreadsheetbench", public, "", call, runtime=PROFILE)
    assert old == new and observed[0] == observed[1] and calls[0] == calls[1]
    assert observed[0] == ("Solve the supplied task using only the supplied public information."
        "\nReturn one complete Python program in a ```python code block.", json.dumps({
        "instruction": "Keep existing data", "workbook_preview": b._sheet_preview(path),
        "input_path": "input.xlsx", "required_output_path": "output.xlsx",
        "required_answer_position": "A1"}, ensure_ascii=False))


def test_solver_sees_public_formula_fields_not_private_contract_or_gold(tmp_path, monkeypatch):
    path = book(tmp_path, formulas=True)
    prompts, requests = [], []
    monkeypatch.setattr(b, "_native", lambda req, runtime: requests.append(req) or {"status": "fixture"})

    def call(system, user):
        prompts.append((system, user))
        return {"ok": True, "response": "```python\npass\n```", "usage": {}, "finish_reason": "stop"}

    public = {"instruction": "Keep formulas", "input_files": [str(path)], "answer_position": "B1",
              "test_files": ["PRIVATE_GOLD_PATH"], "reference": "PRIVATE_ORACLE"}
    prediction = b.solve("spreadsheetbench", public, "Frozen guidance", call, runtime=PROFILE)
    assert prediction["status"] == "available" and len(prompts) == len(requests) == 1
    text = prompts[0][1]
    assert "=A1:A2*2" in text and "B1:B2" in text and "object at 0x" not in text
    assert "PRIVATE" not in text + json.dumps(requests)
    assert set(requests[0]) == {"operation", "code", "input_base64"}


@pytest.mark.parametrize("version", [True, False, 1, "unknown-v2", {}, []])
def test_unknown_profile_fails_before_reading_files_or_calling_solver(monkeypatch, version):
    monkeypatch.setattr(b, "_xlsx_bytes", lambda _: pytest.fail("Must not read workbook"))
    monkeypatch.setattr(b, "_native", lambda *_: pytest.fail("Must not launch native readiness probe"))
    assert b.readiness("spreadsheetbench", {"spreadsheet_preview_version": version}) == {
        "status": "unsupported", "reason": "invalid_spreadsheet_preview_version"}
    with pytest.raises(ValueError, match="preview version"):
        b._sheet_preview("unused", version=version)
    prediction = b.solve("spreadsheetbench", {"instruction": "x", "input_files": ["unused"],
        "answer_position": "A1"}, "", lambda *_: pytest.fail("Must not call model"),
        runtime={"spreadsheet_preview_version": version})
    assert prediction["status"] == "unknown" and prediction["reason"] == "invalid_spreadsheet_preview_version"


def test_missing_formula_capability_is_unsupported_and_never_starts_native_or_model(tmp_path, monkeypatch):
    path = book(tmp_path)
    original_import = builtins.__import__

    def missing(name, *args, **kwargs):
        if name == "openpyxl.worksheet.formula":
            raise ImportError("Authored missing formula class")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", missing)
    monkeypatch.setattr(b, "_native", lambda *_: pytest.fail("No native probe"))
    assert b.readiness("spreadsheetbench", PROFILE) == {
        "status": "unsupported", "reason": "public_formula_preview_classes_unavailable"}
    assert b._sheet_preview(path)[0]["preview"][0][0] == "3"  # Legacy does not need this capability.
    result = b.solve("spreadsheetbench", {"instruction": "x", "input_files": [str(path)],
        "answer_position": "A1"}, "", lambda *_: pytest.fail("No model call"), runtime=PROFILE)
    assert result["status"] == "unknown" and result["reason"] == "public_input_or_runtime_error:ValueError"


def test_version_is_frozen_in_existing_plan_identity(tmp_path):
    config = fixture_config(tmp_path/"data")
    old = freeze_plan(config, tmp_path/"old")
    updated = deepcopy(config)
    updated["runtime"]["spreadsheetbench"] = {**updated["runtime"].get("spreadsheetbench", {}), **PROFILE}
    new = freeze_plan(updated, tmp_path/"new")
    assert new == load_plan(tmp_path/"new") and old["record_hash"] != new["record_hash"]
    assert new["config"]["runtime"]["spreadsheetbench"]["spreadsheet_preview_version"] == b.SHEET_PREVIEW_VERSION
    assert "spreadsheet_preview_version" not in old["config"]["runtime"].get("spreadsheetbench", {})


def test_original_sheet_row_column_limits_are_unchanged(tmp_path):
    wb = openpyxl.Workbook()
    for index in range(20):
        wb.create_sheet(f"Extra{index}")
    for sheet in wb.worksheets:
        sheet["U6"] = "OUTSIDE_PREVIEW"
    path = tmp_path/"bounded.xlsx"
    wb.save(path)
    wb.close()
    result = b._sheet_preview(path, version=b.SHEET_PREVIEW_VERSION)
    assert len(result) == 20
    assert all(len(sheet["preview"]) == 5 and all(len(row) == 20 for row in sheet["preview"]) for sheet in result)
    assert "OUTSIDE_PREVIEW" not in json.dumps(result)
