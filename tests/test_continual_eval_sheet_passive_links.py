"""Offline passive-link boundary tests, not a Calc qualification."""
import base64
import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pytest

from skillopt.coevolution_v5.core import seal
from skillopt.continual_eval import sheet_frozen_gold as scoring
from skillopt.continual_eval import sheet_passive_links as links
from skillopt.continual_eval import sheet_recalc as v5
from skillopt.continual_eval import sheet_recalc_functions as v8
from tests.test_continual_eval_sheet_recalc_v7 import _cache

IMAGE = "sha256:" + "a" * 64
SHEET = "xl/worksheets/sheet1.xml"
RELS = "xl/worksheets/_rels/sheet1.xml.rels"


def fixture(path, *, formula="=A1+1", cache="3", target="https://example.org/help"):
    book = openpyxl.Workbook()
    book.active.title = "S"
    book.active["A1"], book.active["B1"], book.active["C1"] = 2, formula, "Documentation"
    book.active["C1"].hyperlink = target
    book.active["C1"].hyperlink.tooltip = "Read documentation"
    book.save(path)
    book.close()
    _cache(path, cache)
    return path


def rewrite(path, member, edit):
    with zipfile.ZipFile(path) as z:
        data = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w") as z:
        for item, raw in data:
            if item.filename == member:
                root = ET.fromstring(raw)
                edit(root)
                raw = ET.tostring(root)
            z.writestr(item, raw)


def native_fixture(self, path):
    proof = links._checked_screen(path)
    manifest = v5._numeric_manifest(path)
    return seal({"status": "available", "reason": "fixture", "input_sha256": v5.sha(path),
                 "identity": self.identity, "passive_hyperlinks": proof,
                 "output_base64": base64.b64encode(path.read_bytes()).decode(), "output_sha256": v5.sha(path),
                 "numeric_literals_verified": len(manifest["cells"]),
                 "numeric_manifest_sha256": hashlib.sha256(json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()).hexdigest(),
                 "duration_seconds": 0.1, "cleanup_confirmed": True,
                 "container_execution_attempted": True, "model_api_calls": 0})


def test_screen_copy_only_links_change_original_formula_cache_and_bytes_remain(tmp_path):
    source = fixture(tmp_path / "original.xlsx")
    before = source.read_bytes()
    screened, proof = links.screening_view(source)
    checked = tmp_path / "never-execute.xlsx"
    checked.write_bytes(screened)
    assert v5.screen(source) == "unsupported_external_relationship"
    assert v5.screen(checked) is None
    assert proof["execution_input_sha256"] == proof["original_input_sha256"] == v5.sha(source)
    assert proof["screening_view_sha256"] != v5.sha(source)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(checked) as copy:
        a = ET.fromstring(original.read(SHEET)).find(links.MAIN + "sheetData")
        b = ET.fromstring(copy.read(SHEET)).find(links.MAIN + "sheetData")
        assert ET.tostring(a) == ET.tostring(b)
    assert source.read_bytes() == before


@pytest.mark.parametrize("url", ["http://example.org", "file:///etc/passwd", "javascript:alert(1)",
    "https://user:password@example.org/", "https://example.org:22/a", "https://example.org/a b",
    "https://example.org\\x", "https:///missing", "https://example.org:bad"])
def test_nonqualified_uri_never_executes(tmp_path, monkeypatch, url):
    source = fixture(tmp_path / "input.xlsx", target=url)
    monkeypatch.setattr(links.platform, "system", lambda: "Linux")
    monkeypatch.setattr(links.shutil, "which", lambda _: "docker")
    monkeypatch.setattr(v5, "_bounded_command", lambda *a: pytest.fail("No command may run"))
    row = links.Recalculator(IMAGE).run(source)
    assert row["status"] == "unknown" and not row["container_execution_attempted"]


@pytest.mark.parametrize("change", ["kind", "unused", "duplicate_id", "duplicate_cell", "wrong_root",
    "nested", "range", "missing_cell", "location", "internal", "bad_namespace", "tail", "outside"])
def test_ambiguous_relationship_or_cell_is_rejected(tmp_path, change):
    path = fixture(tmp_path / "input.xlsx")
    def rel_edit(root):
        item = root[0]
        if change == "kind":
            item.set("Type", links.DOC_REL + "/externalLinkPath")
        elif change == "unused":
            item.set("Id", "unused")
        elif change == "duplicate_id":
            root.append(ET.fromstring(ET.tostring(item)))
        elif change == "wrong_root":
            root.tag = "Relationships"
        elif change == "nested":
            ET.SubElement(item, "payload")
        elif change == "internal":
            item.set("TargetMode", "Internal")
        elif change == "bad_namespace":
            item.tag = "Relationship"
    def sheet_edit(root):
        container = root.find(links.MAIN + "hyperlinks")
        item = container[0]
        if change == "duplicate_cell":
            container.append(ET.fromstring(ET.tostring(item)))
        elif change == "range":
            item.set("ref", "C1:C2")
        elif change == "missing_cell":
            item.set("ref", "C2")
        elif change == "location":
            item.set("location", "S!A1")
        elif change == "tail":
            container.tail = "unexpected-payload"
        elif change == "outside":
            root.find(".//" + links.MAIN + "c[@r='C1']").set("r", "XFE1")
            item.set("ref", "XFE1")
    rewrite(path, RELS, rel_edit)
    rewrite(path, SHEET, sheet_edit)
    with pytest.raises(ValueError):
        links.screening_view(path)


@pytest.mark.parametrize("formula", ["=TODAY()", "=_xlfn.UNIQUE(A1:A2)", "='[x.xlsx]Sheet'!A1"])
def test_links_do_not_hide_other_unsupported_dependencies(tmp_path, formula):
    path = fixture(tmp_path / "input.xlsx", formula=formula)
    with pytest.raises(ValueError, match="unsupported dependency"):
        links._checked_screen(path)


@pytest.mark.parametrize("change", ["target", "tooltip", "display", "cell"])
def test_export_link_semantic_changes_are_rejected(tmp_path, change):
    source, exported = fixture(tmp_path / "source.xlsx"), fixture(tmp_path / "export.xlsx")
    if change == "target":
        rewrite(exported, RELS, lambda root: root[0].set("Target", "https://different.example/"))
    else:
        rewrite(exported, SHEET, lambda root: root.find(links.MAIN + "hyperlinks")[0].set(
            {"tooltip": "tooltip", "display": "display", "cell": "ref"}[change], "A1" if change == "cell" else "changed"))
    with pytest.raises(ValueError, match="semantics changed"):
        links.verify_preserved(source, exported)


def test_transport_id_can_change_without_rewriting_original_or_cached_values(tmp_path):
    source, exported = fixture(tmp_path / "source.xlsx"), fixture(tmp_path / "export.xlsx", cache="99")
    rewrite(exported, RELS, lambda root: root[0].set("Id", "newId"))
    rewrite(exported, SHEET, lambda root: root.find(links.MAIN + "hyperlinks")[0].set(links.RID, "newId"))
    assert links.verify_preserved(source, exported)["hyperlink_count"] == 1
    values = openpyxl.load_workbook(exported, data_only=True)
    assert values.active["B1"].value == 99
    values.close()


def test_native_receipt_binds_original_bytes_and_full_content_guard(tmp_path, monkeypatch):
    path = fixture(tmp_path / "input.xlsx")
    original = path.read_bytes()
    monkeypatch.setattr(links.Recalculator, "_run_passive_native", native_fixture)
    row = links.Recalculator(IMAGE).run(path)
    assert row["status"] == "available" and row["execution_input_sha256"] == v5.sha(path)
    assert row["hyperlinks_preserved"]["hyperlink_count"] == 1
    assert row["execution_receipt"]["input_sha256"] == v5.sha(path)
    assert path.read_bytes() == original


def test_changed_literal_is_not_accepted_just_because_links_preserved(tmp_path, monkeypatch):
    path = fixture(tmp_path / "input.xlsx")
    other = fixture(tmp_path / "other.xlsx")
    rewrite(other, SHEET, lambda root: root.find(".//" + links.MAIN + "c[@r='A1']/" + links.MAIN + "v").__setattr__("text", "8"))
    def modified(self, source):
        row = native_fixture(self, source)
        return seal({**{k: v for k, v in row.items() if k != "record_hash"},
                     "output_base64": base64.b64encode(other.read_bytes()).decode(), "output_sha256": v5.sha(other)})
    monkeypatch.setattr(links.Recalculator, "_run_passive_native", modified)
    result = links.Recalculator(IMAGE).run(path)
    assert result["status"] == "unknown" and result["execution_receipt"]["status"] == "available"


@pytest.mark.parametrize("actual,expected", [("3", "pass"), ("4", "fail")])
def test_same_link_wrong_value_still_fails_frozen_h(tmp_path, monkeypatch, actual, expected):
    reference = fixture(tmp_path / "reference.xlsx", cache="3")
    prediction = fixture(tmp_path / "prediction.xlsx", cache=actual)
    monkeypatch.setattr(links.Recalculator, "_run_passive_native", native_fixture)
    manifest = scoring.freeze_reference(reference, v5.sha(reference), "B1", provenance={
        "source_protocol_hash": "fixture", "task_id": "fixture", "evidence_kind": "engineering_fixture"})
    row = scoring.evaluate(prediction, reference, "B1", links.Recalculator(IMAGE), manifest)
    assert row["status"] == expected


def test_changed_formula_rejected_even_when_link_and_cache_match(tmp_path, monkeypatch):
    source = fixture(tmp_path / "source.xlsx")
    exported = fixture(tmp_path / "exported.xlsx", formula="=A1+2", cache="3")
    def altered(self, path):
        row = native_fixture(self, path)
        return seal({**{k: v for k, v in row.items() if k != "record_hash"},
                     "output_sha256": v5.sha(exported), "output_base64": base64.b64encode(exported.read_bytes()).decode()})
    monkeypatch.setattr(links.Recalculator, "_run_passive_native", altered)
    assert links.Recalculator(IMAGE).run(source)["status"] == "unknown"


def _drop_tooltip(path):
    rewrite(path, SHEET, lambda root: root.find(links.MAIN + "hyperlinks")[0].attrib.pop("tooltip", None))


@pytest.mark.parametrize("text", ["Read documentation", 'Text > < & "quotes"', "First\nSecond\t说明"])
def test_missing_original_tooltip_restored_without_any_formula_cache_or_value_edits(tmp_path, text):
    source = fixture(tmp_path / "source.xlsx")
    rewrite(source, SHEET, lambda root: root.find(links.MAIN + "hyperlinks")[0].set("tooltip", text))
    exported = fixture(tmp_path / "exported.xlsx", cache="99")
    _drop_tooltip(exported)
    original_bytes, exported_bytes = source.read_bytes(), exported.read_bytes()
    raw, proof = links.tooltip_comparison_view(source, exported)
    restored = tmp_path / "view.xlsx"
    restored.write_bytes(raw)
    assert links.verify_preserved(source, restored)["hyperlink_count"] == 1
    assert proof["tooltip_attributes_restored"] == 1
    assert proof["formula_or_value_bytes_changed"] is False
    with zipfile.ZipFile(exported) as old, zipfile.ZipFile(restored) as new:
        for member in old.namelist():
            a, b = old.read(member), new.read(member)
            if member == SHEET:
                # Literal caches and all cells are byte-exact, not merely equal
                # after openpyxl reparsing; only hyperlink start-tag grew.
                closing = re.search(rb"</(?:[A-Za-z_][A-Za-z0-9_.-]*:)?sheetData\s*>", a)
                assert closing is not None
                cell_end = closing.end()
                assert a[:cell_end] == b[:cell_end]
            else:
                assert a == b
    assert source.read_bytes() == original_bytes and exported.read_bytes() == exported_bytes
    values = openpyxl.load_workbook(restored, data_only=True)
    assert values.active["B1"].value == 99
    values.close()


@pytest.mark.parametrize("change", ["tooltip_changed", "tooltip_added", "display", "target", "observer", "name"])
def test_tooltip_view_cannot_hide_other_changes_or_metadata_observers(tmp_path, change):
    source, exported = fixture(tmp_path / "source.xlsx"), fixture(tmp_path / "export.xlsx")
    if change == "tooltip_changed":
        rewrite(exported, SHEET, lambda root: root.find(links.MAIN + "hyperlinks")[0].set("tooltip", "wrong"))
    elif change == "tooltip_added":
        _drop_tooltip(source)
    else:
        _drop_tooltip(exported)
        if change == "display":
            rewrite(exported, SHEET, lambda root: root.find(links.MAIN + "hyperlinks")[0].set("display", "different"))
        elif change == "target":
            rewrite(exported, RELS, lambda root: root[0].set("Target", "https://different.example/"))
        else:
            from openpyxl.workbook.defined_name import DefinedName
            for path in (source, exported):
                wb = openpyxl.load_workbook(path)
                if change == "observer":
                    wb.active["D1"] = '=CELL("contents",C1)'
                else:
                    wb.defined_names.add(DefinedName("Observer", attr_text="S!A1"))
                wb.save(path)
                wb.close()
    with pytest.raises(ValueError):
        links.tooltip_comparison_view(source, exported)


def display_fixture(path, *, exported=False, cache="3", text="Documentation", target="https://example.org/help"):
    fixture(path, cache=cache, target=target)
    def change(root):
        link = root.find(links.MAIN + "hyperlinks")[0]
        link.attrib.pop("tooltip", None)
        link.set("display", text if exported else target)
        root.find(".//" + links.MAIN + "c[@r='C1']/" + links.MAIN + "is/" + links.MAIN + "t").text = text
    rewrite(path, SHEET, change)
    return path


@pytest.mark.parametrize("text", ["Documentation", 'Title > < & "quotes"', "表格链接\n下一行"])
def test_display_href_to_unchanged_literal_restores_only_property_bytes(tmp_path, text):
    target = "https://example.org/help?a=1&b=2"
    source = display_fixture(tmp_path / "source.xlsx", text=text, target=target)
    exported = display_fixture(tmp_path / "export.xlsx", exported=True, cache="99", text=text, target=target)
    a, b = source.read_bytes(), exported.read_bytes()
    raw, proof = links.display_comparison_view(source, exported)
    view = tmp_path / "view.xlsx"
    view.write_bytes(raw)
    assert links.verify_preserved(source, view)["hyperlink_count"] == 1
    assert proof["display_attributes_restored"] == 1 and proof["formula_or_value_bytes_changed"] is False
    with zipfile.ZipFile(exported) as old, zipfile.ZipFile(view) as new:
        for name in old.namelist():
            left, right = old.read(name), new.read(name)
            if name == SHEET:
                # Exact deletion of the changed display token recovers the
                # raw engine XML; all cells, formulas and caches stay byte exact.
                def strip(value):
                    return re.sub(rb"\sdisplay=([\"']).*?\1", b"", value)
                assert strip(left) == strip(right)
            else:
                assert left == right
    assert source.read_bytes() == a and exported.read_bytes() == b
    book = openpyxl.load_workbook(view, data_only=True)
    assert book.active["B1"].value == 99 and book.active["C1"].value == text
    book.close()


@pytest.mark.parametrize("change", [
    "original_label_not_href", "new_display_arbitrary", "new_display_missing", "href", "cell_text",
    "missing_payload", "numeric_cell", "formula_cell", "cell_metadata", "array_overlap", "observer", "name",
    "tooltip_changed", "coordinate", "duplicate_payload", "table_overlap",
])
def test_display_read_view_rejects_any_other_semantic_change(tmp_path, change):
    source = display_fixture(tmp_path / "source.xlsx")
    exported = display_fixture(tmp_path / "export.xlsx", exported=True)
    def mutate(root):
        link = root.find(links.MAIN + "hyperlinks")[0]
        cell = root.find(".//" + links.MAIN + "c[@r='C1']")
        if change == "original_label_not_href":
            link.set("display", "Original custom label")
        elif change == "new_display_arbitrary":
            link.set("display", "Other text")
        elif change == "new_display_missing":
            link.attrib.pop("display")
        elif change == "cell_text":
            cell.find(links.MAIN + "is/" + links.MAIN + "t").text = "Changed visible text"
        elif change == "missing_payload":
            cell.remove(cell.find(links.MAIN + "is"))
        elif change == "duplicate_payload":
            ET.SubElement(ET.SubElement(cell, links.MAIN + "is"), links.MAIN + "t").text = "Documentation"
        elif change == "numeric_cell":
            cell.attrib["t"] = "n"
            cell.remove(cell.find(links.MAIN + "is"))
            ET.SubElement(cell, links.MAIN + "v").text = "5"
        elif change == "formula_cell":
            ET.SubElement(cell, links.MAIN + "f").text = '"Documentation"'
        elif change == "cell_metadata":
            cell.set("cm", "1")
        elif change == "array_overlap":
            formula = root.find(".//" + links.MAIN + "c[@r='B1']/" + links.MAIN + "f")
            formula.set("t", "array")
            formula.set("ref", "B1:C1")
        elif change == "tooltip_changed":
            link.set("tooltip", "New help")
        elif change == "coordinate":
            link.set("ref", "A1")
    if change == "href":
        rewrite(exported, RELS, lambda root: root[0].set("Target", "https://example.org/changed"))
    elif change in {"observer", "name", "table_overlap"}:
        from openpyxl.workbook.defined_name import DefinedName
        for path in (source, exported):
            wb = openpyxl.load_workbook(path)
            if change == "observer":
                wb.active["D1"] = '=CELL("contents",C1)'
            elif change == "name":
                wb.defined_names.add(DefinedName("Observer", attr_text="S!A1"))
            else:
                wb.active.add_table(openpyxl.worksheet.table.Table(displayName="Labels", ref="C1:C2"))
            wb.save(path)
            wb.close()
    else:
        rewrite(source if change == "original_label_not_href" else exported, SHEET, mutate)
    with pytest.raises(ValueError):
        links.display_comparison_view(source, exported)


@pytest.mark.parametrize("cache,expected", [("3", "pass"), ("4", "fail")])
def test_display_adaptation_keeps_frozen_h_correct_and_wrong_distinct(tmp_path, monkeypatch, cache, expected):
    reference = fixture(tmp_path / "reference.xlsx", cache="3")
    prediction = display_fixture(tmp_path / "prediction.xlsx", cache=cache)
    exported = display_fixture(tmp_path / "export.xlsx", exported=True, cache=cache)
    original = prediction.read_bytes()
    def transformed(self, path):
        receipt = native_fixture(self, path)
        return seal({**{k: v for k, v in receipt.items() if k != "record_hash"},
                     "output_base64": base64.b64encode(exported.read_bytes()).decode(), "output_sha256": v5.sha(exported)})
    monkeypatch.setattr(links.Recalculator, "_run_passive_native", transformed)
    manifest = scoring.freeze_reference(reference, v5.sha(reference), "B1", provenance={
        "source_protocol_hash": "fixture", "task_id": "fixture", "evidence_kind": "engineering_fixture"})
    row = scoring.evaluate(prediction, reference, "B1", links.Recalculator(IMAGE), manifest)
    assert row["status"] == expected and prediction.read_bytes() == original


def test_display_and_tooltip_restoration_compose_without_changing_formula_guard(tmp_path, monkeypatch):
    source = display_fixture(tmp_path / "source.xlsx")
    rewrite(source, SHEET, lambda root: root.find(links.MAIN + "hyperlinks")[0].set("tooltip", "Original help"))
    exported = display_fixture(tmp_path / "export.xlsx", exported=True)
    def transformed(self, path):
        receipt = native_fixture(self, path)
        return seal({**{k: v for k, v in receipt.items() if k != "record_hash"},
                     "output_base64": base64.b64encode(exported.read_bytes()).decode(), "output_sha256": v5.sha(exported)})
    monkeypatch.setattr(links.Recalculator, "_run_passive_native", transformed)
    row = links.Recalculator(IMAGE).run(source)
    assert row["status"] == "available"
    assert row["display_comparison_view"]["display_attributes_restored"] == 1
    assert row["tooltip_comparison_view"]["tooltip_attributes_restored"] == 1
    rewrite(exported, SHEET, lambda root: root.find(".//" + links.MAIN + "f").__setattr__("text", "A1+2"))
    assert links.Recalculator(IMAGE).run(source)["status"] == "unknown"


def test_nonlink_scope_delegates_to_frozen_parent(tmp_path, monkeypatch):
    path = fixture(tmp_path / "input.xlsx")
    calls = []
    monkeypatch.setattr(v5, "screen", lambda p: None)
    monkeypatch.setattr(v8.Recalculator, "run", lambda self, p: calls.append(p) or {"status": "sentinel"})
    assert links.Recalculator(IMAGE).run(path) == {"status": "sentinel"} and calls == [path]


def test_worker_mount_is_original_and_failure_cleans_up(tmp_path, monkeypatch):
    path = fixture(tmp_path / "input.xlsx")
    calls = []
    monkeypatch.setattr(links.platform, "system", lambda: "Linux")
    monkeypatch.setattr(links.shutil, "which", lambda _: "docker")
    def command(args, *unused):
        calls.append(args)
        if args[1:3] == ["image", "inspect"]:
            return v5._Command(0, stdout=json.dumps({"Id": IMAGE, "Os": "linux", "Config": {}}))
        if args[1] == "run":
            assert "--network=none" in args and "--read-only" in args
            mount = next(a for a in args if a.startswith("type=bind,src=") and "dst=/input," in a)
            root = Path(mount.split("src=", 1)[1].split(",", 1)[0])
            assert (root / "book.xlsx").read_bytes() == path.read_bytes()
            assert v5.screen(root / "book.xlsx") == "unsupported_external_relationship"
            return v5._Command(1)
        assert args[1:3] == ["rm", "-f"]
        return v5._Command(0)
    monkeypatch.setattr(v5, "_bounded_command", command)
    result = links.Recalculator(IMAGE).run(path)
    assert result["status"] == "unknown" and result["cleanup_confirmed"]
    assert len(calls) == 3 and result["container_execution_attempted"]
