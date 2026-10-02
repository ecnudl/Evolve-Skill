"""Offline controls; no model calls, no LibreOffice required for this suite."""
import base64
import hashlib
import io
import json
import math
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pytest

from skillopt.coevolution_v5.core import seal
from skillopt.continual_eval import sheet_recalc as recalc
from skillopt.continual_eval import sheet_recalc_worker as worker
from skillopt.continual_eval import spreadsheet_compat as compat
from skillopt.continual_eval.core import read_json


def workbook(path, formula="=SUM(A1:A2)", cache=None, *, cache_type="n", extras=None):
    book = openpyxl.Workbook()
    book.active.title = "S"
    book.active["A1"], book.active["A2"], book.active["B1"] = 2, 3, formula
    for cell, value in (extras or {}).items():
        book.active[cell] = value
    raw = io.BytesIO()
    book.save(raw)
    book.close()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(path, "w") as dest:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml" and cache is not None:
                root = ET.fromstring(data)
                cell = root.find(f'.//{{{compat.MAIN_NS}}}c[@r="B1"]')
                cell.set("t", cache_type)
                value = cell.find(f"{{{compat.MAIN_NS}}}v")
                if value is None:
                    value = ET.SubElement(cell, f"{{{compat.MAIN_NS}}}v")
                value.text = str(cache)
                data = ET.tostring(root)
            dest.writestr(item, data)
    return path


class Fake:
    identity = {"fixture": True}

    def __init__(self, mapping):
        self.mapping, self.calls = mapping, []

    def run(self, path):
        self.calls.append(recalc.sha(path))
        raw = self.mapping[recalc.sha(path)].read_bytes()
        return seal({"status": "available", "reason": "fixture", "identity": self.identity,
            "input_sha256": recalc.sha(path), "output_sha256": hashlib.sha256(raw).hexdigest(),
            "output_base64": base64.b64encode(raw).decode(), "cleanup_confirmed": True,
            "model_api_calls": 0})


@pytest.mark.parametrize("formula,reason", [
    ("=SUM(A1:A2)", None),
    ('=IF(A1=2,"",1)', None),
    ('=WEBSERVICE("https://example.invalid")', "unsupported_external_structured_or_volatile_formula"),
    ("=NOW()", "unsupported_external_structured_or_volatile_formula"),
    ("='[book.xlsx]S'!A1", "unsupported_external_structured_or_volatile_formula"),
    ("=_xlfn.MYFUNCTION(A1)", "unsupported_extended_function"),
])
def test_formula_safety_scope(tmp_path, formula, reason):
    assert recalc.screen(workbook(tmp_path / "in.xlsx", formula)) == reason


@pytest.mark.parametrize("name,body,reason", [
    ("xl/vbaProject.bin", b"macro", "unsupported_active_or_external_workbook"),
    ("xl/connections.xml", b"<x/>", "unsupported_active_or_external_workbook"),
    ("xl/_rels/external.rels", b'<Relationships><Relationship TargetMode="External" Target="file:///etc/passwd"/></Relationships>',
     "unsupported_external_relationship"),
    ("bad.xml", b'<!DOCTYPE x [<!ENTITY unsafe "v">]><x/>', "unsupported_workbook_structure"),
])
def test_active_content_rejected_before_engine(tmp_path, name, body, reason):
    path = workbook(tmp_path / "in.xlsx")
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr(name, body)
    assert recalc.screen(path) == reason


def test_pair_independent_values_cache_replay_and_no_original_mutation(tmp_path):
    gold = workbook(tmp_path / "gold.xlsx", cache=5)
    prediction = workbook(tmp_path / "pred.xlsx", "=A1+A2")
    converted = workbook(tmp_path / "calc.xlsx", "=A1+A2", cache=5)
    before = gold.read_bytes(), prediction.read_bytes()
    engine = Fake({recalc.sha(gold): gold, recalc.sha(prediction): converted})
    first = recalc.evaluate_pair(prediction, gold, "B1", engine, tmp_path / "cache")
    assert first["status"] == "pass" and len(engine.calls) == 2
    assert recalc.evaluate_pair(prediction, gold, "B1", engine, tmp_path / "cache") == first
    assert len(engine.calls) == 2
    assert before == (gold.read_bytes(), prediction.read_bytes())


def test_wrong_formula_is_not_repaired_with_gold(tmp_path):
    gold = workbook(tmp_path / "gold.xlsx", cache=5)
    prediction = workbook(tmp_path / "pred.xlsx", "=A1+A2+1")
    converted = workbook(tmp_path / "calc.xlsx", "=A1+A2+1", cache=6)
    engine = Fake({recalc.sha(gold): gold, recalc.sha(prediction): converted})
    assert recalc.evaluate_pair(prediction, gold, "B1", engine, tmp_path / "cache")["status"] == "fail"


def test_native_reference_drift_cannot_become_model_failure(tmp_path):
    gold = workbook(tmp_path / "gold.xlsx", cache=10)
    converted_gold = workbook(tmp_path / "gold-calc.xlsx", cache=5)
    prediction = workbook(tmp_path / "pred.xlsx", "=A1+A2", cache=5)
    engine = Fake({recalc.sha(gold): converted_gold, recalc.sha(prediction): prediction})
    score = recalc.evaluate_pair(prediction, gold, "B1", engine, tmp_path / "cache")
    assert score["status"] == "unknown" and "drift" in score["reason"]


def test_empty_string_is_real_computed_cache(tmp_path):
    gold = workbook(tmp_path / "gold.xlsx", '=IF(1,"",2)', cache="", cache_type="str")
    prediction = workbook(tmp_path / "pred.xlsx", '=IF(A1=2,"",1)')
    converted = workbook(tmp_path / "calc.xlsx", '=IF(A1=2,"",1)', cache="", cache_type="str")
    engine = Fake({recalc.sha(gold): gold, recalc.sha(prediction): converted})
    result = recalc.evaluate_pair(prediction, gold, "B1", engine, tmp_path / "cache")
    assert result["status"] == "pass"
    assert result["evidence"]["prediction_empty_string_caches_used"] == 1


def test_unsupported_function_and_removed_formula_stay_unknown(tmp_path):
    original = workbook(tmp_path / "orig.xlsx", "=SKILLOPT_NOT_A_FUNCTION(A1)")
    error = workbook(tmp_path / "error.xlsx", "=SKILLOPT_NOT_A_FUNCTION(A1)", cache="#NAME?", cache_type="e")
    constant = workbook(tmp_path / "constant.xlsx", 5)
    assert recalc._recalc_valid(original, error) == "unsupported_formula_result"
    assert recalc._recalc_valid(original, constant) == "recalculation_removed_formula"


def test_no_host_fallback_and_no_mutable_image(tmp_path, monkeypatch):
    with pytest.raises(ValueError, match="Pinned"):
        recalc.Recalculator("libreoffice:latest")
    monkeypatch.setattr(recalc.platform, "system", lambda: "Darwin")
    result = recalc.Recalculator("sha256:" + "a" * 64).run(workbook(tmp_path / "book.xlsx"))
    assert result["status"] == "unknown" and result["reason"] == "linux_docker_unavailable"


def test_docker_isolation_and_cleanup_boundaries(tmp_path, monkeypatch):
    import json

    from skillopt.skill_validation.sandbox import _Command
    path = workbook(tmp_path / "in.xlsx")
    image = "sha256:" + "a" * 64
    commands = []
    monkeypatch.setattr(recalc.platform, "system", lambda: "Linux")
    monkeypatch.setattr(recalc.shutil, "which", lambda _: "/usr/bin/docker")

    def command(argv, timeout, *args):
        commands.append(argv)
        if argv[1:3] == ["image", "inspect"]:
            return _Command(0, json.dumps({"Id": image, "Os": "linux", "Config": {}}).encode())
        if argv[1] == "run":
            assert "--network=none" in argv and "--read-only" in argv
            assert "--user=65534:65534" in argv and "--cap-drop=ALL" in argv
            assert "--pull=never" in argv and "--security-opt=no-new-privileges" in argv
            assert sum("readonly" in arg for arg in argv) == 2
            return _Command(None, timed_out=True)
        return _Command(0)

    monkeypatch.setattr(recalc, "_bounded_command", command)
    result = recalc.Recalculator(image).run(path)
    assert result["reason"] == "recalculation_timeout" and result["cleanup_confirmed"]
    assert commands[-1][1:3] == ["rm", "-f"]


def test_lo_must_not_change_literal_cells(tmp_path):
    before = workbook(tmp_path / "before.xlsx", cache=5)
    after = workbook(tmp_path / "after.xlsx", cache=5)
    book = openpyxl.load_workbook(after)
    book["S"]["A1"] = 100
    book.save(after)
    book.close()
    assert recalc._recalc_valid(before, after) == "recalculation_changed_literal_cell"


def test_cleanup_exception_retained_as_unknown(tmp_path, monkeypatch):
    import json
    import subprocess

    from skillopt.skill_validation.sandbox import _Command
    image = "sha256:" + "a" * 64
    monkeypatch.setattr(recalc.platform, "system", lambda: "Linux")
    monkeypatch.setattr(recalc.shutil, "which", lambda _: "/usr/bin/docker")

    def command(argv, *args):
        if argv[1:3] == ["image", "inspect"]:
            return _Command(0, json.dumps({"Id": image, "Os": "linux", "Config": {}}).encode())
        raise subprocess.TimeoutExpired(argv[:2], 10)

    monkeypatch.setattr(recalc, "_bounded_command", command)
    result = recalc.Recalculator(image).run(workbook(tmp_path / "in.xlsx"))
    assert result["status"] == "unknown" and not result["cleanup_confirmed"]
    assert result["reason"] == "container_cleanup_unconfirmed"


def test_defined_name_and_malformed_output_are_unknown(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName
    path = workbook(tmp_path / "in.xlsx")
    book = openpyxl.load_workbook(path)
    book.defined_names.add(DefinedName("HiddenNetwork", attr_text='WEBSERVICE("https://example.invalid")'))
    book.save(path)
    book.close()
    assert recalc.screen(path) == "unsupported_external_or_volatile_defined_name"
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"not a workbook")
    assert recalc._recalc_valid(path, broken) == "recalculation_invalid_workbook_structure"


def test_recalculation_cannot_change_defined_name_or_literal_to_formula(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName
    before = workbook(tmp_path / "before.xlsx", 5)
    after = workbook(tmp_path / "after.xlsx", "=SUM(A1:A2)", cache=5)
    assert recalc._recalc_valid(before, after) == "recalculation_changed_literal_cell"
    after = workbook(tmp_path / "after-name.xlsx", 5)
    book = openpyxl.load_workbook(after)
    book.defined_names.add(DefinedName("AddedName", attr_text="S!A1"))
    book.save(after)
    book.close()
    assert recalc._recalc_valid(before, after) == "recalculation_changed_defined_names"


def test_array_formula_support_keeps_exact_text_and_range(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    formula = ArrayFormula(ref="B1", text="=SUM(A1:A2*2)")
    original = workbook(tmp_path / "array.xlsx", formula)
    calculated = workbook(tmp_path / "array-calc.xlsx", formula, cache=10)
    assert recalc.screen(original) is None
    assert recalc._recalc_valid(original, calculated) is None
    changed = workbook(tmp_path / "changed.xlsx", ArrayFormula(ref="B1", text="=SUM(A1:A2*2)+1"), cache=11)
    assert recalc._recalc_valid(original, changed) == "recalculation_changed_formula"
    range_changed = workbook(tmp_path / "range.xlsx", ArrayFormula(ref="B1:B2", text="=SUM(A1:A2*2)"), cache=10)
    assert recalc._recalc_valid(original, range_changed) == "recalculation_changed_array_range"
    external = workbook(tmp_path / "external-array.xlsx", ArrayFormula(ref="B1", text='=WEBSERVICE("https://example.invalid")'))
    assert "unsupported" in recalc.screen(external)


def test_array_single_cell_range_spelling_is_not_a_semantic_change(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    text = "=SUM(A1:A2*2)+1"
    original = workbook(tmp_path / "before.xlsx", ArrayFormula(ref="B1", text=text))
    converted = workbook(tmp_path / "after.xlsx", ArrayFormula(ref="B1:B1", text=text), cache=11)
    assert recalc._formula_spec(ArrayFormula(ref="$b$1", text=text)) == recalc._formula_spec(
        ArrayFormula(ref="B1:B1", text=text))
    assert recalc._recalc_valid(original, converted) is None
    shifted = workbook(tmp_path / "shifted.xlsx", ArrayFormula(ref="B1:B2", text=text), cache=11)
    assert recalc._recalc_valid(original, shifted) == "recalculation_changed_array_range"
    changed = workbook(tmp_path / "changed-text.xlsx", ArrayFormula(ref="B1:B1", text="=SUM(A1:A2*2)"), cache=10)
    assert recalc._recalc_valid(original, changed) == "recalculation_changed_formula"


def test_wrong_array_formula_fails_without_changing_input(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    gold = workbook(tmp_path / "gold-array.xlsx", ArrayFormula(ref="B1", text="=SUM(A1:A2*2)"), cache=10)
    prediction = workbook(tmp_path / "wrong-array.xlsx", ArrayFormula(ref="B1", text="=SUM(A1:A2*2)+1"))
    calculated = workbook(tmp_path / "wrong-array-calc.xlsx", ArrayFormula(ref="B1", text="=SUM(A1:A2*2)+1"), cache=11)
    engine = Fake({recalc.sha(gold): gold, recalc.sha(prediction): calculated})
    assert recalc.evaluate_pair(prediction, gold, "B1", engine, tmp_path / "cache")["status"] == "fail"


def test_fixed_array_derived_cells_can_change_but_not_unrelated_cells(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    original = workbook(tmp_path / "array.xlsx", ArrayFormula(ref="B1:B2", text="=A1:A2*2"))
    computed = workbook(tmp_path / "array-calc.xlsx", ArrayFormula(ref="B1:B2", text="=A1:A2*2"), cache=4, extras={"B2": 6})
    assert recalc._recalc_valid(original, computed) is None
    book = openpyxl.load_workbook(computed)
    book["S"]["A2"] = 99
    book.save(computed)
    book.close()
    assert recalc._recalc_valid(original, computed) == "recalculation_changed_literal_cell"


@pytest.mark.parametrize("follower,reason", [("#NAME?", "unsupported_formula_result"),
                                           (None, "recalculation_array_result_unavailable")])
def test_array_followers_missing_or_unsupported_are_unknown(tmp_path, follower, reason):
    from openpyxl.worksheet.formula import ArrayFormula
    formula = ArrayFormula(ref="B1:B2", text="=A1:A2*2")
    original = workbook(tmp_path / "original.xlsx", formula)
    after = workbook(tmp_path / "after.xlsx", formula, cache=4, extras={"B2": follower})
    assert recalc._recalc_valid(original, after) == reason


def test_engine_cannot_repair_wrong_formula_or_add_new_semantics(tmp_path):
    gold = workbook(tmp_path / "gold.xlsx", cache=5)
    original = workbook(tmp_path / "wrong.xlsx", "=A1+A2+1")
    repaired = workbook(tmp_path / "repaired.xlsx", "=A1+A2", cache=5)
    engine = Fake({recalc.sha(gold): gold, recalc.sha(original): repaired})
    score = recalc.evaluate_pair(original, gold, "B1", engine, tmp_path / "cache")
    assert score["status"] == "unknown" and score["reason"] == "prediction:recalculation_changed_formula"
    book = openpyxl.load_workbook(gold)
    book["S"]["D1"] = 10
    changed = tmp_path / "added.xlsx"
    book.save(changed)
    book.close()
    assert recalc._recalc_valid(gold, changed) == "recalculation_added_cell_content"


def source_panel(tmp_path):
    from skillopt.continual_eval import runner
    from skillopt.continual_eval.core import BENCHMARKS, freeze_plan, load_checkpoint, panel_tasks, write_json
    from skillopt.continual_eval.fixtures import fixture_panel
    from skillopt.validator_pilot.api import digest
    source = tmp_path / "original"
    gold = workbook(tmp_path / "gold.xlsx", cache=5)
    input_path = workbook(tmp_path / "input.xlsx", 0)
    panel = fixture_panel("spreadsheetbench")
    task = panel["tasks"][0]
    task["partition"] = "development"
    task["private"] = {"test_files": [str(gold)], "answer_position": "B1", "asset_sha256": {str(gold): recalc.sha(gold)}}
    task["public"] = {"instruction": "Fixture", "input_files": [str(input_path)], "answer_position": "B1"}
    path = tmp_path / "panel.json"
    write_json(path, panel)
    config = {"version": "continual-eval-v1", "order": list(BENCHMARKS),
        "panels": {b: str(path) if b == "spreadsheetbench" else None for b in BENCHMARKS}, "partition": "development",
        "methods": ["no_skill"], "histories": ["h0"], "repeats": 2,
        "model": {"provider": "fixture", "name": "fixture", "max_tokens": 4096, "reasoning_effort": "low"},
        "runtime": {}, "project_disjoint": False, "exposure_manifest": None}
    plan = freeze_plan(config, source)
    checkpoint = load_checkpoint(source, "no_skill", "h0", 0, plan)
    task = panel_tasks(plan, "spreadsheetbench")[0]
    for repeat in range(2):
        base, request = runner.position(source, checkpoint, "spreadsheetbench", task, repeat)
        write_json(base / "intent.json", seal(request))
        pred = {"status": "available" if repeat == 0 else "unknown", "reason": "fixture_missing",
                "output": {"cases": [{"status": "available", "output_base64": base64.b64encode(gold.read_bytes()).decode()}]}
                if repeat == 0 else None}
        record = seal({"request": request, "prediction": pred})
        write_json(base / "prediction.json", record)
        write_json(source / "host_only/scores" / (digest(request) + ".json"), seal({
            "benchmark": "spreadsheetbench", "prediction_hash": record["record_hash"],
            "task_id": task["task_id"], "repeat": repeat, "checkpoint_hash": checkpoint["record_hash"],
            "plan_hash": plan["record_hash"], "status": "pass" if repeat == 0 else "unknown"}))
    engine = recalc.Recalculator("sha256:" + "a" * 64)
    qualification = tmp_path / "qualification.json"
    write_json(qualification, seal({"version": recalc.VERSION, "status": "qualified", "engine": engine.identity}))
    return source, qualification, gold


def test_full_denominator_unknowns_and_idempotent_offline_replay(tmp_path, monkeypatch):
    source, qualification, gold = source_panel(tmp_path)
    output = tmp_path / "rescore"
    assert recalc.prepare(source, output, qualification)["positions"] == 2
    calls = []

    def fake_run(self, path):
        calls.append(path)
        raw = gold.read_bytes()
        return seal({"status": "available", "reason": "fixture", "identity": self.identity,
                     "input_sha256": recalc.sha(path), "output_sha256": recalc.sha(gold),
                     "output_base64": base64.b64encode(raw).decode(), "cleanup_confirmed": True})

    monkeypatch.setattr(recalc.Recalculator, "run", fake_run)
    first = recalc.run(output)
    assert first["positions"] == first["completed"] == 2
    assert first["counts"] == {"pass": 1, "unknown": 1}
    assert first["model_api_calls"] == 0 and not first["historical_scores_replaced"]
    assert len(calls) == 1  # Identical workbook bytes are recalculated once.
    assert recalc.run(output) == first and len(calls) == 1


def test_source_mutation_blocks_before_execution(tmp_path):
    source, qualification, gold = source_panel(tmp_path)
    output = tmp_path / "rescore"
    recalc.prepare(source, output, qualification)
    workbook(gold, cache=100)
    with pytest.raises(ValueError, match="Frozen source evidence"):
        recalc.run(output)


def test_failed_qualification_and_nested_output_rejected(tmp_path):
    source, qualification, _ = source_panel(tmp_path)
    with pytest.raises(ValueError, match="Separate"):
        recalc.prepare(source, source / "new", qualification)
    from skillopt.continual_eval.core import write_json
    failed = read_json(qualification, sealed=True)
    failed.pop("record_hash")
    failed["status"] = "rejected"
    path = tmp_path / "failed.json"
    write_json(path, seal(failed))
    with pytest.raises(ValueError, match="Successful qualification"):
        recalc.prepare(source, tmp_path / "rescore", path)


def numeric_proof(manifest):
    return {"numeric_manifest_sha256": hashlib.sha256(
        json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()).hexdigest(),
        "numeric_verified_before": len(manifest["cells"]), "numeric_verified_after": len(manifest["cells"])}


def test_numeric_manifest_excludes_formulas_array_results_and_other_types(tmp_path):
    import datetime

    from openpyxl.worksheet.formula import ArrayFormula
    path = workbook(tmp_path / "manifest.xlsx", ArrayFormula(ref="B1:B2", text="=A1:A2*2"),
                    cache=4, extras={"B2": 6, "C1": True, "D1": "3", "E1": datetime.date(2020, 1, 1),
                                     "F1": "=1+2", "G1": 1.25e-20})
    manifest = recalc._numeric_manifest(path)
    assert {(item["sheet"], item["coordinate"]) for item in manifest["cells"]} == {
        ("S", "A1"), ("S", "A2"), ("S", "G1")}
    assert manifest["input_sha256"] == recalc.sha(path)
    assert next(item for item in manifest["cells"] if item["coordinate"] == "G1")["binary64"] == float(1.25e-20).hex()


class NumericDocument:
    """Only read-only UNO calls supported; any attempted setter fails the test."""
    def __init__(self, cells):
        self.cells = cells

    def getSheets(self):
        return self

    def hasByName(self, name):
        return name == "S"

    def getByName(self, name):
        assert name == "S"
        return self

    def getCellByPosition(self, column, row):
        from types import SimpleNamespace
        kind, number = self.cells.get((column, row), ("EMPTY", 0.0))
        return SimpleNamespace(getType=lambda: SimpleNamespace(value=kind), getValue=lambda: number)


@pytest.mark.parametrize("actual,kind,reason", [
    (1.0, "VALUE", "value_changed"), (float("inf"), "VALUE", "value_changed"),
    (1.0000000000000002, "FORMULA", "type_changed"), (0, "EMPTY", "type_changed"),
])
def test_threshold_sensitive_input_cannot_be_restored_after_import_rounding(actual, kind, reason):
    value = math.nextafter(1.0, 2.0)
    # Native answer tolerance would equate these, yet a branch can differ.
    assert compat._legacy()._compare_cell_value(value, 1.0)
    assert (value > 1.0) != (1.0 > 1.0)
    manifest = {"cells": [{"sheet": "S", "column": 0, "row": 0, "binary64": value.hex()}]}
    for stage in ("before_calculate", "after_calculate"):
        with pytest.raises(worker.NumericProofError, match=reason):
            worker.verify_numeric_literals(NumericDocument({(0, 0): (kind, actual)}), manifest, stage)


def test_numeric_proof_requires_exact_finite_bits_and_correct_sheet():
    manifest = {"cells": [{"sheet": "S", "column": 0, "row": 0, "binary64": (-0.0).hex()}]}
    with pytest.raises(worker.NumericProofError, match="value_changed"):
        worker.verify_numeric_literals(NumericDocument({(0, 0): ("VALUE", 0.0)}), manifest, "before_calculate")
    assert worker.verify_numeric_literals(NumericDocument({(0, 0): ("VALUE", -0.0)}), manifest, "after_calculate") == 1
    manifest["cells"][0]["sheet"] = "missing"
    with pytest.raises(worker.NumericProofError, match="sheet_missing"):
        worker.verify_numeric_literals(NumericDocument({}), manifest, "before_calculate")


def test_only_verified_literal_export_text_changes_no_formula_or_cache(tmp_path):
    original = workbook(tmp_path / "original.xlsx", "=A1+A2+1", extras={"A1": 0.1234567890123456})
    exported = workbook(tmp_path / "exported.xlsx", "=A1+A2+1", cache=4.12345678901235,
                        extras={"A1": 0.123456789012346})
    manifest = recalc._numeric_manifest(original)
    before = original.read_bytes(), exported.read_bytes()
    raw, changed = recalc._restore_numeric_export(exported.read_bytes(), manifest, numeric_proof(manifest))
    assert changed == 1
    restored = tmp_path / "restored.xlsx"
    restored.write_bytes(raw)
    assert recalc._recalc_valid(original, restored) is None
    assert before == (original.read_bytes(), exported.read_bytes())
    # No semantic correction of a wrong formula or alteration of formula cache.
    values = openpyxl.load_workbook(restored, data_only=True)
    assert values["S"]["B1"].value == 4.12345678901235
    assert values["S"]["A1"].value == 0.1234567890123456
    values.close()
    with zipfile.ZipFile(exported) as old, zipfile.ZipFile(restored) as new:
        for name in old.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert old.read(name) == new.read(name)
        old_xml = old.read("xl/worksheets/sheet1.xml")
        new_xml = new.read("xl/worksheets/sheet1.xml")
        assert new_xml == old_xml.replace(b"0.123456789012346", b"0.1234567890123456")


@pytest.mark.parametrize("changed", ["hash", "before", "after"])
def test_no_numeric_restoration_without_bound_before_and_after_proof(tmp_path, changed):
    path = workbook(tmp_path / "in.xlsx")
    manifest = recalc._numeric_manifest(path)
    proof = numeric_proof(manifest)
    if changed == "hash":
        proof["numeric_manifest_sha256"] = "0" * 64
    else:
        proof["numeric_verified_" + changed] -= 1
    with pytest.raises(ValueError, match="computation input proof"):
        recalc._restore_numeric_export(path.read_bytes(), manifest, proof)


@pytest.mark.parametrize("xml", [
    '<c r="A1" t="str"><v>2</v></c>', '<c r="A1"><f>1+1</f><v>2</v></c>',
    '<c r="A1"><v/></c>', '<c r="A1"><v>2</v><v>2</v></c>', '<c r="A2"><v>2</v></c>',
    '<c r="A1"><v>NaN</v></c>', '<c r="A1"><v>2</v></c><c r="A1"><v>2</v></c>',
])
def test_numeric_export_structural_change_cannot_be_repaired(xml):
    raw = ('<worksheet xmlns="' + compat.MAIN_NS + '"><sheetData><row>' + xml + '</row></sheetData></worksheet>').encode()
    with pytest.raises(ValueError):
        recalc._restore_numeric_xml(raw, [{"coordinate": "A1", "xml_value": "3"}])


def test_numeric_xml_preserves_prefixes_and_exponent_text():
    raw = ('<x:worksheet xmlns:x="' + compat.MAIN_NS + '"><x:sheetData><x:row>'
           '<x:c r="A1"><x:v>1.2e-20</x:v></x:c>'
           '<x:c r="B1"><x:f>A1</x:f><x:v>1.2e-20</x:v></x:c>'
           '</x:row></x:sheetData></x:worksheet>').encode()
    restored, changed = recalc._restore_numeric_xml(raw, [{"coordinate": "A1", "xml_value": "1.23E-20"}])
    assert changed == 1 and restored == raw.replace(b"1.2e-20", b"1.23E-20", 1)


@pytest.mark.parametrize("valid_proof", [True, False])
def test_recalculator_binds_proof_and_retains_raw_export(tmp_path, monkeypatch, valid_proof):
    from pathlib import Path

    from skillopt.skill_validation.sandbox import _Command
    source = workbook(tmp_path / "source.xlsx", extras={"A1": 0.1234567890123456})
    exported = workbook(tmp_path / "export.xlsx", cache=3.12345678901235,
                        extras={"A1": 0.123456789012346})
    image = "sha256:" + "a" * 64
    monkeypatch.setattr(recalc.platform, "system", lambda: "Linux")
    monkeypatch.setattr(recalc.shutil, "which", lambda _: "/usr/bin/docker")

    def command(argv, *args):
        if argv[1:3] == ["image", "inspect"]:
            return _Command(0, json.dumps({"Id": image, "Os": "linux", "Config": {}}).encode())
        if argv[1] != "run":
            return _Command(0)
        mount = next(arg for arg in argv if arg.startswith("type=bind,src=") and "dst=/input," in arg)
        directory = Path(mount.split("src=", 1)[1].split(",", 1)[0])
        manifest = json.loads((directory / "numeric_literals.json").read_bytes())
        assert manifest["input_sha256"] == recalc.sha(source)
        response = {"status": "available", "output_base64": base64.b64encode(exported.read_bytes()).decode(),
                    "output_sha256": recalc.sha(exported), "calculate_all": True, "macros": "never_execute",
                    "links": "no_update", "libreoffice_version": "fixture"}
        if valid_proof:
            response.update(numeric_proof(manifest))
        return _Command(0, json.dumps(response).encode())

    monkeypatch.setattr(recalc, "_bounded_command", command)
    receipt = recalc.Recalculator(image).run(source)
    assert receipt["cleanup_confirmed"]
    if not valid_proof:
        assert receipt["status"] == "unknown" and receipt["reason"] == "invalid_recalculation_receipt"
        return
    assert receipt["status"] == "available" and receipt["numeric_literal_texts_restored"] == 1
    assert receipt["raw_export_sha256"] == recalc.sha(exported)
    assert base64.b64decode(receipt["raw_export_base64"]) == exported.read_bytes()
    restored = tmp_path / "verified.xlsx"
    restored.write_bytes(base64.b64decode(receipt["output_base64"]))
    assert recalc._recalc_valid(source, restored) is None


@pytest.mark.parametrize("before,after", [
    ("=FALSE", "=FALSE()"), ("=TRUE", "=TRUE()"),
    ("=IF(FALSE,1,2)", "=IF(FALSE(),1,2)"),
    ("=IF(TRUE,1,2)", "=IF(TRUE(),1,2)"),
    ("=IF(A1=TRUE,FALSE,TRUE)", "=IF(A1=TRUE(),FALSE(),TRUE())"),
    ("=SUM(FALSE,1)+TRUE", "=SUM(FALSE(),1)+TRUE()"),
    ("=-(FALSE)+TRUE%", "=-(FALSE())+TRUE()%"),
    ('=IF(TRUE,"FALSE",FALSE)', '=IF(TRUE(),"FALSE",FALSE())'),
    ("=IF( TRUE, 1, FALSE )", "=IF( TRUE(), 1, FALSE() )"),
])
def test_only_boolean_empty_calls_in_operand_positions_are_equivalent(before, after):
    assert recalc._same_formula(before, after)
    assert recalc._same_formula(after, before)


@pytest.mark.parametrize("before,after", [
    ('="FALSE"', '="FALSE()"'), ('="TRUE"', '="TRUE()"'),
    ('=IF(TRUE,"FALSE",1)', '=IF(TRUE(),"FALSE()",1)'),
    ('="say ""FALSE"""', '="say ""FALSE()"""'),
    ("=FALSE!A1", "=FALSE()!A1"), ("='FALSE'!A1", "='FALSE()'!A1"),
    ("=TRUE:A1", "=TRUE():A1"), ("=A1:FALSE", "=A1:FALSE()"),
    ("=FALSE A1", "=FALSE() A1"), ("=A1 TRUE", "=A1 TRUE()"),
    ("=MYFALSE", "=MYFALSE()"), ("=FALSE", "=FALSE(0)"),
    ("=TRUE", "=TRUE(1)"), ("=FALSE", "=TRUE()"),
    ("=IF(FALSE,A1,2)", "=IF(FALSE(),A2,2)"),
    ("=IF(FALSE,A1,2)", "=IF(FALSE(), A1,2)"),
    ("=FALSE", "=FALSE( )"), ("=FALSE", "=false()"),
    ("={FALSE,TRUE}", "={FALSE(),TRUE()}"),
    ("=FALSE)", "=FALSE())"), ("=(TRUE", "=(TRUE()"),
    ("=[book]FALSE", "=[book]FALSE()"),
])
def test_boolean_rewrite_cannot_hide_strings_names_ranges_or_other_changes(before, after):
    assert not recalc._same_formula(before, after)
    assert not recalc._same_formula(after, before)


def test_boolean_formula_guard_and_wrong_result_are_not_corrected(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula
    source = workbook(tmp_path / "before.xlsx", "=IF(FALSE,5,6)")
    output = workbook(tmp_path / "after.xlsx", "=IF(FALSE(),5,6)", cache=6)
    gold = workbook(tmp_path / "gold.xlsx", "=IF(FALSE,6,5)", cache=5)
    assert recalc._recalc_valid(source, output) is None
    engine = Fake({recalc.sha(source): output, recalc.sha(gold): gold})
    assert recalc.evaluate_pair(source, gold, "B1", engine, tmp_path / "cache")["status"] == "fail"
    old = ArrayFormula(ref="B1", text="=IF(FALSE,5,6)")
    equivalent = ArrayFormula(ref="B1:B1", text="=IF(FALSE(),5,6)")
    assert recalc._same_formula(old, equivalent)
    assert not recalc._same_formula(old, ArrayFormula(ref="B1:B2", text="=IF(FALSE(),5,6)"))
    assert not recalc._same_formula(old, "=IF(FALSE(),5,6)")


@pytest.mark.parametrize("defined_name", [False, True])
def test_boolean_value_equivalence_not_assumed_when_formula_text_is_observable(tmp_path, defined_name):
    from openpyxl.workbook.defined_name import DefinedName
    before = workbook(tmp_path / "before.xlsx", "=IF(FALSE,1,2)")
    after = workbook(tmp_path / "after.xlsx", "=IF(FALSE(),1,2)", cache=2)
    for path in (before, after):
        book = openpyxl.load_workbook(path)
        if defined_name:
            book.defined_names.add(DefinedName("Observed", attr_text="GET.CELL(6,S!B1)"))
        else:
            book["S"]["D1"] = "=FORMULATEXT(B1)"
        book.save(path)
        book.close()
    assert recalc._recalc_valid(before, after) == "recalculation_formula_text_dependency_unverified"


def xml_package(raw):
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("customXml/fixture.xml", raw)
    return zipfile.ZipFile(io.BytesIO(data.getvalue()))


def utf16_xml(text, endian):
    return (b"\xff\xfe" if endian == "le" else b"\xfe\xff") + text.encode("utf-16-" + endian)


@pytest.mark.parametrize("endian", ["le", "be"])
@pytest.mark.parametrize("declared", [None, "UTF-16", "endian_specific"])
def test_utf16_metadata_is_strictly_read_without_rewriting_original(tmp_path, endian, declared):
    encoding = "UTF-16" + endian.upper() if declared == "endian_specific" else declared
    declaration = f'<?xml version="1.0" encoding="{encoding}"?>' if encoding else ""
    raw = utf16_xml(declaration + '<fixture><note>安全 &amp; valid</note></fixture>', endian)
    with xml_package(raw) as archive:
        assert recalc._xml(archive, "customXml/fixture.xml").find("note").text == "安全 & valid"
        # The frozen historical reader is not monkeypatched or relaxed.
        with pytest.raises(compat.CompatibilityError):
            compat._xml(archive, "customXml/fixture.xml")
    path = workbook(tmp_path / "in.xlsx")
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("customXml/fixture.xml", raw)
    original = path.read_bytes()
    assert recalc.screen(path) is None
    assert path.read_bytes() == original


@pytest.mark.parametrize("endian", ["le", "be"])
@pytest.mark.parametrize("body,reason", [
    ('<!DOCTYPE x><x/>', "unsupported_workbook_structure"),
    ('<!DOCTYPE x [<!ENTITY payload "v">]><x>&payload;</x>', "unsupported_workbook_structure"),
    ('<!ENTITY payload "v"><x/>', "unsupported_workbook_structure"),
    ('<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>',
     "unsupported_external_relationship"),
    ('<definedNames><definedName name="Remote">[private.xlsx]S!A1</definedName></definedNames>',
     "unsupported_external_or_volatile_defined_name"),
])
def test_utf16_cannot_hide_forbidden_semantics(tmp_path, endian, body, reason):
    path = workbook(tmp_path / "in.xlsx")
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("customXml/fixture.xml", utf16_xml(body, endian))
    assert recalc.screen(path) == reason


@pytest.mark.parametrize("raw", [
    b"\xff\xfe\x00\x00" + "<x/>".encode("utf-32-le"),
    b"\x00\x00\xfe\xff" + "<x/>".encode("utf-32-be"),
    "<x/>".encode("utf-16-le"), "<x/>".encode("utf-16-be"),
    utf16_xml("<x>\x00</x>", "le"),
    b"\xff\xfe" + b"\x00\xd8",  # Unpaired high surrogate.
    b"\xfe\xff" + b"\xdc\x00",  # Unpaired low surrogate.
    b"\xff\xfe<",  # Odd byte count.
    utf16_xml('<?xml version="1.0" encoding="UTF-8"?><x/>', "le"),
    utf16_xml('<?xml version="1.0" encoding="UTF-16BE"?><x/>', "le"),
    utf16_xml('<?xml version="1.0" encoding="UTF-16LE"?><x/>', "be"),
    utf16_xml('<?xml version="1.0" encoding="UTF-16-LE"?><x/>', "le"),
    utf16_xml('\ufeff<?xml version="1.0" encoding="UTF-8"?><x/>', "le"),
    utf16_xml('\ufeff<?xml version="1.0" encoding="UTF-8"?><x/>', "be"),
    utf16_xml('\ufeff<?xml version="1.0" encoding="UTF-16"?><x/>', "le"),
    utf16_xml('\ufeff<x/>', "be"),
    utf16_xml("<x>", "be"),
])
def test_invalid_or_unsupported_xml_encoding_is_unknown(tmp_path, raw):
    path = workbook(tmp_path / "in.xlsx")
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("customXml/fixture.xml", raw)
    assert recalc.screen(path) == "unsupported_workbook_structure"


def test_utf16_has_raw_and_decoded_xml_size_budgets(monkeypatch):
    raw = utf16_xml("<x>安全安全安全安全安全安全安全安全</x>", "le")
    decoded_size = len(raw.decode("utf-16").encode("utf-8"))
    assert len(raw) < decoded_size
    with xml_package(raw) as archive:
        monkeypatch.setattr(compat, "MAX_XML_BYTES", len(raw) - 1)
        with pytest.raises(compat.CompatibilityError, match="xml_part_size_limit"):
            recalc._xml(archive, "customXml/fixture.xml")
        monkeypatch.setattr(compat, "MAX_XML_BYTES", len(raw))
        with pytest.raises(compat.CompatibilityError, match="xml_decoded_size_limit"):
            recalc._xml(archive, "customXml/fixture.xml")
        monkeypatch.setattr(compat, "MAX_XML_BYTES", decoded_size)
        assert recalc._xml(archive, "customXml/fixture.xml").tag == "x"


def test_utf8_reader_retains_existing_behavior():
    raw = b'<?xml version="1.0" encoding="UTF-8"?><x>&lt;!DOCTYPE literal</x>'
    with xml_package(raw) as archive:
        assert ET.tostring(recalc._xml(archive, "customXml/fixture.xml")) == ET.tostring(
            compat._xml(archive, "customXml/fixture.xml"))
