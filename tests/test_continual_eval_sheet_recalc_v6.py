"""Offline format-30 controls, not evidence of real LibreOffice qualification."""
import base64
import hashlib
import io
import json
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pytest
from openpyxl.styles import numbers
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName

from skillopt.coevolution_v5.core import seal
from skillopt.continual_eval import sheet_recalc as v5
from skillopt.continual_eval import sheet_recalc_v6 as v6
from skillopt.continual_eval import spreadsheet_compat as compat


def workbook(path, value=46000, *, date=False, formula=None, epoch=None, second_date=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet["A1"] = value
    if date:
        sheet["A1"].number_format = "yyyy-mm-dd"
    else:
        sheet["A1"].number_format = numbers.FORMAT_NUMBER
        sheet["A1"]._style.numFmtId = 30
    if formula:
        sheet["B1"] = formula
    if second_date:
        sheet["C1"] = 46010
        sheet["C1"].number_format = "yyyy-mm-dd"
    if epoch:
        book.epoch = epoch
    book.save(path)
    book.close()
    return path


def write_view(tmp_path, before, after):
    raw, proof = v6.comparison_view(before, after)
    target = tmp_path / "view.xlsx"
    target.write_bytes(raw)
    return target, proof


def fake_receipt(before, after, identity):
    manifest = v5._numeric_manifest(before)
    encoded = json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
    return seal({"status": "available", "reason": "fixture", "identity": identity,
        "input_sha256": v5.sha(before), "output_sha256": v5.sha(after),
        "output_base64": base64.b64encode(after.read_bytes()).decode(),
        "numeric_manifest_sha256": hashlib.sha256(encoded).hexdigest(),
        "numeric_literals_verified": len(manifest["cells"]),
        "cleanup_confirmed": True, "model_api_calls": 0})


def test_format30_scalar_view_retains_originals_and_v5_rejection(tmp_path):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", date=True)
    frozen = before.read_bytes(), after.read_bytes()
    assert v5._recalc_valid(before, after) == "recalculation_changed_literal_cell"
    view, proof = write_view(tmp_path, before, after)
    assert v5._recalc_valid(before, view) is None
    assert len(proof["adapted_cells"]) == 1
    assert proof["export_sha256"] == v5.sha(after)
    assert proof["view_sha256"] == v5.sha(view)
    assert frozen == (before.read_bytes(), after.read_bytes())


def test_view_changes_only_existing_selected_style_attribute(tmp_path):
    before = workbook(tmp_path / "before.xlsx", second_date=True)
    after = workbook(tmp_path / "after.xlsx", date=True, second_date=True)
    view, _ = write_view(tmp_path, before, after)
    with zipfile.ZipFile(after) as original, zipfile.ZipFile(view) as changed:
        assert original.namelist() == changed.namelist()
        for name in original.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert original.read(name) == changed.read(name)
            else:
                assert original.read(name).replace(b'r="A1" s="1"', b'r="A1" s="0"') == changed.read(name)
    book = openpyxl.load_workbook(view)
    assert book["S"]["A1"].value == 46000
    assert book["S"]["C1"].is_date
    book.close()


@pytest.mark.parametrize("value", [46001, 46000.00001, -46000])
def test_changed_numeric_literal_is_not_coerced(tmp_path, value):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", value, date=True)
    with pytest.raises(ValueError, match="binary64 changed"):
        v6.comparison_view(before, after)


@pytest.mark.parametrize("formula", ['=CELL("format",A1)', '=GET.CELL(7,A1)',
    '=TEXT(A1,"yyyy")', '=FORMULATEXT(B2)', '=DOLLAR(A1)', '=FIXED(A1)', '=GET.WORKBOOK(1)'])
def test_format_observers_leave_v5_rejection(tmp_path, formula):
    before = workbook(tmp_path / "before.xlsx", formula=formula)
    after = workbook(tmp_path / "after.xlsx", date=True, formula=formula)
    view, proof = write_view(tmp_path, before, after)
    assert not proof["adapted_cells"]
    assert view.read_bytes() == after.read_bytes()
    assert v5._recalc_valid(before, view) == "recalculation_changed_literal_cell"


def test_defined_name_observer_and_epoch_change_are_not_authorized(tmp_path):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", date=True)
    for path in (before, after):
        book = openpyxl.load_workbook(path)
        book.defined_names.add(DefinedName("ObservedFormat", attr_text='CELL("format",S!A1)'))
        book.save(path)
        book.close()
    _, proof = v6.comparison_view(before, after)
    assert not proof["adapted_cells"]
    before = workbook(tmp_path / "epoch-original.xlsx")
    after = workbook(tmp_path / "epoch-export.xlsx", date=True, epoch=CALENDAR_MAC_1904)
    view, proof = write_view(tmp_path, before, after)
    assert not proof["adapted_cells"]
    assert v5._recalc_valid(before, view) == "recalculation_changed_literal_cell"


def test_other_numeric_formats_are_not_date_equivalence(tmp_path):
    before = workbook(tmp_path / "before.xlsx")
    book = openpyxl.load_workbook(before)
    book["S"]["A1"].number_format = "0"
    book.save(before)
    book.close()
    after = workbook(tmp_path / "after.xlsx", date=True)
    _, proof = v6.comparison_view(before, after)
    assert not proof["adapted_cells"]


def test_formula_changes_are_still_rejected_after_view(tmp_path):
    before = workbook(tmp_path / "before.xlsx", formula="=A1+1")
    after = workbook(tmp_path / "after.xlsx", date=True, formula="=A1+2")
    view, proof = write_view(tmp_path, before, after)
    assert len(proof["adapted_cells"]) == 1
    assert v5._recalc_valid(before, view) == "recalculation_changed_formula"


def test_literal_to_formula_is_not_patched(tmp_path):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", "=46000", date=True)
    view, proof = write_view(tmp_path, before, after)
    assert not proof["adapted_cells"]
    assert v5._recalc_valid(before, view) == "recalculation_changed_literal_cell"


def test_empty_string_and_error_formula_changes_stay_unknown(tmp_path):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", date=True)
    # An explicit OOXML string is distinct from an absent/blank cell.
    def add_cell(path, xml):
        raw = io.BytesIO(path.read_bytes())
        with zipfile.ZipFile(raw) as source, zipfile.ZipFile(path, "w") as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    content = content.replace(b"</row>", xml + b"</row>")
                target.writestr(item, content)
    add_cell(before, b'<c r="B1" t="inlineStr"><is><t></t></is></c>')
    view, _ = write_view(tmp_path, before, after)
    assert v5._recalc_valid(before, view) == "recalculation_changed_literal_cell"
    before = workbook(tmp_path / "error-before.xlsx")
    after = workbook(tmp_path / "error-after.xlsx", date=True)
    add_cell(before, b'<c r="B1" t="e"><v>#DIV/0!</v></c>')
    add_cell(after, b'<c r="B1" t="e"><f>1/0</f><v>#DIV/0!</v></c>')
    view, _ = write_view(tmp_path, before, after)
    assert v5._recalc_valid(before, view) == "recalculation_changed_literal_cell"


def test_recalculator_has_separate_identity_and_retains_export(tmp_path, monkeypatch):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", date=True)
    engine = v6.Recalculator("sha256:" + "a" * 64)
    monkeypatch.setattr(v5.Recalculator, "run", lambda self, path: fake_receipt(path, after, self.identity))
    result = engine.run(before)
    assert result["status"] == "available"
    assert result["engine_export_sha256"] == v5.sha(after)
    assert result["output_sha256"] != result["engine_export_sha256"]
    assert result["identity"]["version"] == v6.VERSION != v5.VERSION
    assert result["record_hash"] != fake_receipt(before, after, engine.identity)["record_hash"]
    assert result["model_api_calls"] == 0


@pytest.mark.parametrize("mutation", ["missing", "wrong_hash", "wrong_count", "bool_count"])
def test_numeric_proof_missing_or_mismatched_is_unknown(tmp_path, monkeypatch, mutation):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", date=True)
    engine = v6.Recalculator("sha256:" + "a" * 64)
    receipt = fake_receipt(before, after, engine.identity)
    if mutation == "missing":
        del receipt["numeric_manifest_sha256"]
    elif mutation == "wrong_hash":
        receipt["numeric_manifest_sha256"] = "0" * 64
    elif mutation == "bool_count":
        receipt["numeric_literals_verified"] = True
    else:
        receipt["numeric_literals_verified"] += 1
    monkeypatch.setattr(v5.Recalculator, "run", lambda self, path: receipt)
    result = engine.run(before)
    assert result["status"] == "unknown"
    assert result["reason"] == "format30_comparison_view_unverified"


def test_no_docker_remains_unsupported_and_never_host_fallback(tmp_path, monkeypatch):
    monkeypatch.setattr(v5.platform, "system", lambda: "Darwin")
    result = v6.Recalculator("sha256:" + "a" * 64).run(workbook(tmp_path / "in.xlsx"))
    assert result["status"] == "unknown" and result["reason"] == "linux_docker_unavailable"


def test_view_does_not_hide_wrong_answer_or_reference_drift(tmp_path, monkeypatch):
    gold = workbook(tmp_path / "gold.xlsx")
    wrong = workbook(tmp_path / "wrong.xlsx", 46001)
    gold_export = workbook(tmp_path / "gold-export.xlsx", date=True)
    wrong_export = workbook(tmp_path / "wrong-export.xlsx", 46001, date=True)
    mapping = {v5.sha(gold): gold_export, v5.sha(wrong): wrong_export}
    monkeypatch.setattr(v5.Recalculator, "run",
                        lambda self, path: fake_receipt(path, mapping[v5.sha(path)], self.identity))
    engine = v6.Recalculator("sha256:" + "a" * 64)
    result = v5.evaluate_pair(wrong, gold, "A1", engine, tmp_path / "receipts")
    assert result["status"] == "fail"
    assert len(result["receipts"]) == 2


def test_reserved_custom_style_and_duplicate_cells_fail_closed(tmp_path):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", date=True)
    original_bytes = after.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original_bytes)) as source, zipfile.ZipFile(after, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/styles.xml":
                tree = ET.fromstring(content)
                custom = tree.find(f"{{{compat.MAIN_NS}}}numFmts")
                ET.SubElement(custom, f"{{{compat.MAIN_NS}}}numFmt", numFmtId="30", formatCode="0")
                custom.set("count", str(len(custom)))
                content = ET.tostring(tree)
            target.writestr(item, content)
    with pytest.raises(ValueError, match="Reserved"):
        v6.comparison_view(before, after)
    with zipfile.ZipFile(io.BytesIO(original_bytes)) as source, zipfile.ZipFile(after, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"</row>", b'<c r="A1" s="1" t="n"><v>46000</v></c></row>')
            target.writestr(item, content)
    with pytest.raises(ValueError, match="Duplicate"):
        v6.comparison_view(before, after)


@pytest.mark.parametrize("target", ["original", "exported"])
@pytest.mark.parametrize("damage", ["doctype", "entity", "zip_expansion", "oversized_archive", "unsafe_member"])
def test_archive_and_xml_screening_precedes_openpyxl(tmp_path, monkeypatch, target, damage):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", date=True)
    path = before if target == "original" else after
    if damage == "oversized_archive":
        path.write_bytes(b"x" * (compat.MAX_ZIP_BYTES + 1))
    else:
        with zipfile.ZipFile(path, "a", compression=zipfile.ZIP_DEFLATED) as archive:
            if damage == "doctype":
                archive.writestr("extra.xml", b'<!DOCTYPE x [<!ENTITY e "x">]><x/>')
            elif damage == "entity":
                archive.writestr("extra.xml", b'<!ENTITY e "x"><x/>')
            elif damage == "zip_expansion":
                monkeypatch.setattr(compat, "MAX_EXPANDED_BYTES", 10000)
                archive.writestr("extra.bin", b"x" * 10001)
            else:
                archive.writestr("../outside.xml", b"<x/>")
    calls = []
    def forbidden(*args, **kwargs):
        calls.append(1)
        raise AssertionError("openpyxl must not parse a rejected archive")
    monkeypatch.setattr(v6.openpyxl, "load_workbook", forbidden)
    with pytest.raises(compat.CompatibilityError):
        v6.comparison_view(before, after)
    assert not calls


def calc_custom_general(path, *, format_code="General", duplicate=False):
    """Reproduce the styles from the rejected real Calc 25.2 q20 export.

    That authored control exported 164=General and 165=m/d/yyyy, with no
    builtin-0 cellXf. No benchmark workbook or answer data is embedded here.
    """
    raw = io.BytesIO(path.read_bytes())
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/styles.xml":
                tree = ET.fromstring(content)
                custom = tree.find(f"{{{compat.MAIN_NS}}}numFmts")
                for fmt in custom:
                    assert fmt.get("numFmtId") == "164"
                    fmt.set("numFmtId", "165")
                    fmt.set("formatCode", "m/d/yyyy")
                ET.SubElement(custom, f"{{{compat.MAIN_NS}}}numFmt", numFmtId="164", formatCode=format_code)
                if duplicate:
                    ET.SubElement(custom, f"{{{compat.MAIN_NS}}}numFmt", numFmtId="164", formatCode="mm/dd/yy")
                custom.set("count", str(len(custom)))
                for style in tree.find(f"{{{compat.MAIN_NS}}}cellXfs"):
                    style.set("numFmtId", "164" if style.get("numFmtId") == "0" else "165")
                content = ET.tostring(tree)
            target.writestr(item, content)


def test_authentic_calc_custom_general_representation_supported(tmp_path):
    before = workbook(tmp_path / "before.xlsx")
    after = workbook(tmp_path / "after.xlsx", date=True)
    calc_custom_general(after)
    view, proof = write_view(tmp_path, before, after)
    assert len(proof["adapted_cells"]) == 1
    assert v5._recalc_valid(before, view) is None
    with zipfile.ZipFile(after) as a, zipfile.ZipFile(view) as b:
        assert a.read("xl/styles.xml") == b.read("xl/styles.xml")


@pytest.mark.parametrize("code", ["general", "GENERAL", "General;General", "mm/dd/yy"])
def test_similar_or_date_custom_format_is_not_general(tmp_path, code):
    after = workbook(tmp_path / "after.xlsx", date=True)
    calc_custom_general(after, format_code=code)
    with compat._safe_archive(after) as archive:
        with pytest.raises(ValueError, match="No existing General"):
            v6._general_style(archive)


def test_duplicate_custom_format_is_not_disambiguated(tmp_path):
    after = workbook(tmp_path / "after.xlsx", date=True)
    calc_custom_general(after, duplicate=True)
    with compat._safe_archive(after) as archive:
        with pytest.raises(ValueError, match="Duplicate custom"):
            v6._general_style(archive)
