"""Opt-in profile/roster controls. All receipts here are engineering fixtures."""
import datetime
import json

import openpyxl
import pytest

from scripts import replay_sheet_v6 as replay
from skillopt.coevolution_v5.core import seal
from skillopt.continual_eval.core import read_json, write_json
from tests.test_continual_eval_sheet_recalc import workbook
from tests.test_continual_eval_sheet_recalc_v7 import _cache
from tests.test_sheet_v6_replay import IMAGE, Fake, qualification_fixture


def _setup(tmp_path, monkeypatch, profile):
    original = tmp_path / "original.xlsx"
    if not original.exists():
        workbook(original, cache=5)
    lock = tmp_path / "native.lock"
    lock.touch(exist_ok=True)
    old_controls = [{"name": name, "path": str(original), "expected": 5}
                    for name in sorted(replay.CONTROLS)]
    monkeypatch.setattr(replay, "_legacy_controls",
                        lambda *args: (old_controls.copy(), replay._files([original])))
    old_path = None
    if profile == "v7":
        old_root = tmp_path / "v6"
        if not old_root.exists():
            replay._qual_setup(old_root, IMAGE, tmp_path / "old" / "qualification.json", lock)
        protocol = read_json(old_root / "protocol.json", sealed=True)
        fake = Fake()
        fake.identity = protocol["engine"]
        engine = replay.DurableEngine(fake, old_root, lock)
        rows = []
        for spec in protocol["controls"]:
            receipt = engine.run(spec["path"])
            row = seal({"name": spec["name"], "protocol_hash": protocol["record_hash"],
                        "qualified": False, "receipt_hash": receipt["record_hash"], "cleanup_confirmed": True})
            write_json(old_root / "controls" / (spec["name"] + ".json"), row)
            rows.append(row)
        old_path = old_root / "qualification.json"
        write_json(old_path, seal({"version": replay.VERSION, "status": "rejected", "engine": protocol["engine"],
            "engine_profile": "v6", "protocol_hash": protocol["record_hash"], "controls": rows}))
    return replay._qual_setup(tmp_path / profile, IMAGE, tmp_path / "old" / "qualification.json", lock, profile, old_path)


def test_v7_explicit_profile_preserves_original_twenty_control_inputs(tmp_path, monkeypatch):
    old = _setup(tmp_path, monkeypatch, "v6")
    new = _setup(tmp_path, monkeypatch, "v7")
    assert len(old["controls"]) == 20 and len(new["controls"]) == 23
    assert old["engine_profile"] == "v6" and new["engine_profile"] == "v7"
    assert new["engine"]["version"] == replay.v7.VERSION
    assert {c["name"] for c in new["controls"]} == replay._controls("v7")
    assert not new["old_authorization_inherited"]
    names = {p.rsplit("/", 1)[-1] for p in new["engine"]["sources"]}
    assert {"sheet_recalc.py", "sheet_recalc_v6.py", "sheet_recalc_v7.py"} <= names
    for source in old["controls"]:
        candidate = next(c for c in new["controls"] if c["name"] == source["name"])
        assert source["expected"] == candidate["expected"]
        assert source["path"] == candidate["path"]
        assert replay.v5.sha(source["path"]) == new["inventory"][candidate["path"]]
        left = openpyxl.load_workbook(source["path"])
        right = openpyxl.load_workbook(candidate["path"])
        assert [(c.coordinate, c.value, c.data_type, c._style) for c in left.active._cells.values()] == [
            (c.coordinate, c.value, c.data_type, c._style) for c in right.active._cells.values()]
        left.close()
        right.close()
    for name in ("original_date_literal", "original_date_formula"):
        control = next(c for c in new["controls"] if c["name"] == name)
        assert control["expected"] == {"type": "datetime", "iso8601": "2025-12-09T00:00:00"}
    replay._engine(new, tmp_path / "v7")  # Does not execute a container.


@pytest.mark.parametrize("damage", ["profile_only", "missing_controls"])
def test_old_authorization_cannot_become_v7(tmp_path, damage):
    q, _ = qualification_fixture(tmp_path)
    protocol_path = q.parent / "protocol.json"
    protocol = read_json(protocol_path, sealed=True)
    protocol.pop("record_hash")
    protocol["engine_profile"] = "v7"
    if damage == "missing_controls":
        protocol["engine"] = replay.v7.Recalculator(IMAGE).identity
    protocol_path.write_text(json.dumps(seal(protocol)))  # Authored tampering control.
    with pytest.raises(ValueError, match="control roster"):
        replay.qualify(q.parent)


def test_frozen_v7_resume_cannot_switch_profile(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch, "v7")
    with pytest.raises(ValueError, match="Frozen engine profile"):
        replay.qualify(tmp_path / "v7", engine_version="v6")


def test_unrecognized_profile_rejected_before_creation(tmp_path):
    with pytest.raises(ValueError, match="Unknown engine profile"):
        replay._qual_setup(tmp_path / "bad", IMAGE, "unused", "unused", "v9")
    assert not (tmp_path / "bad").exists()


@pytest.mark.parametrize("damage", ["input", "expected", "role", "receipt", "pending"])
def test_prior_comparison_fixtures_are_bound_not_new_authorization(tmp_path, monkeypatch, damage):
    protocol = _setup(tmp_path, monkeypatch, "v7")
    path = tmp_path / "v6" / "qualification.json"
    assert protocol["comparison_fixture_qualification"] == str(path)
    q = read_json(path, sealed=True)
    prior_path = path.parent / "protocol.json"
    prior = read_json(prior_path, sealed=True)
    legacy = [c for c in prior["controls"] if c["name"] in replay.CONTROLS]
    spec = next(c for c in prior["controls"] if c["name"] == "format30_correct")
    if damage == "input":
        from pathlib import Path
        Path(spec["path"]).write_bytes(b"authored corruption")
    elif damage in {"expected", "role"}:
        spec["expected" if damage == "expected" else "name"] = 0 if damage == "expected" else "unrelated"
        prior.pop("record_hash")
        prior = seal(prior)
        prior_path.write_text(json.dumps(prior))
        q["protocol_hash"] = prior["record_hash"]
    elif damage == "receipt":
        key = replay.digest({"input": replay.v5.sha(spec["path"]), "engine": q["engine"]})
        receipt_path = path.parent / "recalculations" / (key + ".json")
        receipt = read_json(receipt_path, sealed=True)
        receipt.pop("record_hash")
        receipt["input_sha256"] = "a" * 64
        receipt_path.write_text(json.dumps(seal(receipt)))
    else:
        q["status"] = "pending"
    q.pop("record_hash")
    path.write_text(json.dumps(seal(q)))
    with pytest.raises(ValueError, match="fixture|q20"):
        replay._comparison_fixtures(path, IMAGE, legacy)


@pytest.mark.parametrize("damage", ["bytes", "expected"])
def test_individually_valid_legacy_controls_cannot_be_mixed_with_another_q20(tmp_path, monkeypatch, damage):
    _setup(tmp_path, monkeypatch, "v7")
    old = read_json(tmp_path / "v6" / "protocol.json", sealed=True)
    legacy = [dict(c) for c in old["controls"] if c["name"] in replay.CONTROLS]
    changed = next(c for c in legacy if c["name"] == "correct")
    if damage == "bytes":
        # Both are structurally valid fixture workbooks, not a damaged ZIP.
        alternative = workbook(tmp_path / "another-valid-fixture.xlsx", cache=6)
        changed["path"] = str(alternative)
    else:
        changed["expected"] = 6
    monkeypatch.setattr(replay, "_legacy_controls", lambda *args: (legacy, replay._files([c["path"] for c in legacy])))
    with pytest.raises(ValueError, match="Legacy q17 and frozen q20"):
        replay._qual_setup(tmp_path / "mixed", IMAGE, tmp_path / "another-q17.json", tmp_path / "native.lock",
                           "v7", tmp_path / "v6" / "qualification.json")
    assert not (tmp_path / "mixed").exists()


@pytest.mark.parametrize("name", ["original_date_literal", "original_date_formula"])
def test_typed_date_controls_require_date_not_same_numeric_serial(tmp_path, name):
    source = tmp_path / "source.xlsx"
    exported = tmp_path / "exported.xlsx"
    for path in (source, exported):
        book = openpyxl.Workbook()
        book.active.title = "S"
        book.active["B1"] = (46000 if name == "original_date_literal" else "=DATE(2025,12,9)")
        book.active["B1"].number_format = "yyyy-mm-dd"
        book.save(path)
        book.close()
    if name == "original_date_formula":
        _cache(exported, "46000")
    control = {"name": name, "path": str(source),
               "expected": {"type": "datetime", "iso8601": "2025-12-09T00:00:00"}}
    receipt = Fake().run(exported)
    result = replay._control_result(control, receipt, tmp_path)
    assert result["qualified"]
    book = openpyxl.load_workbook(exported, data_only=True)
    assert type(book["S"]["B1"].value) is datetime.datetime
    book.close()
    # A bare serial is not allowed to masquerade as the required original date.
    book = openpyxl.load_workbook(exported)
    book["S"]["B1"] = 46000
    book["S"]["B1"].number_format = "General"
    book.save(exported)
    book.close()
    assert not replay._control_result(control, Fake().run(exported), tmp_path)["qualified"]


@pytest.mark.parametrize("actual,qualified", [(46002, True), (46001, False)])
def test_added_wrong_control_keeps_original_wrong_value(tmp_path, actual, qualified):
    source = workbook(tmp_path / "source.xlsx", "=A1+2")
    exported = workbook(tmp_path / "exported.xlsx", "=A1+2", cache=actual)
    control = {"name": "general_wrong_numeric", "path": str(source), "expected": 46002}
    assert replay._control_result(control, Fake().run(exported), tmp_path)["qualified"] is qualified
