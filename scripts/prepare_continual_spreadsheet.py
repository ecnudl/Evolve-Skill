"""Prepare private Verified400 panels without repairing ambiguous upstream data.

The original archive/split/staging script are immutable inputs. Missing named
gold partners remain missing: tasks retain their denominator and are blocked by
dataset readiness. No model is called and final workbooks are not scored.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
from collections import Counter
from pathlib import Path, PurePosixPath

from scripts import prepare_verified400 as source
from skillopt.continual_eval import backends, datasets
from skillopt.continual_eval.core import source_identity

VERSION = "continual-verified400-private-preparation-v1"
PARTITIONS = {"train": "development", "val": "skill_confirmation", "test": "final"}


def _require(ok, message):
    if not ok:
        raise ValueError(message)


def _publish(path, value):
    source.immutable(path, source.encoded(value))


def _inventory(files, indices):
    tables = source._metadata_candidates(files)
    _require(len(tables) == 1, "Exactly one official metadata table required")
    metadata_path, rows = tables[0]
    by_id = {}
    for row in rows:
        identifier = str(row.get("id", ""))
        _require(identifier not in by_id, "Duplicate official task ID")
        by_id[identifier] = row
    _require(set(by_id) == set(indices), "Task identities differ from frozen source split")
    file_map = dict(files)
    directories = {str(PurePosixPath(name).parent) for name in file_map if name.endswith(".xlsx")}
    result = []
    for identifier, registration in indices.items():
        row = by_id[identifier]
        for field in ("instruction", "instruction_type", "spreadsheet_path", "answer_position"):
            _require(isinstance(row.get(field), str) and row[field].strip(), "Incomplete official task metadata")
        _require(row["instruction_type"] == registration["index"]["instruction_type"], "Instruction type differs from source split")
        expected = source.member_path(row["spreadsheet_path"]).rstrip("/")
        candidates = [name for name in directories if name == expected or name.endswith("/" + expected)]
        _require(len(candidates) == 1, "Explicit source directory is missing or ambiguous")
        directory = candidates[0]
        workbooks = sorted(name for name in file_map if str(PurePosixPath(name).parent) == directory and name.endswith(".xlsx"))
        pairs, used, unresolved = [], set(), []
        for name in workbooks:
            base = PurePosixPath(name).name
            if base.endswith("_input.xlsx"):
                partner = name[:-len("_input.xlsx")] + "_answer.xlsx"
            elif base.endswith("_init.xlsx"):
                partner = name[:-len("_init.xlsx")] + "_golden.xlsx"
            elif base == "initial.xlsx":
                partner = str(PurePosixPath(name).with_name("golden.xlsx"))
            else:
                continue
            used.add(name)
            if partner in file_map:
                used.add(partner)
            else:
                unresolved.append({"kind": "missing_named_gold_partner", "input": name, "expected_gold": partner})
            pairs.append((name, partner))
        _require(pairs, "No recognizable input workbook; refuse to invent task assets")
        extras = sorted(set(workbooks) - used)
        _require(not extras or unresolved, "Unexpected additional workbook without a missing partner")
        # Extras are recorded, never reinterpreted as the correct partner.
        result.append({"task_id": identifier, "split": registration["split"], "metadata": row,
                       "pairs": pairs, "unresolved": unresolved, "unpaired_files": extras})
    return metadata_path, result


def prepare(repo, source_root, output):
    repo, source_root, output = Path(repo).absolute(), Path(source_root).absolute(), Path(output).absolute()
    for path in (repo, source_root, output):
        source.no_symlinks(path)
    _require(output != source_root and not output.is_relative_to(source_root), "Use a new preparation output, not the source archive directory")
    _require((source_root / "source_snapshot.json").is_file(), "Pinned source snapshot must already be downloaded")
    snapshot = source.download(source_root, fetch=lambda *_: (_ for _ in ()).throw(ValueError("No network in preparation")))
    indices, index_hashes = source.source_indices(repo)
    files = source.regular_members((source_root / source.FILENAME).read_bytes())
    metadata_path, rows = _inventory(files, indices)
    # Every archive member and every task identity is validated before writes.
    for name, body in files:
        source.immutable(output / "assets" / name, body)
    file_hashes = {name: hashlib.sha256(body).hexdigest() for name, body in files}
    panels, summaries, panel_data = {}, {}, {}
    for split, partition in PARTITIONS.items():
        tasks = []
        for row in rows:
            if row["split"] != split:
                continue
            metadata = row["metadata"]
            public_inputs = [str(output / "assets" / pair[0]) for pair in row["pairs"]]
            private_golds = [str(output / "assets" / pair[1]) for pair in row["pairs"]]
            public = {"instruction": metadata["instruction"], "input_files": public_inputs,
                      "answer_position": metadata["answer_position"]}
            task = {"task_id": row["task_id"],
                    "family_id": datasets._family(row["task_id"], {"instruction": metadata["instruction"]}, None),
                    "project_id": "", "partition": partition, "public": public,
                    "private": {"test_files": private_golds, "answer_position": metadata["answer_position"],
                                "source_sha256": file_hashes[metadata_path],
                                "asset_sha256": {str(output / "assets" / name): file_hashes[name]
                                                 for pair in row["pairs"] for name in pair if name in file_hashes}}}
            tasks.append(task)
        panel = {"version": datasets.VERSION, "benchmark": "spreadsheetbench", "dataset_revision": source.REVISION,
                 "provenance": "natural", "tasks": tasks}
        datasets.validate_panel(panel)
        panel_path = output / "panels" / f"{partition}.json"
        panels[split] = str(panel_path)
        panel_data[split] = panel
        summaries[split] = datasets.readiness(panel)
    combined = {**panel, "tasks": [task for split in PARTITIONS for task in panel_data[split]["tasks"]]}
    datasets.validate_panel(combined)  # exact families cannot straddle source splits
    for split in PARTITIONS:
        _publish(panels[split], panel_data[split])
    unresolved = [{"task_id": row["task_id"], "split": row["split"], "issues": row["unresolved"],
                   "unpaired_files": row["unpaired_files"]} for row in rows if row["unresolved"]]
    private_manifest = source.seal({"version": VERSION, "source_snapshot_hash": snapshot["record_hash"],
        "archive_sha256": snapshot["files"][source.FILENAME]["sha256"], "source_revision": source.REVISION,
        "source_index_hashes": index_hashes, "metadata_source": metadata_path,
        "panels": panels, "partition_mapping": PARTITIONS, "unresolved_source_assets": unresolved,
        "asset_inventory": [{"path": name, "sha256": file_hashes[name], "bytes": len(body)} for name, body in files],
        "model_api_calls": 0, "final_scoring_performed": False,
        "exposure_note": "Source split retained. No assertion of globally unexposed tasks or semantic family independence."})
    _publish(output / "private_preparation_manifest.json", private_manifest)
    summary = source.seal({"version": VERSION, "source_revision": source.REVISION,
        "private_manifest_hash": private_manifest["record_hash"], "panels": summaries,
        "source_tasks": len(rows), "workbook_cases": sum(len(row["pairs"]) for row in rows),
        "unresolved_tasks": len(unresolved), "unresolved_by_split": dict(Counter(row["split"] for row in unresolved)),
        "model_api_calls": 0, "final_scoring_performed": False,
        "readiness_not_method_effect": True, "split_policy": "original_80_40_280_unchanged"})
    _publish(output / "preparation_summary.json", summary)
    return summary


def qualify_development(output):
    """Check public previews, target syntax/sheets and cache availability only.

    Reads development workbooks only. This is not solver success, a baseline
    score, a reason to filter tasks, or a claim of fresh formula caches.
    """
    from contextlib import ExitStack

    import openpyxl

    output = Path(output).absolute()
    panel = datasets.load_panel(output / "panels/development.json")
    scorer = backends._sheet_scorer()
    rows = []
    for task in panel["tasks"]:
        record = {"task_id": task["task_id"], "preview_available": True, "target_parseable": True,
                  "target_sheets_present": True, "gold_readable": True, "target_cells": 0,
                  "formula_cells": 0, "missing_formula_caches": 0}
        for path in task["public"]["input_files"]:
            try:
                backends._sheet_preview(path)
            except Exception:
                record["preview_available"] = False
        for path in task["private"]["test_files"]:
            try:
                with ExitStack() as stack:
                    wb_formula = openpyxl.load_workbook(path, read_only=True, data_only=False)
                    stack.callback(wb_formula.close)
                    wb_values = openpyxl.load_workbook(path, read_only=True, data_only=True)
                    stack.callback(wb_values.close)
                    try:
                        targets = scorer._answer_targets(task["public"]["answer_position"], wb_formula.sheetnames[0])
                    except Exception:
                        record["target_parseable"] = False
                        continue
                    for sheet, region in targets:
                        if sheet not in wb_formula.sheetnames:
                            record["target_sheets_present"] = False
                            continue
                        (min_col, min_row), (max_col, max_row) = scorer._range_bounds(region)
                        bounds = {"min_col": min_col, "min_row": min_row, "max_col": max_col, "max_row": max_row}
                        for formulas, values in zip(wb_formula[sheet].iter_rows(**bounds), wb_values[sheet].iter_rows(**bounds)):
                            for formula, value in zip(formulas, values):
                                record["target_cells"] += 1
                                if formula.data_type == "f":
                                    record["formula_cells"] += 1
                                    if value.value is None:
                                        record["missing_formula_caches"] += 1
            except Exception:
                record["gold_readable"] = False
        record["status"] = "ready" if all(record[field] for field in
            ("preview_available", "target_parseable", "target_sheets_present", "gold_readable")) and not record["missing_formula_caches"] else "unknown_risk"
        rows.append(record)
    report = {"version": VERSION + "-development-qualification", "panel_hash": datasets.panel_hash(panel),
              "tasks": len(rows), "ready": sum(row["status"] == "ready" for row in rows),
              "unknown_risk": sum(row["status"] != "ready" for row in rows),
              "preview_unavailable": sum(not row["preview_available"] for row in rows),
              "unparseable_target": sum(not row["target_parseable"] for row in rows),
              "missing_target_sheet": sum(not row["target_sheets_present"] for row in rows),
              "unreadable_gold": sum(not row["gold_readable"] for row in rows),
              "tasks_with_missing_formula_cache": sum(row["missing_formula_caches"] > 0 for row in rows),
              "model_api_calls": 0, "final_workbooks_opened": 0, "tasks_filtered": 0,
              "formula_recalculation_performed": False, "details": rows}
    _publish(output / "development_qualification.json", report)
    return {key: value for key, value in report.items() if key != "details"}


def native_smoke(output, runtime):
    """Three authored controls only; never score any official final workbook."""
    import openpyxl

    output = Path(output).absolute()
    source.no_symlinks(output)
    sources = source_identity()
    preparation_hash = datasets.file_hash(__file__)
    expected = {"correct": "pass", "wrong": "fail", "dependency_unavailable": "unknown"}
    prior_path = output / "native_qualification.json"
    if prior_path.exists():
        prior = json.loads(prior_path.read_bytes())
        _require(prior.get("version") == VERSION + "-native-controls", "Existing native control version differs")
        _require(prior.get("runtime") == runtime, "Existing native control runtime differs; use a new output")
        _require(prior.get("source_identity") == sources and prior.get("preparation_script_sha256") == preparation_hash,
                 "Existing native control source identity differs; use a new output")
        _require({key: value.get("status") for key, value in prior.get("controls", {}).items()} == expected
                 and all(value.get("cleanup_confirmed") is True for value in prior["controls"].values()),
                 "Existing native control statuses or cleanup differ")
        _require(type(prior.get("fixture_hashes")) is dict
                 and set(prior["fixture_hashes"]) == {"input.xlsx", "gold.xlsx"},
                 "Existing native fixture identities differ")
        for name, digest in prior["fixture_hashes"].items():
            _require(datasets.file_hash(output / "native_controls" / name) == digest, "Existing native fixture changed")
        return prior
    ready = backends.readiness("spreadsheetbench", runtime)
    _require(ready["status"] == "ready", "Native spreadsheet backend not ready")
    directory = output / "native_controls"
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    raw = {}
    for name, value in (("input", 1), ("gold", 2)):
        book = openpyxl.Workbook()
        book.active["A1"] = value
        import io
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        raw[name] = stream.getvalue()
        # XLSX ZIP timestamps vary; existing authored controls may be reused,
        # but are never silently overwritten on rerun.
        path = directory / f"{name}.xlsx"
        if not path.exists():
            source.immutable(path, raw[name])
        raw[name] = path.read_bytes()
    controls = {
        "correct": "import openpyxl\nw=openpyxl.load_workbook('input.xlsx')\nw.active['A1']=2\nw.save('output.xlsx')",
        "wrong": "import openpyxl\nw=openpyxl.load_workbook('input.xlsx')\nw.active['A1']=3\nw.save('output.xlsx')",
        "dependency_unavailable": "import continual_eval_missing_dependency_fixture",
    }
    results = {}
    for label, code in controls.items():
        generation = backends._native({"operation": "spreadsheet_generate", "code": code,
                                      "input_base64": base64.b64encode(raw["input"]).decode()}, runtime)
        prediction = {"status": "available", "output": {"cases": [generation]}}
        result = backends.score("spreadsheetbench", {"answer_position": "A1"},
                                {"test_files": [str(directory / "gold.xlsx")], "answer_position": "A1"}, prediction,
                                runtime=runtime)
        results[label] = {"status": result["status"], "execution": generation.get("execution_costs", {}),
                          "cleanup_confirmed": generation.get("cleanup_confirmed", False)}
    _require({key: value["status"] for key, value in results.items()} == expected, "Native controls did not match expected statuses")
    _require(all(value["cleanup_confirmed"] for value in results.values()), "Native control cleanup not confirmed")
    report = {"version": VERSION + "-native-controls", "provenance": "fixture", "model_api_calls": 0,
              "official_tasks_scored": 0, "runtime": runtime, "readiness": ready, "controls": results,
              "source_identity": sources, "preparation_script_sha256": preparation_hash,
              "fixture_hashes": {f"{name}.xlsx": hashlib.sha256(body).hexdigest() for name, body in raw.items()}}
    _publish(output / "native_qualification.json", report)
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, default=source.REPO)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--native-smoke", action="store_true")
    parser.add_argument("--image", default="")
    args = parser.parse_args(argv)
    summary = prepare(args.repo, args.source_root, args.output)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(json.dumps(qualify_development(args.output), ensure_ascii=False, indent=2))
    if args.native_smoke:
        report = native_smoke(args.output, {"image": args.image, "timeout_seconds": 300,
                                           "memory_mb": 4096, "cpus": 1})
        print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
