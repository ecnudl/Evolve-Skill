"""New passive-link qualification plus frozen C/160 candidate-only replay.

No model calls, answer changes, inherited new-scope authorization, or old writes.
The old 41 engine fixtures are reused byte-for-byte and re-executed; all 18 H
and two link-scoring controls are rerun. Two href-display controls are added.
"""
from __future__ import annotations

import argparse
import base64
import json
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl

from scripts import replay_sheet_frozen_gold as frozen
from scripts import replay_sheet_v6 as replay
from skillopt.coevolution_v5.core import seal
from skillopt.continual_eval import sheet_frozen_gold as scoring
from skillopt.continual_eval import sheet_passive_links as links
from skillopt.continual_eval import sheet_recalc as recalc
from skillopt.continual_eval.core import output_lock, read_json, require, safe_path, write_json
from skillopt.continual_eval.truncation_recovery import _readonly_parent_lock

VERSION = "passive-hyperlinks-frozen-h-replay-v3"
NEW_EXPECTED = {"passive_correct": 3, "passive_wrong": 4, "passive_file": None,
                "passive_http": None, "passive_credentials": None,
                "passive_volatile": None, "passive_orphan": None}
DISPLAY_EXPECTED = {"display_href_correct": 3, "display_href_wrong": 4}


def _sources():
    return replay._files([__file__, scoring.__file__, frozen.__file__, replay.__file__, links.__file__])


def _engine(protocol, root):
    require(protocol["version"] == VERSION and protocol["root"] == str(root)
            and protocol["sources"] == _sources(), "New replay source/root identity changed")
    replay._verify_files(protocol["inventory"])
    engine = links.Recalculator(protocol["engine"]["image_id"], protocol["engine"]["timeout_seconds"])
    require(engine.identity == protocol["engine"], "Passive-link engine identity changed")
    return replay.DurableEngine(engine, root, protocol["native_lock"])


def _old_controls(path):
    """Verify old fixture evidence, without inheriting new engine authorization."""
    path = safe_path(path)
    q, p = read_json(path, sealed=True), read_json(path.parent / "protocol.json", sealed=True)
    require(p["version"] == replay.VERSION and p["kind"] == "qualification"
            and p["root"] == str(path.parent) and replay._profile(p) == "v8"
            and q["status"] == "qualified" and q["protocol_hash"] == p["record_hash"]
            and q["engine"] == p["engine"] and len(q["controls"]) == 34,
            "Original complete v8b fixture qualification required")
    replay._verify_files(p["inventory"])
    replay._verify_files(p["engine"]["sources"])
    replay._published(path.parent, p)
    replay._validate_function_controls(p)
    require(p["script_sha256"] == recalc.sha(replay.__file__)
            and len(p["controls"]) == 34 and {c["name"] for c in p["controls"]} == replay._controls("v8"),
            "Frozen old 34-control roster changed")
    class Identity:
        identity = p["engine"]
    old = replay.DurableEngine(Identity(), path.parent, p["native_lock"])
    receipts = {r["record_hash"]: r for r in old.receipts()}
    require(q["execution"] == replay._execution(old) and not old.unsafe_cleanup()
            and not replay._execution(old)["open_workbook_intents"], "Old qualification has unclosed executions")
    rows = {r["name"]: r for r in q["controls"]}
    require(set(rows) == replay._controls("v8") and all(r["qualified"] for r in rows.values()),
            "Old qualification missing controls")
    for c in p["controls"]:
        saved = rows[c["name"]]
        receipt = receipts[saved["receipt_hash"]]
        require(receipt["input_sha256"] == recalc.sha(c["path"]), "Old control bytes changed")
        with tempfile.TemporaryDirectory(prefix="passive-old-control-") as temp:
            actual = seal({"protocol_hash": p["record_hash"],
                           **replay._control_result(c, receipt, Path(temp).resolve())})
        require(actual == saved, "Old control judgment changed")
    inventory = {**p["inventory"], **p["engine"]["sources"], **replay._files(list(path.parent.rglob("*.json")))}
    return p, inventory


def _fixture(path, name):
    """Only authors labelled engineering controls, never real artifacts."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 2
    ws["B1"] = "=A1+2" if name in {"passive_wrong", "display_href_wrong"} else "=TODAY()" if name == "passive_volatile" else "=A1+1"
    ws["C1"] = "Documentation"
    ws["C1"].hyperlink = {"passive_file": "file:///etc/passwd", "passive_http": "http://example.org/",
                           "passive_credentials": "https://user:secret@example.org/"}.get(name, "https://example.org/help")
    ws["C1"].hyperlink.display = "Documentation"
    ws["C1"].hyperlink.tooltip = "Read documentation"
    if name in DISPLAY_EXPECTED:
        ws["C1"].hyperlink.display = ws["C1"].hyperlink.target
        ws["C1"].hyperlink.tooltip = None
    wb.save(path)
    wb.close()
    if name == "passive_orphan":
        with zipfile.ZipFile(path) as z:
            parts = [(i, z.read(i.filename)) for i in z.infolist()]
        with zipfile.ZipFile(path, "w") as z:
            for item, raw in parts:
                if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                    tree = ET.fromstring(raw)
                    tree[0].set("Id", "unbound")
                    raw = ET.tostring(tree)
                z.writestr(item, raw)


def _passive_fixtures(path, legacy_controls, image, *, prior=False):
    """Reuse all seven original inputs exactly; never donor authorization."""
    path = safe_path(path)
    q, p = read_json(path, sealed=True), read_json(path.parent / "protocol.json", sealed=True)
    version = "passive-hyperlinks-frozen-h-replay-v2" if prior else "passive-hyperlinks-frozen-h-replay-v1"
    require(q["version"] == p["version"] == version
            and p["kind"] == "qualification" and p["root"] == str(path.parent)
            and q["status"] in {"qualified", "rejected"} and q["protocol_hash"] == p["record_hash"]
            and q["engine"] == p["engine"] and q["engine"]["image_id"] == image
            and q["model_api_calls"] == 0, "Original A fixture evidence required")
    if prior:
        require(q["status"] == "qualified" and q["score_qualification"]["status"] == "qualified"
                and len(q["score_qualification"]["controls"]) == 18
                and len(q["link_score_controls"]) == 2
                and {r["name"]: (r["expected"], r["qualified"], r["result"]["status"])
                     for r in q["link_score_controls"]}
                    == {"passive_correct": ("pass", True, "pass"), "passive_wrong": ("fail", True, "fail")},
                "Complete B fixture evidence required")
    replay._verify_files(p["sources"])
    replay._verify_files(p["inventory"])
    replay._verify_files(p["engine"]["sources"])
    controls = {c["name"]: c for c in p["controls"]}
    rows = {r["name"]: r for r in q["controls"]}
    require(len(controls) == len(p["controls"]) == len(rows) == len(q["controls"]) == 41
            and set(controls) == set(rows) == replay._controls("v8") | set(NEW_EXPECTED), "Original A roster changed")
    legacy = {c["name"]: c for c in legacy_controls}
    require(len(legacy) == len(legacy_controls) == 34 and set(legacy) == replay._controls("v8"), "Legacy roster changed")
    for name, expected in legacy.items():
        actual = controls[name]
        require(actual["expected"] == expected["expected"] and recalc.sha(actual["path"]) == recalc.sha(expected["path"]),
                "A and original 34 fixture donors differ")
    class Identity:
        identity = p["engine"]
    engine = replay.DurableEngine(Identity(), path.parent, p["native_lock"])
    receipts = {r["record_hash"]: r for r in engine.receipts()}
    require(q["execution"] == replay._execution(engine) and not engine.unsafe_cleanup()
            and not replay._execution(engine)["open_workbook_intents"], "A donor has unclosed execution")
    for name, control in controls.items():
        row = rows[name]
        require(read_json(path.parent / "controls" / (name + ".json"), sealed=True) == row
                and row["protocol_hash"] == p["record_hash"] and row["receipt_hash"] in receipts
                and receipts[row["receipt_hash"]]["input_sha256"] == recalc.sha(control["path"]), "A control binding changed")
        if name in NEW_EXPECTED:
            require(control["expected"] == NEW_EXPECTED[name], "A passive expected changed")
    inventory = {**p["sources"], **p["inventory"], **p["engine"]["sources"],
                 **replay._files(list(path.parent.rglob("*.json")))}
    return [controls[name] for name in NEW_EXPECTED], inventory


def _control(control, receipt, temp):
    expected = {**NEW_EXPECTED, **DISPLAY_EXPECTED}
    if control["name"] not in expected:
        return replay._control_result(control, receipt, temp)
    require(control["expected"] == expected[control["name"]], "New control expected changed")
    row = {"name": control["name"], "qualified": False, "reason": receipt["reason"],
           "receipt_hash": receipt["record_hash"], "cleanup_confirmed": receipt["cleanup_confirmed"]}
    if control["expected"] is None:
        row["qualified"] = (receipt["status"] == "unknown" and receipt["reason"] == "unsupported_external_relationship"
                            and receipt.get("container_execution_attempted") is False)
    elif receipt["status"] == "available":
        path = temp / "result.xlsx"
        path.write_bytes(base64.b64decode(receipt["output_base64"], validate=True))
        require(recalc.sha(path) == receipt["output_sha256"], "Link control output differs")
        preserved = links.verify_preserved(control["path"], path)
        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            actual = wb["S"]["B1"].value
            row["qualified"] = (recalc._recalc_valid(control["path"], path) is None and preserved["hyperlink_count"] == 1
                and type(actual) in (int, float) and actual == control["expected"]
                and (control["name"] not in {"passive_wrong", "display_href_wrong"} or actual != NEW_EXPECTED["passive_correct"])
                and receipt.get("execution_input_sha256") == recalc.sha(control["path"])
                and receipt.get("container_execution_attempted") is True)
            if control["name"] in DISPLAY_EXPECTED:
                proof = receipt.get("display_comparison_view", {})
                row["qualified"] = (row["qualified"] and proof.get("version") == links.DISPLAY_VIEW
                                    and type(proof.get("display_attributes_restored")) is int
                                    and proof["display_attributes_restored"] == 1)
        finally:
            wb.close()
    row["qualified"] = bool(row["qualified"] and row["cleanup_confirmed"] is True)
    return row


def _link_scores(protocol, engine):
    """Same passive links, real computed correct/wrong candidates, one fixed H."""
    reference = Path(protocol["link_score_reference"])
    manifest = scoring.freeze_reference(reference, protocol["inventory"][str(reference)], "S!B1",
        provenance={"source_protocol_hash": protocol["record_hash"], "task_id": "passive-link-score",
                    "evidence_kind": "engineering_fixture"})
    rows = []
    controls = {c["name"]: c for c in protocol["controls"]}
    pairs = [("passive_correct", "pass"), ("passive_wrong", "fail")]
    if protocol.get("version") == VERSION:
        pairs.extend((("display_href_correct", "pass"), ("display_href_wrong", "fail")))
    for name, expected in pairs:
        control = controls[name]
        actual = scoring.evaluate(control["path"], reference, "S!B1", engine, manifest)
        rows.append({"name": name, "expected": expected, "qualified": actual["status"] == expected,
                     "reference_manifest_hash": manifest["record_hash"], "result": actual})
    return rows


def qualify(output, original_qualification=None, native_lock=None, passive_fixture_qualification=None,
            prior_passive_qualification=None):
    root = safe_path(output)
    if not root.exists():
        require(original_qualification and native_lock and prior_passive_qualification
                and passive_fixture_qualification is None, "New C qualification requires the frozen B fixture donor")
        oldpath = safe_path(original_qualification)
        passive_path = safe_path(prior_passive_qualification)
        require(all(not root.is_relative_to(p) and not p.is_relative_to(root)
                    for p in (oldpath.parent, passive_path.parent)), "Separate qualification output required")
        old, inventory = _old_controls(oldpath)
        additional, donor_inventory = _passive_fixtures(passive_path, old["controls"], old["engine"]["image_id"], prior=True)
        inventory.update(donor_inventory)
        lock = replay._lock_path(native_lock, (root, oldpath.parent, passive_path.parent))
        engine = links.Recalculator(old["engine"]["image_id"], old["engine"]["timeout_seconds"])
        root.mkdir(parents=True, mode=0o700)
        controls = list(old["controls"]) + additional
        for name, expected in DISPLAY_EXPECTED.items():
            path = root / (name + ".xlsx")
            _fixture(path, name)
            inventory.update(replay._files([path]))
            controls.append({"name": name, "path": str(path), "expected": expected})
        score_reference = root / "passive_score_reference.xlsx"
        scoring._fixture_book(score_reference, {"B1": 3})
        inventory.update(replay._files([score_reference]))
        write_json(root / "protocol.json", seal({"version": VERSION, "kind": "qualification", "root": str(root),
            "sources": _sources(), "engine": engine.identity, "native_lock": str(lock), "controls": controls,
            "inventory": inventory, "link_score_reference": str(score_reference),
            "prior_passive_qualification": str(passive_path),
            "old_fixture_protocol_hash": old["record_hash"], "model_api_calls": 0}))
    with output_lock(root):
        p = read_json(root / "protocol.json", sealed=True)
        require(p["kind"] == "qualification" and len(p["controls"]) == 43
                and {c["name"] for c in p["controls"]} == replay._controls("v8") | set(NEW_EXPECTED) | set(DISPLAY_EXPECTED),
                "Control roster changed")
        engine = _engine(p, root)
        require(not engine.unsafe_cleanup() and not replay._execution(engine)["open_workbook_intents"], "Unclosed engine execution")
        rows = []
        for c in p["controls"]:
            receipt = engine.run(Path(c["path"]))
            with tempfile.TemporaryDirectory(prefix="passive-control-") as temp:
                row = seal({"protocol_hash": p["record_hash"], **_control(c, receipt, Path(temp).resolve())})
            location = root / "controls" / (c["name"] + ".json")
            if location.exists():
                require(read_json(location, sealed=True) == row, "Existing control judgment changed")
            else:
                write_json(location, row)
            rows.append(row)
            if engine.unsafe_cleanup():
                break
        scored = scoring.qualify(root / "score_fixtures", engine) if not engine.unsafe_cleanup() else None
        link_scores = _link_scores(p, engine) if not engine.unsafe_cleanup() else []
        status = ("qualified" if len(rows) == 43 and all(c["qualified"] for c in rows)
                  and scored and scored["status"] == "qualified" and len(scored["controls"]) == 18
                  and len(link_scores) == 4 and all(c["qualified"] for c in link_scores) else "rejected")
        value = seal({"version": VERSION, "protocol_hash": p["record_hash"], "status": status, "controls": rows,
                      "score_qualification": scored, "link_score_controls": link_scores,
                      "engine": engine.identity, "sources": _sources(),
                      "execution": replay._execution(engine), "evidence_kind": "engineering_fixture", "model_api_calls": 0})
        destination = root / "qualification.json"
        if destination.exists():
            require(read_json(destination, sealed=True) == value, "Published qualification changed")
        else:
            write_json(destination, value)
        return value


def _previous(root):
    root = safe_path(root)
    p = read_json(root / "protocol.json", sealed=True)
    require(p["version"] == frozen.VERSION and p["kind"] == "replay" and p["root"] == str(root), "Frozen C replay required")
    replay._verify_files(p["inventory"])
    replay._verify_files(p["sources"])
    replay._verify_files(p["engine"]["sources"])
    require({Path(k).name: h for k, h in p["sources"].items()} ==
            {Path(k).name: h for k, h in frozen._sources().items()}, "C scorer policy differs")
    frozen._published(root, p)
    class Identity:
        identity = p["engine"]
    engine = replay.DurableEngine(Identity(), root, p["native_lock"])
    require(not engine.unsafe_cleanup() and not replay._execution(engine)["open_workbook_intents"], "C execution incomplete")
    receipt_ids = {r["record_hash"] for r in engine.receipts()}
    require(len(p["slots"]) == len({s["id"] for s in p["slots"]}) == 160
            and len({s["task_id"] for s in p["slots"]}) == 80
            and {f.stem for f in (root / "positions").glob("*.json")} == {s["id"] for s in p["slots"]}, "Full C roster required")
    rows = {}
    for slot in p["slots"]:
        row = read_json(root / "positions" / (slot["id"] + ".json"), sealed=True)
        require(row["protocol_hash"] == p["record_hash"] and row["slot_id"] == slot["id"]
                and row["old_status"] == slot["old_status"] and row["previous_status"] == slot["previous_status"], "C position identity changed")
        require(read_json(slot["prediction_path"], sealed=True)["record_hash"] == slot["prediction_hash"], "Original prediction changed")
        require(all(h in receipt_ids for c in row["cases"] for h in c.get("receipts", {}).values()), "C receipt missing")
        rows[slot["id"]] = row
    require(Counter(r["status"] for r in rows.values()) == {"pass": 79, "fail": 55, "unknown": 26}, "Frozen C counts differ")
    return p, rows


def prepare(previous_replay, qualification, output, native_lock):
    previous, qpath, root = map(safe_path, (previous_replay, qualification, output))
    require(not root.exists() and all(not root.is_relative_to(p) and not p.is_relative_to(root)
            for p in (previous, qpath.parent)), "Separate new output required")
    q = read_json(qpath, sealed=True)
    require(q["status"] == "qualified" and qualify(qpath.parent) == q, "Complete new link qualification required")
    lock = replay._lock_path(native_lock, (root, previous, qpath.parent))
    with _readonly_parent_lock(previous):
        p, rows = _previous(previous)
        inventory = {**p["inventory"], **p["sources"], **p["engine"]["sources"],
                     **replay._files(list(previous.rglob("*.json")))}
        qp = read_json(qpath.parent / "protocol.json", sealed=True)
        inventory.update(qp["inventory"])
        inventory.update(replay._files(list(qpath.parent.rglob("*.json")) + list(qpath.parent.rglob("*.xlsx"))))
        root.mkdir(parents=True, mode=0o700)
        slots = []
        for slot in p["slots"]:
            for key in slot["reference_manifests"]:
                src = previous / "host_only/references" / (key + ".json")
                dst = root / "host_only/references" / (key + ".json")
                if not dst.exists():
                    write_json(dst, read_json(src, sealed=True))
                    inventory.update(replay._files([dst]))
            slots.append({**slot, "previous_status": rows[slot["id"]]["status"], "previous_result_hash": rows[slot["id"]]["record_hash"]})
        write_json(root / "protocol.json", seal({"version": VERSION, "kind": "replay", "root": str(root),
            "sources": _sources(), "engine": q["engine"], "native_lock": str(lock), "inventory": inventory,
            "qualification_hash": q["record_hash"], "previous_replay": str(previous), "previous_protocol_hash": p["record_hash"],
            "original_source": p["original_source"], "slots": slots, "truth_basis": p["truth_basis"],
            "model_api_calls": 0, "historical_scores_replaced": False, "feedback_allowed": False}))
    return {"status": "prepared", "positions": len(slots)}


def report(output):
    root = safe_path(output)
    p = read_json(root / "protocol.json", sealed=True)
    require(p["kind"] == "replay", "Replay protocol required")
    engine = _engine(p, root)
    frozen._published(root, p)
    rows = []
    for slot in p["slots"]:
        path = root / "positions" / (slot["id"] + ".json")
        if path.exists():
            row = read_json(path, sealed=True)
            require(row["protocol_hash"] == p["record_hash"] and row["slot_id"] == slot["id"]
                    and row["old_status"] == slot["old_status"] and row["previous_status"] == slot["previous_status"], "Position identity differs")
            rows.append(row)
    costs = replay._execution(engine)
    known = [r for r in rows if r["previous_status"] != "unknown"]
    return seal({"version": VERSION, "protocol_hash": p["record_hash"],
        "status": "complete" if len(rows) == 160 and not costs["open_workbook_intents"] and not costs["cleanup_unconfirmed"] else "pending",
        "positions": 160, "completed": len(rows), "counts": dict(Counter(r["status"] for r in rows)),
        "previous_counts": dict(Counter(s["previous_status"] for s in p["slots"])),
        "previous_transitions": dict(Counter(r["previous_status"] + "->" + r["status"] for r in rows)),
        "original_transitions": dict(Counter(r["old_status"] + "->" + r["status"] for r in rows)),
        "known_preservation": {"expected_positions": 134, "replayed": len(known),
                               "unchanged": sum(r["previous_status"] == r["status"] for r in known)},
        "unknown_reasons": dict(Counter(c["reason"] for r in rows for c in r["cases"] if c["status"] == "unknown")),
        "result_hashes": {r["slot_id"]: r["record_hash"] for r in rows}, "execution": costs,
        "model_api_calls": 0, "reference_engine_calls": 0, "historical_scores_replaced": False,
        "feedback_allowed": False, "is_method_effect": False, "truth_basis": p["truth_basis"]})


def run(output):
    root = safe_path(output)
    with output_lock(root):
        p = read_json(root / "protocol.json", sealed=True)
        engine = _engine(p, root)
        require(p["kind"] == "replay" and not engine.unsafe_cleanup()
                and not replay._execution(engine)["open_workbook_intents"], "Wrong protocol or unclosed execution")
        report(root)
        with _readonly_parent_lock(Path(p["previous_replay"])), _readonly_parent_lock(Path(p["original_source"])):
            for slot in p["slots"]:
                path = root / "positions" / (slot["id"] + ".json")
                if path.exists():
                    continue
                if (root / "PAUSE").exists():
                    break
                saved = read_json(slot["prediction_path"], sealed=True)
                require(saved["record_hash"] == slot["prediction_hash"], "Original prediction changed")
                prediction = saved["prediction"]
                cases = (prediction.get("output") or {}).get("cases", [])
                results, delivery = [], "undelivered"
                if prediction["status"] != "available" or not cases:
                    results = [recalc._undelivered(prediction["reason"])]
                elif len(cases) != len(slot["references"]):
                    results, delivery = [{"status": "unknown", "reason": "original_case_count_mismatch"}], "case_count_mismatch"
                else:
                    require(len(slot["references"]) == len(slot["reference_manifests"]), "Reference roster changed")
                    delivered = sum(c["status"] == "available" for c in cases)
                    delivery = "delivered" if delivered == len(cases) else "partial" if delivered else "undelivered"
                    with tempfile.TemporaryDirectory(prefix="passive-frozen-candidate-") as temp:
                        for index, (case, ref, key) in enumerate(zip(cases, slot["references"], slot["reference_manifests"])):
                            if case["status"] != "available":
                                results.append(recalc._undelivered(case["reason"]))
                                continue
                            raw = base64.b64decode(case["output_base64"], validate=True)
                            require(len(raw) <= recalc.MAX_BYTES, "Workbook size limit")
                            candidate = Path(temp).resolve() / f"{index}.xlsx"
                            candidate.write_bytes(raw)
                            manifest = read_json(root / "host_only/references" / (key + ".json"), sealed=True)
                            results.append(scoring.evaluate(candidate, ref, slot["answer_position"], engine, manifest))
                            if engine.unsafe_cleanup():
                                break
                counts = Counter(c["status"] for c in results)
                status = "unknown" if engine.unsafe_cleanup() else "fail" if counts["fail"] else "unknown" if counts["unknown"] else "pass"
                write_json(path, seal({"protocol_hash": p["record_hash"], "slot_id": slot["id"],
                    "old_status": slot["old_status"], "previous_status": slot["previous_status"],
                    "status": status, "delivery_status": delivery, "cases": results, "model_api_calls": 0}))
                if engine.unsafe_cleanup():
                    break
        value = report(root)
        write_json(root / "reports" / (value["record_hash"] + ".json"), value)
        return value


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("qualify", "prepare", "run", "report"))
    parser.add_argument("--output", required=True)
    parser.add_argument("--original-qualification")
    parser.add_argument("--passive-fixture-qualification")
    parser.add_argument("--prior-passive-qualification")
    parser.add_argument("--qualification")
    parser.add_argument("--previous-replay")
    parser.add_argument("--native-lock")
    args = parser.parse_args(argv)
    if args.command == "qualify":
        value = qualify(args.output, args.original_qualification, args.native_lock, args.passive_fixture_qualification,
                        args.prior_passive_qualification)
    elif args.command == "prepare":
        require(all((args.previous_replay, args.qualification, args.native_lock)), "Missing preparation arguments")
        value = prepare(args.previous_replay, args.qualification, args.output, args.native_lock)
    else:
        value = {"run": run, "report": report}[args.command](args.output)
    print(json.dumps(value, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
