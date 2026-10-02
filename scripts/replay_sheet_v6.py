"""Independent v6 qualification and zero-API full-panel workbook replay.

Old qualification contributes fixture inputs only, never v6 authorization.
Original model outputs and scores are immutable; missing workbooks stay unknown.
"""
from __future__ import annotations

import argparse
import base64
import datetime
import fcntl
import json
import tempfile
from collections import Counter
from pathlib import Path

import openpyxl

from skillopt.coevolution_v5.core import seal
from skillopt.continual_eval import sheet_recalc as v5
from skillopt.continual_eval import sheet_recalc_functions as v8
from skillopt.continual_eval import sheet_recalc_v6 as v6
from skillopt.continual_eval import sheet_recalc_v7 as v7
from skillopt.continual_eval.core import output_lock, panel_tasks, read_json, require, safe_path, write_json
from skillopt.continual_eval.sheet_recalc_adapter import CONTROLS
from skillopt.continual_eval.truncation_recovery import _readonly_parent_lock
from skillopt.validator_pilot.api import digest

VERSION = "spreadsheet-v6-independent-replay-v1"
NEW_CONTROLS = {"format30_correct", "format30_wrong", "format30_observer"}
V7_CONTROLS = {"original_date_literal", "original_date_formula", "general_wrong_numeric"}
V8_EXPECTED = {
    "ifna_fallback": 5, "ifna_passthrough": 7, "ifna_error_not_na": "#DIV/0!",
    "ifna_fixed_array": {"B1": 7, "B2": 5}, "aggregate_ignore_error": 2,
    "aggregate_array_filter": 1, "aggregate_wrong": 4,
    "unknown_extended_rejected": None, "alias_observer_rejected": None,
    "alias_named_formula_rejected": None, "alias_text_not_rewritten": None,
}


def _profile(protocol):
    profile = protocol.get("engine_profile", "v6")
    require(profile in {"v6", "v7", "v8"}, "Unknown engine profile")
    return profile


def _engine_type(profile):
    require(profile in {"v6", "v7", "v8"}, "Unknown engine profile")
    return {"v6": v6.Recalculator, "v7": v7.Recalculator, "v8": v8.Recalculator}[profile]


def _controls(profile):
    _engine_type(profile)
    return (CONTROLS | NEW_CONTROLS | (V7_CONTROLS if profile in {"v7", "v8"} else set())
            | (set(V8_EXPECTED) if profile == "v8" else set()))


def _files(paths):
    return {str(safe_path(p)): v5.sha(p) for p in paths}


def _verify_files(inventory):
    require(all(v5.sha(p) == expected for p, expected in inventory.items()), "Frozen input changed")


def _lock_path(path, excluded=()):
    path = safe_path(path)
    require(path.is_file() and not any(path.is_relative_to(p) for p in excluded), "External native lock required")
    return path


class DurableEngine:
    """Exactly one attempt per immutable input and engine, including crashes."""
    def __init__(self, engine, root, native_lock):
        self.engine, self.root = engine, safe_path(root)
        self.lock = _lock_path(native_lock, (self.root,))

    @property
    def identity(self):
        return self.engine.identity

    def receipts(self):
        rows = []
        for path in sorted((self.root / "recalculations").glob("*.json")):
            row = read_json(path, sealed=True)
            expected = {"input": row["input_sha256"], "engine": self.identity}
            require(row["identity"] == self.identity and path.stem == digest(expected)
                    and read_json(self.root / "intents" / path.name, sealed=True) == seal(expected),
                    "Workbook receipt binding changed")
            rows.append(row)
        return rows

    def unsafe_cleanup(self):
        return any(row.get("cleanup_confirmed") is not True for row in self.receipts())

    def run(self, path):
        require(not self.unsafe_cleanup(), "Previous workbook cleanup unconfirmed")
        require(_execution(self)["open_workbook_intents"] == 0, "Open workbook intent; no reexecution")
        request = {"input": v5.sha(path), "engine": self.identity}
        name = digest(request) + ".json"
        terminal, intent = self.root / "recalculations" / name, self.root / "intents" / name
        if terminal.exists():
            self.receipts()
            return read_json(terminal, sealed=True)
        require(not intent.exists(), "Open workbook intent; no reexecution")
        write_json(intent, seal(request))
        with self.lock.open("rb") as lock:
            fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
            row = self.engine.run(path)
        require(row["input_sha256"] == request["input"] and row["identity"] == self.identity
                and row["model_api_calls"] == 0, "Engine receipt mismatch")
        write_json(terminal, row)
        return row


def _execution(engine):
    rows = engine.receipts()
    closed = {digest({"input": row["input_sha256"], "engine": row["identity"]}) for row in rows}
    intents = {p.stem for p in (engine.root / "intents").glob("*.json")}
    require(closed <= intents, "Workbook execution missing intent")
    return {"closed_workbooks": len(rows), "open_workbook_intents": len(intents - closed),
            "containers_attempted": sum(row.get("container_execution_attempted", False) for row in rows),
            "reported_wall_seconds": sum(row.get("duration_seconds", 0.) for row in rows),
            "cleanup_unconfirmed": sum(row.get("cleanup_confirmed") is not True for row in rows),
            "model_api_calls": 0}


def _engine(protocol, root):
    require(protocol["version"] == VERSION and protocol["script_sha256"] == v5.sha(__file__),
            "Replay source changed")
    _verify_files(protocol["inventory"])
    engine = _engine_type(_profile(protocol))(protocol["engine"]["image_id"], protocol["engine"]["timeout_seconds"])
    require(engine.identity == protocol["engine"], "Engine/source/image changed")
    return DurableEngine(engine, root, protocol["native_lock"])


def _published(root, protocol):
    """Published evidence is immutable while subsequent positions may finish."""
    for path in (root / "reports").glob("*.json"):
        saved = read_json(path, sealed=True)
        require(saved["protocol_hash"] == protocol["record_hash"], "Published protocol differs")
        if protocol["kind"] == "qualification":
            for row in saved["controls"]:
                require(read_json(root / "controls" / (row["name"] + ".json"), sealed=True) == row,
                        "Published qualification control changed")
        else:
            for key, expected in saved["result_hashes"].items():
                require(read_json(root / "positions" / (key + ".json"), sealed=True)["record_hash"] == expected,
                        "Published replay result changed")


def _legacy_controls(path, image):
    path = safe_path(path)
    q = read_json(path, sealed=True)
    require(q["version"] == v5.VERSION and q["status"] == "qualified"
            and q["engine"]["image_id"] == image and q["model_api_calls"] == 0
            and len(q["controls"]) == len(CONTROLS)
            and {row["name"] for row in q["controls"]} == CONTROLS,
            "Complete original 17-control fixture source required")
    old_sources = q["engine"]["sources"]
    require({Path(p).name: h for p, h in old_sources.items()}
            == {Path(p).name: h for p, h in v5._sources().items()}, "Historical v5 implementation differs")
    _verify_files(old_sources)
    inventory, controls = _files([path]), []
    inventory.update(old_sources)
    for row in q["controls"]:
        name = row["name"]
        fixture, receipt_path = path.parent / "fixtures" / (name + ".xlsx"), path.parent / "receipts" / (name + ".json")
        receipt = read_json(receipt_path, sealed=True)
        require(row["qualified"] is True and receipt["record_hash"] == row["receipt_hash"]
                and receipt["input_sha256"] == v5.sha(fixture) and receipt["identity"] == q["engine"]
                and receipt["cleanup_confirmed"] is True, "Historical fixture/receipt differs")
        inventory.update(_files([fixture, receipt_path]))
        controls.append({"name": name, "path": str(fixture), "expected": row["expected"]})
    return controls, inventory


def _comparison_fixtures(path, image, legacy_controls):
    """Reuse rejected/qualified q20 inputs, never its engine authorization."""
    path = safe_path(path)
    q = read_json(path, sealed=True)
    protocol_path = path.parent / "protocol.json"
    protocol = read_json(protocol_path, sealed=True)
    require(q["version"] == protocol["version"] == VERSION
            and q["status"] in {"qualified", "rejected"} and protocol["kind"] == "qualification"
            and protocol["root"] == str(path.parent) and _profile(q) == _profile(protocol) == "v6"
            and q["protocol_hash"] == protocol["record_hash"] and q["engine"] == protocol["engine"]
            and q["engine"]["image_id"] == image, "Frozen q20 fixture provenance differs")
    specs = {c["name"]: c for c in protocol["controls"]}
    rows = {c["name"]: c for c in q["controls"]}
    require(len(specs) == len(protocol["controls"]) == len(rows) == len(q["controls"]) == 20
            and set(specs) == set(rows) == _controls("v6"), "Complete frozen q20 fixture roster required")
    inventory, controls = _files([path, protocol_path]), []
    legacy = {c["name"]: c for c in legacy_controls}
    require(len(legacy) == len(legacy_controls) == len(CONTROLS) and set(legacy) == CONTROLS,
            "Complete legacy fixture roster required")
    for name in sorted(CONTROLS):
        original, current = specs[name], legacy[name]
        source = safe_path(original["path"])
        require(protocol["inventory"].get(str(source)) == v5.sha(source) == v5.sha(current["path"])
                and digest(original["expected"]) == digest(current["expected"]),
                "Legacy q17 and frozen q20 fixture bytes/expected differ")
        inventory.update(_files([source]))
    for name in sorted(NEW_CONTROLS):
        spec, row = specs[name], rows[name]
        source = safe_path(spec["path"])
        require(spec["expected"] == (46002 if name == "format30_wrong" else 46001)
                and protocol["inventory"].get(str(source)) == v5.sha(source), "Frozen comparison fixture changed")
        row_path = path.parent / "controls" / (name + ".json")
        require(read_json(row_path, sealed=True) == row and row["protocol_hash"] == protocol["record_hash"],
                "Frozen comparison fixture judgment differs")
        request = {"input": v5.sha(source), "engine": q["engine"]}
        key = digest(request) + ".json"
        receipt_path, intent_path = path.parent / "recalculations" / key, path.parent / "intents" / key
        receipt = read_json(receipt_path, sealed=True)
        require(receipt["record_hash"] == row["receipt_hash"] and receipt["input_sha256"] == request["input"]
                and receipt["identity"] == q["engine"] and receipt["model_api_calls"] == 0
                and read_json(intent_path, sealed=True) == seal(request), "Frozen comparison fixture binding differs")
        inventory.update(_files([source, row_path, receipt_path, intent_path]))
        controls.append({"name": name, "path": str(source), "expected": spec["expected"]})
    return controls, inventory


def _prior_profile_fixtures(path, image, legacy_controls):
    """Reuse exact q23 roles/inputs/expected values, not its authorization."""
    path = safe_path(path)
    q = read_json(path, sealed=True)
    protocol_path = path.parent / "protocol.json"
    protocol = read_json(protocol_path, sealed=True)
    require(q["version"] == protocol["version"] == VERSION
            and q["status"] in {"qualified", "rejected"} and protocol["kind"] == "qualification"
            and protocol["root"] == str(path.parent) and _profile(q) == _profile(protocol) == "v7"
            and q["protocol_hash"] == protocol["record_hash"] and q["engine"] == protocol["engine"]
            and q["engine"]["version"] == v7.VERSION and q["engine"]["image_id"] == image
            and q["model_api_calls"] == protocol["model_api_calls"] == 0,
            "Frozen q23 fixture provenance differs")
    specs, rows = {c["name"]: c for c in protocol["controls"]}, {c["name"]: c for c in q["controls"]}
    require(len(specs) == len(protocol["controls"]) == len(rows) == len(q["controls"]) == 23
            and set(specs) == set(rows) == _controls("v7"), "Complete frozen q23 fixture roster required")
    prior, inventory = _comparison_fixtures(protocol["comparison_fixture_qualification"], image, legacy_controls)
    original_twenty = {c["name"]: c for c in [*legacy_controls, *prior]}
    inventory.update(_files([path, protocol_path]))
    _verify_files(q["engine"]["sources"])
    inventory.update(q["engine"]["sources"])
    controls = []
    for name in sorted(specs):
        spec, row = specs[name], rows[name]
        source = safe_path(spec["path"])
        require(protocol["inventory"].get(str(source)) == v5.sha(source), "Frozen q23 fixture bytes differ")
        if name in original_twenty:
            original = original_twenty[name]
            require(v5.sha(original["path"]) == v5.sha(source)
                    and digest(original["expected"]) == digest(spec["expected"]),
                    "Original q17/q20 and frozen q23 fixture bytes/expected differ")
        else:
            expected = 46002 if name == "general_wrong_numeric" else {
                "type": "datetime", "iso8601": "2025-12-09T00:00:00"}
            require(digest(spec["expected"]) == digest(expected), "Frozen q23 expected value differs")
        row_path = path.parent / "controls" / (name + ".json")
        require(read_json(row_path, sealed=True) == row and row["protocol_hash"] == protocol["record_hash"],
                "Frozen q23 control judgment differs")
        request = {"input": v5.sha(source), "engine": q["engine"]}
        key = digest(request) + ".json"
        receipt_path, intent_path = path.parent / "recalculations" / key, path.parent / "intents" / key
        receipt = read_json(receipt_path, sealed=True)
        require(receipt["record_hash"] == row["receipt_hash"] and receipt["input_sha256"] == request["input"]
                and receipt["identity"] == q["engine"] and receipt["model_api_calls"] == 0
                and read_json(intent_path, sealed=True) == seal(request), "Frozen q23 fixture receipt binding differs")
        inventory.update(_files([source, row_path, receipt_path, intent_path]))
        controls.append({"name": name, "path": str(source), "expected": spec["expected"]})
    return controls, inventory


def _function_fixture(path, name):
    """Authored standard-function controls, not real benchmark outcomes."""
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.formula import ArrayFormula

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet["A1"] = 4
    formulas = {"ifna_fallback": "=_xlfn.IFNA(NA(),5)", "ifna_passthrough": "=_xlfn.IFNA(7,5)",
        "ifna_error_not_na": "=_xlfn.IFNA(1/0,5)",
        "aggregate_ignore_error": "=_xlfn.AGGREGATE(15,6,A1:A3,1)",
        "aggregate_wrong": "=_xlfn.AGGREGATE(15,6,A1:A3,2)",
        "aggregate_array_filter": "=_xlfn.AGGREGATE(15,6,ROW(A1:A3)/(A1:A3=1),1)",
        "unknown_extended_rejected": "=_xlfn.UNIQUE(A1:A2)",
        "alias_observer_rejected": "=_xlfn.IFNA(NA(),5)",
        "alias_named_formula_rejected": "=_xlfn.IFNA(MyValue,5)",
        "alias_text_not_rewritten": '="_xlfn.IFNA("'}
    if name == "ifna_fixed_array":
        sheet["A1"], sheet["A2"] = 1, 3
        sheet["D1"], sheet["E1"], sheet["D2"], sheet["E2"] = 1, 7, 2, 9
        sheet["B1"] = ArrayFormula(ref="B1:B2", text="=_xlfn.IFNA(VLOOKUP(A1:A2,D1:E2,2,FALSE()),5)")
    else:
        sheet["B1"] = formulas[name]
    if name in {"aggregate_ignore_error", "aggregate_wrong"}:
        sheet["A2"], sheet["A3"] = "=1/0", 2  # Formula, not unsupported error-literal coercion.
    if name == "aggregate_array_filter":
        sheet["A1"], sheet["A2"], sheet["A3"] = 1, 0, 1
    if name == "alias_observer_rejected":
        sheet["C1"] = "=FORMULATEXT(B1)"
    if name == "alias_named_formula_rejected":
        book.defined_names.add(DefinedName("MyValue", attr_text="'S'!$A$1"))
    book.save(path)
    book.close()


def _validate_function_controls(protocol):
    """A resealed proposal cannot relax a v8 control's role/expected value."""
    if _profile(protocol) != "v8":
        return
    legacy, _ = _legacy_controls(protocol["legacy_qualification"], protocol["engine"]["image_id"])
    prior, _ = _prior_profile_fixtures(protocol["prior_profile_qualification"], protocol["engine"]["image_id"], legacy)
    specs = {c["name"]: c for c in protocol["controls"]}
    require(all(specs[c["name"]] == c for c in prior), "Frozen q23 role/bytes/expected changed in v8")
    require(all(digest(specs[name]["expected"]) == digest(expected) for name, expected in V8_EXPECTED.items()),
            "Frozen v8 expected values changed")


def _qual_setup(output, image, legacy_qualification, native_lock, engine_version="v6", comparison_fixture_qualification=None,
                prior_profile_qualification=None):
    root = safe_path(output)
    require(not root.exists(), "New qualification directory required")
    engine = _engine_type(engine_version)(image)
    require(engine_version == "v7" or comparison_fixture_qualification is None,
            "Only v7 accepts prior comparison fixtures")
    require(engine_version == "v8" or prior_profile_qualification is None, "Only v8 accepts prior profile fixtures")
    controls, inventory = _legacy_controls(legacy_qualification, image)
    require(engine_version != "v7" or comparison_fixture_qualification,
            "v7 requires frozen q20 comparison fixtures")
    if engine_version == "v7":
        prior_controls, prior_inventory = _comparison_fixtures(comparison_fixture_qualification, image, controls)
        controls.extend(prior_controls)
        inventory.update(prior_inventory)
    if engine_version == "v8":
        require(prior_profile_qualification, "v8 requires frozen q23 profile fixtures")
        controls, prior_inventory = _prior_profile_fixtures(prior_profile_qualification, image, controls)
        inventory.update(prior_inventory)
    lock = _lock_path(native_lock, (root, safe_path(legacy_qualification).parent))
    root.mkdir(parents=True, mode=0o700)
    for name in sorted(NEW_CONTROLS if engine_version == "v6" else ()):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "S"
        sheet["A1"] = 46000
        sheet["A1"].number_format = "0"
        sheet["A1"]._style.numFmtId = 30
        sheet["B1"] = ('=CELL("format",A1)' if name == "format30_observer"
                       else "=A1+2" if name == "format30_wrong" else "=A1+1")
        path = root / (name + ".xlsx")
        book.save(path)
        book.close()
        inventory.update(_files([path]))
        controls.append({"name": name, "path": str(path), "expected":
                         46002 if name == "format30_wrong" else 46001})
    if engine_version == "v7":
        for name in sorted(V7_CONTROLS):
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "S"
            sheet["A1"] = 46000
            if name == "general_wrong_numeric":
                sheet["B1"] = "=A1+2"
                expected = 46002
            else:
                sheet["B1"] = 46000 if name == "original_date_literal" else "=DATE(2025,12,9)"
                sheet["B1"].number_format = "yyyy-mm-dd"
                expected = {"type": "datetime", "iso8601": "2025-12-09T00:00:00"}
            path = root / (name + ".xlsx")
            book.save(path)
            book.close()
            inventory.update(_files([path]))
            controls.append({"name": name, "path": str(path), "expected": expected})
    if engine_version == "v8":
        for name, expected in sorted(V8_EXPECTED.items()):
            path = root / (name + ".xlsx")
            _function_fixture(path, name)
            inventory.update(_files([path]))
            controls.append({"name": name, "path": str(path), "expected": expected})
    protocol = seal({"version": VERSION, "kind": "qualification", "root": str(root),
        "script_sha256": v5.sha(__file__), "engine": engine.identity, "engine_profile": engine_version,
        "native_lock": str(lock),
        "inventory": inventory, "controls": controls, "evidence_kind": "engineering_fixture",
        "comparison_fixture_qualification": str(safe_path(comparison_fixture_qualification)) if engine_version == "v7" else None,
        "model_api_calls": 0, "old_authorization_inherited": False})
    if engine_version == "v8":
        protocol.pop("record_hash")
        protocol = seal({**protocol, "prior_profile_qualification": str(safe_path(prior_profile_qualification)),
                         "legacy_qualification": str(safe_path(legacy_qualification))})
    write_json(root / "protocol.json", protocol)
    return protocol


def _control_result(control, receipt, directory):
    name, expected, source = control["name"], control["expected"], Path(control["path"])
    result = {"name": name, "qualified": False, "receipt_hash": receipt["record_hash"],
              "cleanup_confirmed": receipt["cleanup_confirmed"], "reason": receipt["reason"]}
    if name in {"utf16_external", "utf16_doctype"}:
        expected_reason = "unsupported_external_relationship" if name == "utf16_external" else "unsupported_workbook_structure"
        result["qualified"] = receipt["status"] == "unknown" and receipt["reason"] == expected_reason and not receipt["container_execution_attempted"]
    elif name in V8_EXPECTED and V8_EXPECTED[name] is None:
        result["qualified"] = (receipt["status"] == "unknown" and receipt["reason"] == "unsupported_extended_function"
                               and receipt["container_execution_attempted"] is False)
    elif receipt["status"] == "available":
        target = directory / "control.xlsx"
        target.write_bytes(base64.b64decode(receipt["output_base64"], validate=True))
        require(v5.sha(target) == receipt["output_sha256"], "Control output differs")
        reason = v5._recalc_valid(source, target)
        result["reason"] = reason
        book = openpyxl.load_workbook(target, data_only=True)
        try:
            actual = book["S"]["B1"].value
            if name == "format30_observer":
                result["qualified"] = (reason == "recalculation_changed_literal_cell"
                    and not receipt["comparison_view"]["adapted_cells"]
                    and not receipt["comparison_view"].get("formula_adapted_cells", []))
            elif name == "unsupported":
                result["qualified"] = actual == "#NAME?" and reason in {"unsupported_formula_result", "recalculation_changed_formula"}
            else:
                compare = v5.compat._legacy()._compare_cell_value
                if name in {"original_date_literal", "original_date_formula"}:
                    matches = (isinstance(expected, dict) and set(expected) == {"type", "iso8601"}
                               and expected["type"] == "datetime" and type(actual) is datetime.datetime
                               and actual.isoformat() == expected["iso8601"])
                else:
                    matches = all(compare(book["S"][cell].value, val) for cell, val in expected.items()) if isinstance(expected, dict) else compare(actual, expected)
                result["qualified"] = reason is None and matches
                if name == "empty_string":
                    result["qualified"] &= ("S", "B1") in v5.compat.explicit_empty_formula_strings(target)
                if name in {"format30_correct", "format30_wrong"}:
                    result["qualified"] &= (len(receipt["comparison_view"]["adapted_cells"]) == 1
                                             and book["S"]["A1"].value == 46000)
                wrong_gold = {"wrong": 5, "array_wrong": 10, "numeric_wrong": 1.234567890123456,
                              "boolean_wrong": 5, "boolean_string_wrong": "FALSE", "format30_wrong": 46001,
                              "general_wrong_numeric": 46001, "aggregate_wrong": 2}
                if name in wrong_gold:
                    result["qualified"] &= not compare(actual, wrong_gold[name])
                if name in V8_EXPECTED:
                    proof = receipt.get("function_alias_view", {})
                    executed = receipt.get("execution_receipt", {})
                    result["qualified"] &= (proof.get("version") == v8.ALIAS_VERSION
                        and proof.get("original_input_sha256") == v5.sha(source)
                        and executed.get("input_sha256") == proof.get("execution_input_sha256")
                        and executed.get("identity") == receipt["identity"]
                        and executed.get("cleanup_confirmed") is True and executed.get("container_execution_attempted") is True)
        finally:
            book.close()
    result["qualified"] = bool(result["qualified"] and receipt["cleanup_confirmed"] is True)
    return result


def qualify(output, image=None, legacy_qualification=None, native_lock=None, engine_version=None,
            comparison_fixture_qualification=None, prior_profile_qualification=None):
    root = safe_path(output)
    if not root.exists():
        require(all((image, legacy_qualification, native_lock)), "Qualification preparation arguments required")
        _qual_setup(root, image, legacy_qualification, native_lock, engine_version or "v6", comparison_fixture_qualification,
                    prior_profile_qualification)
    with output_lock(root):
        protocol = read_json(root / "protocol.json", sealed=True)
        require(protocol["kind"] == "qualification" and protocol["root"] == str(root), "Wrong qualification directory")
        profile = _profile(protocol)
        require(engine_version is None or engine_version == profile, "Frozen engine profile differs")
        require(comparison_fixture_qualification is None
                or str(safe_path(comparison_fixture_qualification)) == protocol.get("comparison_fixture_qualification"),
                "Frozen comparison fixture qualification differs")
        require(prior_profile_qualification is None
                or str(safe_path(prior_profile_qualification)) == protocol.get("prior_profile_qualification"),
                "Frozen prior profile qualification differs")
        require(len(protocol["controls"]) == len(_controls(profile))
                and {c["name"] for c in protocol["controls"]} == _controls(profile),
                "Frozen qualification control roster differs")
        _validate_function_controls(protocol)
        engine = _engine(protocol, root)
        require(not engine.unsafe_cleanup(), "Previous workbook cleanup unconfirmed")
        require(_execution(engine)["open_workbook_intents"] == 0, "Open workbook intent; no reexecution")
        _published(root, protocol)
        for control in protocol["controls"]:
            path = root / "controls" / (control["name"] + ".json")
            if path.exists():
                saved = read_json(path, sealed=True)
                receipts = {r["record_hash"]: r for r in engine.receipts()}
                require(saved["receipt_hash"] in receipts, "Qualification receipt missing")
                receipt = receipts[saved["receipt_hash"]]
                require(receipt["input_sha256"] == v5.sha(control["path"]), "Qualification input changed")
                with tempfile.TemporaryDirectory(prefix="sheet-v6-control-check-") as temp:
                    verified = seal({"protocol_hash": protocol["record_hash"],
                                     **_control_result(control, receipt, Path(temp).resolve())})
                require(verified == saved, "Qualification control judgment changed")
                continue
            if (root / "PAUSE").exists():
                break
            receipt = engine.run(Path(control["path"]))
            with tempfile.TemporaryDirectory(prefix="sheet-v6-control-") as temp:
                result = _control_result(control, receipt, Path(temp).resolve())
            write_json(path, seal({"protocol_hash": protocol["record_hash"], **result}))
            if engine.unsafe_cleanup():
                break
        rows = [read_json(root / "controls" / (c["name"] + ".json"), sealed=True)
                for c in protocol["controls"] if (root / "controls" / (c["name"] + ".json")).exists()]
        require(all(row["protocol_hash"] == protocol["record_hash"] for row in rows), "Control binding changed")
        costs = _execution(engine)
        status = ("pending" if len(rows) != len(protocol["controls"]) or costs["open_workbook_intents"]
                  or costs["cleanup_unconfirmed"] else "qualified" if all(r["qualified"] for r in rows) else "rejected")
        result = seal({"version": VERSION, "status": status, "protocol_hash": protocol["record_hash"],
            "engine": engine.identity, "engine_profile": profile, "controls": rows, "execution": costs,
            "evidence_kind": "engineering_fixture", "model_api_calls": 0, "excel_equivalence_proven": False})
        write_json(root / "reports" / (result["record_hash"] + ".json"), result)
        if status != "pending":
            write_json(root / "qualification.json", result)
        return result


def prepare(source, source_root, output, qualification, native_lock):
    source, source_root, root, qpath = map(safe_path, (source, source_root, output, qualification))
    require(not root.exists() and not any(root.is_relative_to(p) or p.is_relative_to(root)
            for p in (source, source_root, qpath.parent)), "Separate new replay directory required")
    lock = _lock_path(native_lock, (source, source_root, root, qpath.parent))
    q = read_json(qpath, sealed=True)
    qprotocol = read_json(qpath.parent / "protocol.json", sealed=True)
    profile = _profile(qprotocol)
    roster = _controls(profile)
    qengine = _engine(qprotocol, qpath.parent)
    _validate_function_controls(qprotocol)
    _published(qpath.parent, qprotocol)
    require(qprotocol["kind"] == "qualification" and qprotocol["root"] == str(qpath.parent)
            and q["version"] == VERSION and q["status"] == "qualified" and q["engine"] == qengine.identity
            and _profile(q) == profile
            and q["protocol_hash"] == qprotocol["record_hash"] and q["model_api_calls"] == 0
            and {c["name"] for c in q["controls"]} == roster
            and len(q["controls"]) == len(roster) and all(c["qualified"] for c in q["controls"])
            and not qengine.unsafe_cleanup() and _execution(qengine)["open_workbook_intents"] == 0,
            f"Independent complete {profile} qualification required")
    receipts = {r["record_hash"]: r for r in qengine.receipts()}
    control_inputs = {c["name"]: c["path"] for c in qprotocol["controls"]}
    require(len(control_inputs) == len(qprotocol["controls"]) == len(roster)
            and set(control_inputs) == roster, "Qualified control roster changed")
    for c in q["controls"]:
        require(c["receipt_hash"] in receipts and c["cleanup_confirmed"] is True
                and receipts[c["receipt_hash"]]["input_sha256"] == v5.sha(control_inputs[c["name"]])
                and read_json(qpath.parent / "controls" / (c["name"] + ".json"), sealed=True) == c,
                "Qualified control binding differs")
    with _readonly_parent_lock(source):
        plan = read_json(source / "plan.json", sealed=True)
        require(plan["version"] == "continual-eval-v2" and plan["config"]["partition"] == "development"
                and plan["config"]["methods"] == ["no_skill"] and plan["config"]["histories"] == ["h0"]
                and plan["repeats"] == 2, "Original complete No-Skill development plan required")
        tasks = panel_tasks(plan, "spreadsheetbench")
        require(len(tasks) == 80 and len({t["family_id"] for t in tasks}) == 80
                and {r["task_hash"] for r in plan["tasks"]} == {digest(t) for t in tasks}, "Full 80-task roster required")
        cp = read_json(source / "checkpoints/no_skill/h0/s0.json", sealed=True)
        require(cp["skill_text"] == "" and cp["plan_hash"] == plan["record_hash"], "Original empty checkpoint required")
        inventory = _files(list(source.rglob("*.json")) + list(qpath.parent.rglob("*.json")))
        inventory.update(qprotocol["inventory"])
        inventory.update(_files([plan["panels"]["spreadsheetbench"]["path"]]))
        for name, expected in plan["source_identity"].items():
            path = safe_path(source_root / "skillopt" / name)
            require(path.is_relative_to(source_root / "skillopt") and v5.sha(path) == expected, "Original source changed")
            inventory[str(path)] = expected
        slots = []
        for task in tasks:
            require(task["partition"] == "development" and task["public"]["answer_position"] == task["private"]["answer_position"], "Task boundary differs")
            refs = task["private"]["test_files"]
            require(refs and all(task["private"]["asset_sha256"].get(p) == v5.sha(p) for p in refs), "Reference changed")
            inventory.update(_files(refs))
            for repeat in range(2):
                request = {"plan_hash": plan["record_hash"], "checkpoint_hash": cp["record_hash"],
                    "benchmark": "spreadsheetbench", "task_hash": digest(task), "repeat": repeat}
                key = digest(request)
                path = source / "predictions" / key / "prediction.json"
                prediction = read_json(path, sealed=True)
                old = read_json(source / "host_only/scores" / (key + ".json"), sealed=True)
                require(prediction["request"] == request and old["prediction_hash"] == prediction["record_hash"]
                        and all(old[k] == v for k, v in {"plan_hash": plan["record_hash"], "checkpoint_hash": cp["record_hash"], "task_id": task["task_id"], "repeat": repeat}.items())
                        and read_json(path.parent / "intent.json", sealed=True) == seal(request), "Original position binding differs")
                slots.append({"id": key, "prediction_path": str(path), "prediction_hash": prediction["record_hash"],
                    "old_score_hash": old["record_hash"], "old_status": old["status"], "task_id": task["task_id"],
                    "repeat": repeat, "references": refs, "answer_position": task["private"]["answer_position"]})
        require({p.name for p in (source / "predictions").iterdir()} == {s["id"] for s in slots}
                and {p.stem for p in (source / "host_only/scores").glob("*.json")} == {s["id"] for s in slots},
                "Extra or missing original positions")
        protocol = seal({"version": VERSION, "kind": "replay", "root": str(root), "source": str(source),
            "script_sha256": v5.sha(__file__), "engine": qengine.identity, "engine_profile": profile,
            "native_lock": str(lock),
            "qualification_hash": q["record_hash"], "slots": slots, "inventory": inventory,
            "model_api_calls": 0, "historical_scores_replaced": False, "feedback_allowed": False})
        root.mkdir(parents=True, mode=0o700)
        write_json(root / "protocol.json", protocol)
    return {"status": "prepared", "positions": len(slots), "protocol_hash": protocol["record_hash"]}


def report(output):
    root = safe_path(output)
    protocol = read_json(root / "protocol.json", sealed=True)
    require(protocol["kind"] == "replay" and protocol["root"] == str(root), "Wrong replay directory")
    engine = _engine(protocol, root)
    _published(root, protocol)
    rows = []
    for slot in protocol["slots"]:
        path = root / "positions" / (slot["id"] + ".json")
        if path.exists():
            row = read_json(path, sealed=True)
            require(row["protocol_hash"] == protocol["record_hash"] and row["slot_id"] == slot["id"], "Position result binding changed")
            rows.append(row)
    costs = _execution(engine)
    return seal({"version": VERSION, "protocol_hash": protocol["record_hash"], "engine_profile": _profile(protocol),
        "status": "complete" if len(rows) == len(protocol["slots"]) and not costs["open_workbook_intents"]
                  and not costs["cleanup_unconfirmed"] else "pending",
        "positions": len(protocol["slots"]), "completed": len(rows),
        "old_counts": dict(Counter(s["old_status"] for s in protocol["slots"])),
        "counts": dict(Counter(r["status"] for r in rows)),
        "transitions": dict(Counter(r["old_status"] + "->" + r["status"] for r in rows)),
        "unknown_reasons": dict(Counter(c["reason"] for r in rows for c in r["cases"] if c["status"] == "unknown")),
        "result_hashes": {r["slot_id"]: r["record_hash"] for r in rows},
        "execution": costs, "model_api_calls": 0, "historical_scores_replaced": False,
        "feedback_allowed": False, "is_method_effect": False, "excel_equivalence_proven": False})


def run(output):
    root = safe_path(output)
    with output_lock(root):
        protocol = read_json(root / "protocol.json", sealed=True)
        require(protocol["kind"] == "replay" and protocol["root"] == str(root), "Wrong replay directory")
        engine = _engine(protocol, root)
        require(not engine.unsafe_cleanup(), "Previous workbook cleanup unconfirmed")
        require(_execution(engine)["open_workbook_intents"] == 0, "Open workbook intent; no reexecution")
        report(root)
        with _readonly_parent_lock(Path(protocol["source"])):
            for slot in protocol["slots"]:
                path = root / "positions" / (slot["id"] + ".json")
                if path.exists():
                    continue
                if (root / "PAUSE").exists():
                    break
                saved = read_json(slot["prediction_path"], sealed=True)
                require(saved["record_hash"] == slot["prediction_hash"], "Frozen prediction changed")
                prediction = saved["prediction"]
                cases = (prediction.get("output") or {}).get("cases", [])
                results = []
                if prediction["status"] != "available" or not cases:
                    results = [v5._undelivered(prediction["reason"])]
                elif len(cases) != len(slot["references"]):
                    results = [{"status": "unknown", "reason": "original_case_count_mismatch"}]
                else:
                    with tempfile.TemporaryDirectory(prefix="sheet-v6-replay-") as temp:
                        for index, (case, reference) in enumerate(zip(cases, slot["references"])):
                            if case["status"] != "available":
                                results.append(v5._undelivered(case["reason"]))
                                continue
                            raw = base64.b64decode(case["output_base64"], validate=True)
                            require(len(raw) <= v5.MAX_BYTES, "Workbook size limit")
                            candidate = Path(temp).resolve() / f"{index}.xlsx"
                            candidate.write_bytes(raw)
                            results.append(v5.evaluate_pair(candidate, reference, slot["answer_position"], engine, root / "recalculations"))
                            if engine.unsafe_cleanup():
                                break
                counts = Counter(c["status"] for c in results)
                status = "unknown" if engine.unsafe_cleanup() else "fail" if counts["fail"] else "unknown" if counts["unknown"] else "pass"
                write_json(path, seal({"protocol_hash": protocol["record_hash"], "slot_id": slot["id"],
                    "old_status": slot["old_status"], "status": status, "cases": results, "model_api_calls": 0}))
                if engine.unsafe_cleanup():
                    break
        result = report(root)
        write_json(root / "reports" / (result["record_hash"] + ".json"), result)
        return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("qualify", "prepare", "run", "report"))
    parser.add_argument("--output", required=True)
    parser.add_argument("--image")
    parser.add_argument("--legacy-qualification")
    parser.add_argument("--native-lock")
    parser.add_argument("--source")
    parser.add_argument("--source-root")
    parser.add_argument("--qualification")
    parser.add_argument("--engine-version", choices=("v6", "v7", "v8"))
    parser.add_argument("--comparison-fixture-qualification")
    parser.add_argument("--prior-profile-qualification")
    args = parser.parse_args(argv)
    if args.command == "qualify":
        result = qualify(args.output, args.image, args.legacy_qualification, args.native_lock, args.engine_version,
                         args.comparison_fixture_qualification, args.prior_profile_qualification)
    elif args.command == "prepare":
        require(args.engine_version is None and args.comparison_fixture_qualification is None and args.prior_profile_qualification is None,
                "Replay inherits the frozen qualification profile")
        require(all((args.source, args.source_root, args.qualification, args.native_lock)), "Prepare arguments missing")
        result = prepare(args.source, args.source_root, args.output, args.qualification, args.native_lock)
    else:
        require(args.engine_version is None and args.comparison_fixture_qualification is None and args.prior_profile_qualification is None,
                "Resume uses the frozen engine profile")
        result = globals()[args.command](args.output)
    print(json.dumps(result, indent=2))
    return 0 if result["status"] in {"complete", "qualified", "prepared"} else 3


if __name__ == "__main__":
    raise SystemExit(main())
