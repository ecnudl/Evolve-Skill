"""Opt-in original-General formula-result read view; not default scoring.

Calc can autoformat a recalculated General formula as a date. A date-formatted
numeric cache is then read by openpyxl as datetime, unlike the original cell's
General read semantics. This view changes only the exported style attribute,
never a formula, cache, or source workbook. It does not infer that a computed
number is correct, and the literal UNO proof does not establish that claim.

This deliberately expands v6's literal-only boundary. It requires a separate
qualification/profile before use; v6 authorization is not inherited.
"""
from __future__ import annotations

import base64
import hashlib
import io
import math
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from xml.parsers import expat

import openpyxl
from openpyxl.styles.numbers import is_date_format

from skillopt.coevolution_v5.core import seal

from . import sheet_recalc as v5
from . import sheet_recalc_v6 as v6
from . import spreadsheet_compat as compat
from .core import require, safe_path

VERSION = "spreadsheet-independent-lo-recalculation-v7"
VIEW_VERSION = "original-general-formula-numeric-read-view-v1"


def _original_general(cell):
    """Unknown reserved formats must not inherit openpyxl's General fallback."""
    identity = cell._style.numFmtId if cell._style is not None else 0
    return identity == 0 or identity >= 164 and cell.number_format == "General"


def _cells(archive):
    for sheet, member in v5._sheet_parts(archive).items():
        seen = set()
        for cell in v5._xml(archive, member).iter(f"{{{compat.MAIN_NS}}}c"):
            coordinate = cell.get("r")
            require(coordinate and coordinate not in seen, "Duplicate or missing cell")
            seen.add(coordinate)
            yield (sheet, coordinate), cell


def _formula_text(cell):
    formulas = cell.findall(f"{{{compat.MAIN_NS}}}f")
    if (len(formulas) == 1 and formulas[0].get("t", "normal") == "normal"
            and formulas[0].text):
        return "=" + formulas[0].text
    return None


def _formula_numeric_caches(archive):
    result = {}
    for identity, cell in _cells(archive):
        formula = _formula_text(cell)
        values = cell.findall(f"{{{compat.MAIN_NS}}}v")
        if (cell.get("t", "n") == "n" and formula and len(values) == 1
                and values[0].text and v5.NUMERIC_TEXT.fullmatch(values[0].text)):
            raw = values[0].text
            value = float(raw)
            if math.isfinite(value):
                result[identity] = (value.hex(), hashlib.sha256(raw.encode()).hexdigest(), formula)
    return result


def _formula_view(original, exported, literal_proof):
    original, exported = safe_path(original), safe_path(exported)
    v6._preflight(original)
    v6._preflight(exported)
    raw = exported.read_bytes()
    proof = {**literal_proof, "version": VIEW_VERSION,
             "literal_view_version": v6.VIEW_VERSION,
             "literal_view_sha256": hashlib.sha256(raw).hexdigest(),
             "formula_adapted_cells": [], "formula_result_correctness_proven": False}
    books = []
    try:
        before = openpyxl.load_workbook(original, data_only=False, keep_links=False)
        books.append(before)
        after = openpyxl.load_workbook(exported, data_only=False, keep_links=False)
        books.append(after)
        if (before.sheetnames != after.sheetnames or before.epoch != after.epoch
                or v6._observes_format(before) or v6._observes_format(after)):
            proof["view_sha256"] = hashlib.sha256(raw).hexdigest()
            return raw, proof
        derived = v5._array_outputs(before) | v5._array_outputs(after)
        eligible = []
        for sheet in before:
            for cell in sheet._cells.values():
                changed = after[sheet.title][cell.coordinate]
                if (cell.data_type == "f" and isinstance(cell.value, str)
                        and _original_general(cell) and is_date_format(changed.number_format)
                        and (sheet.title, cell.coordinate) not in derived):
                    require(changed.data_type == "f" and changed.value == cell.value,
                            "General formula changed before numeric read view")
                    eligible.append((sheet.title, cell.coordinate))
        if not eligible:
            proof["view_sha256"] = hashlib.sha256(raw).hexdigest()
            return raw, proof
        with compat._safe_archive(original) as left, compat._safe_archive(exported) as right:
            v6._general_style(left)  # Reject reserved or duplicate custom definitions.
            general = v6._general_style(right)
            source_formulas = {identity: _formula_text(cell) for identity, cell in _cells(left)}
            caches, parts, changes = _formula_numeric_caches(right), v5._sheet_parts(right), {}
            for sheet, coordinate in eligible:
                require((sheet, coordinate) in caches, "General formula numeric cache unverified")
                binary64, cache_hash, raw_formula = caches[(sheet, coordinate)]
                require(source_formulas.get((sheet, coordinate)) == raw_formula
                        == before[sheet][coordinate].value, "General formula XML binding unverified")
                changes.setdefault(parts[sheet], set()).add(coordinate)
                proof["formula_adapted_cells"].append({
                    "sheet": sheet, "coordinate": coordinate,
                    "source_formula_sha256": hashlib.sha256(before[sheet][coordinate].value.encode()).hexdigest(),
                    "source_number_format": "General", "cache_binary64": binary64,
                    "cache_text_sha256": cache_hash, "comparison_style_index": general})
            target = io.BytesIO()
            with zipfile.ZipFile(target, "w") as dest:
                for info in right.infolist():
                    content = right.read(info.filename)
                    if info.filename in changes:
                        content = v6._replace_styles(content, changes[info.filename], general)
                    dest.writestr(info, content)
            view = target.getvalue()
            require(len(view) <= v5.MAX_BYTES, "Comparison view size limit")
            proof["view_sha256"] = hashlib.sha256(view).hexdigest()
            return view, proof
    finally:
        for book in books:
            book.close()


def comparison_view(original, exported):
    """Offline read adaptation; alone this does not prove engine execution."""
    raw, literal_proof = v6.comparison_view(original, exported)
    with tempfile.TemporaryDirectory(prefix="sheet-v7-literal-view-") as directory:
        intermediate = Path(directory).resolve() / "literal-view.xlsx"
        intermediate.write_bytes(raw)
        return _formula_view(original, intermediate, literal_proof)


class Recalculator(v6.Recalculator):
    """Reuse v5 isolation and v6 literal checks, with a separate identity."""

    @property
    def identity(self):
        identity = super().identity
        identity["version"] = VERSION
        identity["comparison_view"] = VIEW_VERSION
        identity["sources"] = {**identity["sources"], str(Path(__file__)): v5.sha(Path(__file__))}
        return identity

    def run(self, path):
        receipt = super().run(path)
        if receipt["status"] != "available":
            return receipt
        try:
            raw = base64.b64decode(receipt["output_base64"], validate=True)
            require(hashlib.sha256(raw).hexdigest() == receipt["output_sha256"]
                    and receipt["comparison_view"]["version"] == v6.VIEW_VERSION,
                    "Literal view binding changed")
            with tempfile.TemporaryDirectory(prefix="sheet-v7-formula-view-") as directory:
                intermediate = Path(directory).resolve() / "literal-view.xlsx"
                intermediate.write_bytes(raw)
                view, proof = _formula_view(path, intermediate, receipt["comparison_view"])
            return seal({**{key: value for key, value in receipt.items() if key != "record_hash"},
                         "comparison_view": proof, "output_base64": base64.b64encode(view).decode(),
                         "output_sha256": hashlib.sha256(view).hexdigest()})
        except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError,
                expat.ExpatError, zipfile.BadZipFile):
            return seal({**{key: value for key, value in receipt.items() if key != "record_hash"},
                         "status": "unknown", "reason": "general_formula_read_view_unverified"})
