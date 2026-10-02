"""Frozen five-benchmark inference and native scoring, with no judge feedback.

This is deliberately a one-shot profile for QA, rules, code and spreadsheet;
ALFWorld receives only public observations/actions. Native scores are produced
after inference. Missing execution dependencies are unknown, never fake passes.
"""
from __future__ import annotations

import base64
import hashlib
import importlib.util
import io
import json
import multiprocessing
import os
import re
import shutil
import subprocess
import tempfile
import time
import uuid
import zipfile
from datetime import date, datetime, timedelta
from datetime import time as datetime_time
from functools import lru_cache
from pathlib import Path

from skillopt.skill_validation.sandbox import _IMAGE, _bounded_command, _strict_json

VERSION = "continual-eval-native-backends-v1"
PROTOCOL = "continual-eval-native-v1"
BENCHMARKS = {"searchqa", "korbench", "bigcodebench", "spreadsheetbench", "alfworld"}
MAX_RESPONSE = 262144
MAX_WORKBOOK = 8 * 1024 * 1024
ALFWORLD_TRACE_VERSION = "alfworld-public-transitions-v1"
ALFWORLD_TRACE_OBSERVATION_CHARS = 12000
SHEET_PREVIEW_VERSION = "public-formula-preview-v1"
SHEET_PREVIEW_FIELD_CHARS = 200
CODE_EXTRACTION_VERSION = "explicit-python-fences-v1"
_PUBLIC = {
    "searchqa": ("question", "context"), "korbench": ("rule", "question"),
    "bigcodebench": ("prompt", "entry_point"),
    "spreadsheetbench": ("instruction", "input_files", "answer_position"), "alfworld": ("game_file",),
}


def _unknown(reason, **extra):
    return {"status": "unknown", "score": None, "metrics": {}, "reason": reason, **extra}


def _prediction_unknown(reason, **extra):
    return {"status": "unknown", "output": None, "costs": {}, "reason": reason, **extra}


def _public_view(benchmark, public):
    if benchmark not in BENCHMARKS:
        raise ValueError("Unknown benchmark")
    # Never serialize the whole imported task, even if its caller accidentally
    # supplies hidden labels or reference implementations alongside public data.
    view = {key: public[key] for key in _PUBLIC[benchmark]}
    for key, value in view.items():
        if key in {"input_files", "context"} and isinstance(value, list):
            if not all(type(item) is str for item in value):
                raise ValueError("Public sequences must contain text only")
        elif type(value) is not str:
            raise ValueError("Public fields must be text")
    if benchmark == "spreadsheetbench" and not isinstance(view["input_files"], list):
        raise ValueError("Workbook inputs must be a list")
    return view


def _response(call, system, user):
    try:
        receipt = call(system, user)
    except Exception as exc:
        return None, {"calls": 1}, "model_call_exception:" + type(exc).__name__
    if not isinstance(receipt, dict):
        return None, {"calls": 1}, "invalid_model_receipt"
    costs = {"calls": 1, "usage": receipt.get("usage", {})}
    response = receipt.get("response")
    if receipt.get("request", {}).get("service", {}).get("delivery_retry_policy") == "closed_network_error_v1":
        if receipt.get("finish_reason") in {"sensitive", "content_filter"}:
            return None, costs, "model_response_filtered"
        if receipt.get("finish_reason") == "network_error":
            return None, costs, "model_delivery_network_error"
    if receipt.get("finish_reason") in {"length", "max_tokens"}:
        return None, costs, "model_response_truncated"
    if not receipt.get("ok") or not isinstance(response, str):
        return None, costs, "model_call_unavailable"
    if not response.strip() or len(response.encode()) > MAX_RESPONSE:
        return None, costs, "empty_or_oversized_model_response"
    return response, costs, "model_response_available"


def _code(response):
    match = re.search(r"```(?:python)?\s*\n(.*?)```", response, re.S | re.I)
    return match.group(1).strip() if match else response.strip()


def _runtime_limits(runtime):
    timeout = runtime.get("timeout_seconds", 300)
    memory = runtime.get("memory_mb", 4096)
    cpus = runtime.get("cpus", 1)
    if not isinstance(timeout, (int, float)) or not 5 <= timeout <= 600:
        raise ValueError("Native timeout must be 5..600 seconds")
    if type(memory) is not int or not 256 <= memory <= 32768:
        raise ValueError("Native memory must be 256..32768 MiB")
    if not isinstance(cpus, (int, float)) or not 0 < cpus <= 8:
        raise ValueError("Native CPU allocation must be in (0,8]")
    return timeout, memory, cpus


def _image_ready(runtime):
    image = runtime.get("image", "")
    if not isinstance(image, str) or not _IMAGE.fullmatch(image):
        return {"status": "unsupported", "reason": "digest_pinned_image_required"}
    docker = shutil.which("docker")
    if not docker:
        return {"status": "unsupported", "reason": "docker_unavailable"}
    result = _bounded_command([docker, "image", "inspect", image], 10)
    if result.code != 0 or result.timed_out or result.overflow or result.unavailable:
        return {"status": "unsupported", "reason": "pinned_image_or_daemon_unavailable"}
    try:
        info = _strict_json(result.stdout)[0]
        if (info["Os"] != "linux" or info["Config"].get("Volumes")
                or image not in [info["Id"], *(info.get("RepoDigests") or [])]):
            raise ValueError("Image identity/volume mismatch")
    except (ValueError, KeyError, TypeError, IndexError, AttributeError):
        return {"status": "unsupported", "reason": "image_inspection_mismatch"}
    return {"status": "ready", "reason": "pinned_linux_image_available", "image_id": info["Id"],
            "architecture": info.get("Architecture")}


def _native(request, runtime):
    ready = _image_ready(runtime)
    if ready["status"] != "ready":
        return _unknown(ready["reason"], execution_costs={"container_calls": 0, "wall_seconds": 0.0})
    timeout, memory, cpus = _runtime_limits(runtime)
    docker, name = shutil.which("docker"), "continual-eval-" + uuid.uuid4().hex
    payload = json.dumps(request, allow_nan=False).encode()
    if len(payload) > 32 * 1024 * 1024:
        return _unknown("native_request_size_limit")
    with tempfile.TemporaryDirectory(prefix="continual-eval-native-") as directory:
        root = Path(directory)
        root.chmod(0o755)
        (root / "request.json").write_bytes(payload)
        (root / "request.json").chmod(0o444)
        (root / "worker.py").write_bytes(Path(__file__).with_name("native_worker.py").read_bytes())
        (root / "worker.py").chmod(0o444)
        command = [docker, "run", "--pull=never", "--name", name, "--network=none", "--read-only",
                   "--user=65534:65534", "--cap-drop=ALL", "--security-opt=no-new-privileges=true",
                   "--pids-limit=128", f"--memory={memory}m", f"--memory-swap={memory}m", f"--cpus={cpus}",
                   "--ipc=private", "--shm-size=64m", "--no-healthcheck", "--log-driver=none",
                   "--ulimit=nofile=256:256", "--ulimit=core=0:0", "--env", f"CONTINUAL_EVAL_CONTAINER={PROTOCOL}",
                   "--env", "MPLCONFIGDIR=/tmp/matplotlib", "--env", "OMP_NUM_THREADS=1",
                   "--env", "OPENBLAS_NUM_THREADS=1",
                   "--tmpfs=/tmp:rw,noexec,nosuid,nodev,size=256m,mode=1777", "--workdir=/tmp",
                   "--mount", f"type=bind,source={root},target=/input,readonly",
                   "--entrypoint=python", runtime["image"], "-I", "-B", "/input/worker.py"]
        started = time.monotonic()
        result = cleanup = None
        execution_exception = None
        try:
            result = _bounded_command(command, timeout, limit=16 * 1024 * 1024)
        except (OSError, subprocess.TimeoutExpired) as exc:
            # The bounded helper can itself fail while reaping a stalled Docker
            # client. This is infrastructure uncertainty, not a failed test.
            execution_exception = type(exc).__name__
        finally:
            try:
                cleanup = _bounded_command([docker, "rm", "--force", name], 10)
            except (OSError, subprocess.TimeoutExpired):
                pass  # No confirmed cleanup; never expose a successful score.
        cleaned = (cleanup is not None and not cleanup.timed_out and not cleanup.overflow and not cleanup.unavailable
                   and (cleanup.code == 0 or b"no such container" in cleanup.stderr.lower()))
        execution_costs = {"container_calls": 1, "wall_seconds": round(time.monotonic() - started, 6),
                           "includes_cleanup": True, "image_id": ready["image_id"]}
    if not cleaned:
        return _unknown("container_cleanup_unconfirmed", execution_costs=execution_costs)
    if execution_exception is not None:
        return _unknown("native_timeout" if execution_exception == "TimeoutExpired" else "native_execution_unavailable",
                        execution_costs=execution_costs)
    if result.timed_out or result.overflow or result.code != 0 or result.unavailable:
        return _unknown("native_timeout" if result.timed_out else "native_execution_unavailable",
                        execution_costs=execution_costs)
    try:
        envelope = _strict_json(result.stdout)
        if envelope["protocol"] != PROTOCOL or not isinstance(envelope["result"], dict):
            raise ValueError("Invalid native envelope")
        return {**envelope["result"], "runtime_image_id": ready["image_id"], "cleanup_confirmed": True,
                "runtime_architecture": ready.get("architecture"), "execution_costs": execution_costs}
    except (ValueError, KeyError, TypeError, RecursionError):
        return _unknown("invalid_native_receipt", execution_costs=execution_costs)


def _kor_sources(runtime):
    root = Path(runtime.get("kor_repo", ""))
    sources = {}
    for key, relative in (("eval", "eval/eval_utils.py"), ("common", "utils/common.py"),
                          ("config", "config/config_wrapper.py")):
        raw = (root / relative).read_bytes()
        if len(raw) > 1024 * 1024:
            raise ValueError("KOR scorer source exceeds size budget")
        expected = runtime.get("kor_" + key + "_sha256", "")
        if not re.fullmatch(r"[0-9a-f]{64}", expected) or hashlib.sha256(raw).hexdigest() != expected:
            raise ValueError("KOR scorer source is not hash-pinned")
        sources[key + "_source"] = raw.decode()
    return sources


def readiness(benchmark, runtime=None):
    runtime = runtime or {}
    if benchmark not in BENCHMARKS:
        return {"status": "unsupported", "reason": "unknown_benchmark"}
    extraction = runtime.get("code_extraction_version")
    if extraction is not None and (extraction != CODE_EXTRACTION_VERSION
                                   or benchmark not in {"bigcodebench", "spreadsheetbench"}):
        return {"status": "unsupported", "reason": "invalid_code_extraction_version"}
    if benchmark == "searchqa":
        return {"status": "ready", "reason": "repository_native_em_f1"}
    if benchmark == "alfworld":
        missing = [name for name in ("alfworld", "gymnasium", "yaml", "numpy")
                   if importlib.util.find_spec(name) is None]
        root = runtime.get("alfworld_data", os.environ.get("ALFWORLD_DATA", ""))
        if missing or not root or not Path(root, "logic/alfred.pddl").is_file():
            return {"status": "unsupported", "reason": "alfworld_packages_or_data_unavailable", "missing": missing}
        return {"status": "ready", "reason": "alfworld_dependencies_present_episode_not_yet_started"}
    if benchmark == "spreadsheetbench":
        preview_version = runtime.get("spreadsheet_preview_version")
        if preview_version not in (None, SHEET_PREVIEW_VERSION):
            return {"status": "unsupported", "reason": "invalid_spreadsheet_preview_version"}
        if preview_version == SHEET_PREVIEW_VERSION:
            try:
                _sheet_formula_classes()
            except ValueError:
                return {"status": "unsupported", "reason": "public_formula_preview_classes_unavailable"}
    sources = {}
    try:
        _runtime_limits(runtime)
        if benchmark == "korbench":
            sources = _kor_sources(runtime)
        if benchmark == "spreadsheetbench" and importlib.util.find_spec("openpyxl") is None:
            return {"status": "unsupported", "reason": "host_openpyxl_unavailable"}
        if benchmark == "spreadsheetbench":
            if runtime.get("spreadsheet_scorer") in {"qualified_lo_recalc_v5_v1", "qualified_lo_recalc_v7_v1"}:
                if runtime["spreadsheet_scorer"] == "qualified_lo_recalc_v7_v1":
                    from .sheet_numeric_adapter import readiness as recalc_readiness
                else:
                    from .sheet_recalc_adapter import readiness as recalc_readiness

                state = recalc_readiness(runtime)
                if state["status"] != "ready":
                    return state
            else:
                _sheet_scorer(runtime.get("spreadsheet_scorer"))
    except (ValueError, OSError):
        return {"status": "unsupported", "reason": "invalid_runtime_or_missing_pinned_scorer"}
    result = _native({"operation": "probe", "benchmark": benchmark, **sources}, runtime)
    return {**result, "status": "ready" if result.get("status") == "ready" else "unsupported"}


def _xlsx_bytes(path):
    path = Path(path)
    if not path.is_file() or path.stat().st_size > MAX_WORKBOOK:
        raise ValueError("Missing or oversized workbook")
    raw = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 128 * 1024 * 1024:
            raise ValueError("Workbook expanded size limit")
    return raw


def _sheet_formula_classes():
    try:
        from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
    except (ImportError, AttributeError) as exc:
        raise ValueError("Public formula preview classes unavailable") from exc
    return ArrayFormula, DataTableFormula


def _sheet_public_value(value):
    """Stable, bounded public display; never repr arbitrary workbook objects."""
    ArrayFormula, DataTableFormula = _sheet_formula_classes()

    if value is None:
        return None
    if any(type(value) is kind for kind in (str, bool, int, float, date, datetime, datetime_time, timedelta)):
        return str(value)[:SHEET_PREVIEW_FIELD_CHARS]
    if type(value) is ArrayFormula:
        kind, text_fields, boolean_fields = "array_formula", ("text", "ref"), ()
    elif type(value) is DataTableFormula:
        kind, text_fields, boolean_fields = "data_table_formula", ("ref", "r1", "r2"), (
            "ca", "dt2D", "dtr", "del1", "del2")
    else:
        return {"kind": "unsupported_cell", "reason": "unsupported_public_value_type"}
    result = {"kind": kind}
    truncated, unsupported, missing = [], [], []
    for key in (*text_fields, *boolean_fields):
        field = vars(value).get(key)
        if field is None:
            result[key] = None
            missing.append(key)
        elif key in text_fields and type(field) is str:
            result[key] = field[:SHEET_PREVIEW_FIELD_CHARS]
            if len(field) > SHEET_PREVIEW_FIELD_CHARS:
                truncated.append(key)
        elif key in boolean_fields and type(field) is bool:
            result[key] = field
        elif key in boolean_fields and type(field) is str and field in ("0", "1", "false", "true"):
            # openpyxl 3.1.x forwards these XML attributes without conversion.
            result[key] = field in ("1", "true")
        else:
            result[key] = None
            unsupported.append(key)
    result.update(truncated_fields=truncated, unsupported_fields=unsupported, missing_fields=missing)
    return result


def _sheet_preview(path, *, version=None):
    import openpyxl
    if version not in (None, SHEET_PREVIEW_VERSION):
        raise ValueError("Unsupported public workbook preview version")
    if version == SHEET_PREVIEW_VERSION:
        _sheet_formula_classes()
    wb = openpyxl.load_workbook(io.BytesIO(_xlsx_bytes(path)), read_only=True, data_only=False)
    try:
        return [{"sheet": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                 "preview": [[(_sheet_public_value(value) if version == SHEET_PREVIEW_VERSION
                               else str(value)[:200] if value is not None else None) for value in row]
                             for row in sheet.iter_rows(min_row=1, max_row=5, max_col=20, values_only=True)]}
                for sheet in wb.worksheets[:20]]
    finally:
        wb.close()


def solve(benchmark, public, skill_text, call, *, runtime=None):
    """Only public contract/observations enter `call(system, user)`. No grading."""
    runtime = runtime or {}
    extraction = runtime.get("code_extraction_version")
    if extraction is not None and (extraction != CODE_EXTRACTION_VERSION
                                   or benchmark not in {"bigcodebench", "spreadsheetbench"}):
        return _prediction_unknown("invalid_code_extraction_version")
    try:
        view = _public_view(benchmark, public)
    except (KeyError, ValueError, TypeError):
        return _prediction_unknown("invalid_public_contract")
    if benchmark == "alfworld":
        return _solve_alfworld(view, skill_text, call, runtime)
    system = "Solve the supplied task using only the supplied public information."
    if skill_text.strip():
        system += "\n\nFrozen skill guidance (use only where applicable):\n" + skill_text
    if benchmark == "searchqa":
        system += "\nReturn the concise final answer inside <answer>...</answer>."
    elif benchmark == "korbench":
        system += "\nApply the given rules; enclose the final answer in [[...]]."
    else:
        system += "\nReturn one complete Python program in a ```python code block."
    costs = {}
    try:
        if benchmark == "spreadsheetbench":
            preview_version = runtime.get("spreadsheet_preview_version")
            if preview_version not in (None, SHEET_PREVIEW_VERSION):
                return _prediction_unknown("invalid_spreadsheet_preview_version")
            files = view.get("input_files", [])
            if not files:
                return _prediction_unknown("missing_public_workbook")
            # The same script is applied independently to every official case,
            # as in the existing codegen benchmark profile.
            view = {"instruction": view["instruction"],
                    "workbook_preview": _sheet_preview(files[0], version=preview_version),
                    "input_path": "input.xlsx", "required_output_path": "output.xlsx",
                    "required_answer_position": view["answer_position"]}
        response, costs, reason = _response(call, system, json.dumps(view, ensure_ascii=False))
        if response is None:
            return _prediction_unknown(reason, costs=costs)
        delivery = {}
        if extraction == CODE_EXTRACTION_VERSION:
            from .code_delivery import extract_python

            extracted = extract_python(response)
            delivery = {"code_delivery": {key: value for key, value in extracted.items() if key != "code"}}
            if extracted["status"] != "available":
                return _prediction_unknown("code_delivery:" + extracted["reason"], costs=costs, **delivery)
            output = extracted["code"]
        else:
            # Historical profiles retain their exact extraction behavior.
            output = response if benchmark in {"searchqa", "korbench"} else _code(response)
        if benchmark == "spreadsheetbench":
            cases = [_native({"operation": "spreadsheet_generate", "code": output,
                              "input_base64": base64.b64encode(_xlsx_bytes(path)).decode()}, runtime)
                     for path in public["input_files"]]
            return {"status": "available", "output": {"code": output, "cases": cases},
                    "costs": costs, "reason": "generated_once_executed_per_public_case", **delivery}
        return {"status": "available", "output": output, "costs": costs, "reason": reason, **delivery}
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        return _prediction_unknown("public_input_or_runtime_error:" + type(exc).__name__, costs=costs)


@lru_cache(maxsize=2)
def _sheet_scorer(profile=None):
    if profile is not None:
        from . import spreadsheet_compat

        if profile != spreadsheet_compat.VERSION:
            raise ValueError("Unsupported frozen spreadsheet scorer profile")
        return spreadsheet_compat
    # Load the pure scorer without SpreadsheetBench's __init__ importing legacy
    # agents, API clients and the old (host-executing) codegen implementation.
    path = Path(__file__).parents[1] / "envs/spreadsheetbench/evaluator.py"
    spec = importlib.util.spec_from_file_location("continual_native_sheet_scorer", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def score(benchmark, public, private, prediction, *, runtime=None):
    """Hidden audit only; never call the model or feed results into inference."""
    runtime = runtime or {}
    if prediction.get("status") != "available":
        return _unknown(prediction.get("reason", "prediction_unavailable"))
    output = prediction.get("output")
    try:
        if benchmark == "searchqa":
            from skillopt.envs.searchqa.evaluator import evaluate
            answers = private["answers"]
            if not isinstance(output, str) or not answers or not all(isinstance(a, str) for a in answers):
                return _unknown("invalid_qa_score_input")
            result = evaluate(output, answers)
            return {"status": "pass" if result["em"] else "fail", "score": result["em"],
                    "reason": "repository_searchqa_native_em", "metrics": {k: result[k] for k in ("em", "f1", "sub_em")}}
        if benchmark == "korbench":
            required = ("answer", "category", "rule_id", "upstream_index")
            if any(not isinstance(private.get(k), str) for k in required):
                return _unknown("missing_kor_native_scoring_identity")
            if private["category"] not in {"logic", "operation", "puzzle", "cipher", "counterfactual"}:
                return _unknown("unsupported_kor_category_or_mixed_mode")
            return _native({"operation": benchmark, "response": output,
                            **{k: private[k] for k in required}, **_kor_sources(runtime)}, runtime)
        if benchmark == "bigcodebench":
            timeout, memory, _ = _runtime_limits(runtime)
            return _native({"operation": benchmark, "code": output, "test": private["test"],
                            "entry_point": public["entry_point"], "memory_mb": memory,
                            "test_timeout": min(60, max(5, timeout - 15))}, runtime)
        if benchmark == "spreadsheetbench":
            if public.get("answer_position") != private.get("answer_position"):
                return _unknown("public_scoring_region_mismatch")
            if runtime.get("spreadsheet_scorer") in {"qualified_lo_recalc_v5_v1", "qualified_lo_recalc_v7_v1"}:
                if runtime["spreadsheet_scorer"] == "qualified_lo_recalc_v7_v1":
                    from .sheet_numeric_adapter import score as recalc_score
                else:
                    from .sheet_recalc_adapter import score as recalc_score

                return recalc_score(public, private, prediction, runtime=runtime)
            return _score_sheet(private, output, scorer_profile=runtime.get("spreadsheet_scorer"))
        if benchmark == "alfworld":
            if not isinstance(output, dict) or type(output.get("won")) is not bool:
                return _unknown("missing_native_environment_success")
            ok = output["won"]
            return {"status": "pass" if ok else "fail", "score": float(ok),
                    "metrics": {"success": float(ok), "steps": output["steps"]}, "reason": "alfworld_native_won"}
        return _unknown("unknown_benchmark")
    except (KeyError, ValueError, TypeError, OSError, zipfile.BadZipFile) as exc:
        return _unknown("native_scoring_input_error:" + type(exc).__name__)


def _score_sheet(private, output, *, scorer_profile=None):
    scorer = _sheet_scorer(scorer_profile)
    cases, golds = output["cases"], private["test_files"]
    if not cases or len(cases) != len(golds):
        return _unknown("spreadsheet_case_count_mismatch")
    results = []
    with tempfile.TemporaryDirectory(prefix="continual-sheet-score-") as directory:
        for index, (case, gold) in enumerate(zip(cases, golds)):
            if case["status"] != "available":
                results.append({"status": "unknown", "reason": case.get("reason", "no_workbook")})
                continue
            path = Path(directory) / f"{index}.xlsx"
            raw = base64.b64decode(case["output_base64"], validate=True)
            if len(raw) > MAX_WORKBOOK:
                return _unknown("predicted_workbook_size_limit")
            path.write_bytes(raw)
            _xlsx_bytes(path)
            _xlsx_bytes(gold)
            result = scorer.evaluate(str(path), gold, "", private["answer_position"])
            case_result = {"status": "pass" if result["status"] == "passed" else
                           "fail" if result["status"] in {"failed", "missing_output"} else "unknown",
                           "reason": result["reason"]}
            if scorer_profile is not None:
                case_result.update(evaluator_version=result["evaluator_version"], evidence=result["evidence"])
            results.append(case_result)
    passed = sum(r["status"] == "pass" for r in results)
    failed = sum(r["status"] == "fail" for r in results)
    unknown = len(results) - passed - failed
    status = "fail" if failed else "unknown" if unknown else "pass"
    return {"status": status, "score": 0.0 if failed else None if unknown else 1.0,
            "reason": "native_cell_values_hard_all_cases", "metrics": {"cases": results,
            "case_passed": passed, "case_failed": failed, "case_unknown": unknown,
            "soft_score": passed / len(results) if not unknown else None,
            "soft_score_lower_bound": passed / len(results)}}


class _AlfSession:
    """Bounded transport around the existing trusted ALFWorld worker only."""

    def __init__(self, config, game, split, seed, timeout):
        from skillopt.envs.alfworld.vendor.alfworld_envs import _worker_loop
        context = multiprocessing.get_context("spawn")
        self.commands, self.results = context.Queue(1), context.Queue(1)
        self.timeout = timeout
        self.process = context.Process(target=_worker_loop, args=(
            self.commands, self.results, config, seed, split == "train", split, game))
        self.process.start()
        try:
            self._receive()
        except Exception:
            self.close()
            raise

    def _receive(self):
        ok, result = self.results.get(timeout=self.timeout)
        if not ok:
            raise RuntimeError("ALFWorld native worker failed")
        return result

    @staticmethod
    def _info(info):
        return [{key: value[0] for key, value in info.items()}]

    def reset(self):
        self.commands.put(("reset", None), timeout=self.timeout)
        observations, info = self._receive()
        return observations, None, self._info(info)

    def step(self, actions):
        self.commands.put(("step", actions[0]), timeout=self.timeout)
        observations, scores, dones, info = self._receive()
        return observations, None, scores, dones, self._info(info)

    def close(self):
        # Give the native worker a bounded normal-exit opportunity so its own
        # TextWorld children/semaphores can be reclaimed. Always terminating it
        # leaked native semaphore resources even after the outer process exited.
        if self.process.is_alive():
            try:
                self.commands.put(("close", None), timeout=min(1, self.timeout))
            except Exception:
                pass
            self.process.join(timeout=1)
        if self.process.is_alive():
            self.process.terminate()
        self.process.join(timeout=3)
        if self.process.is_alive():
            self.process.kill()
            self.process.join(timeout=2)
        self.commands.cancel_join_thread()
        self.results.cancel_join_thread()
        self.commands.close()
        self.results.close()
        if self.process.is_alive():
            raise RuntimeError("ALFWorld worker cleanup unconfirmed")


def _new_alfworld(public, runtime):
    import yaml
    game_path = Path(public["game_file"]).resolve(strict=True)
    root = runtime.get("alfworld_data", os.environ.get("ALFWORLD_DATA", ""))
    # The library expands ALFWORLD_DATA; don't silently change a shared process
    # environment while concurrently running other episodes.
    if root != os.environ.get("ALFWORLD_DATA"):
        raise ValueError("Set ALFWORLD_DATA to the frozen runtime path before launch")
    if not game_path.is_file() or game_path.name != "game.tw-pddl":
        raise ValueError("ALFWorld requires an existing game.tw-pddl file")
    if not game_path.is_relative_to(Path(root).resolve(strict=True)):
        raise ValueError("ALFWorld game is outside the frozen data root")
    config = Path(__file__).parents[1] / "envs/alfworld/vendor/config_tw.yaml"
    split = runtime.get("alfworld_split", "eval_out_of_distribution")
    split_path_keys = {"train": "data_path", "eval_in_distribution": "eval_id_data_path",
                       "eval_out_of_distribution": "eval_ood_data_path"}
    if split not in split_path_keys:
        raise ValueError("Invalid ALFWorld split")
    max_steps = runtime.get("max_steps", 50)
    if type(max_steps) is not int or not 1 <= max_steps <= 150:
        raise ValueError("Invalid ALFWorld step budget")
    episode_config = yaml.safe_load(config.read_text())
    # The upstream constructor scans its dataset before the worker binds one
    # game. Restrict that scan to this already frozen episode; scanning thousands
    # of unrelated games per episode can exhaust the native startup deadline.
    # This is a per-worker config copy, never a shared environment/config edit.
    episode_config["dataset"][split_path_keys[split]] = str(game_path.parent)
    # Keep the native TimeLimit in agreement with the solver's frozen budget;
    # otherwise the vendor's default silently ends requested >50-step runs.
    episode_config["dagger"]["training"]["max_nb_steps_per_episode"] = max_steps
    episode_config["rl"]["training"]["max_nb_steps_per_episode"] = max_steps
    return _AlfSession(episode_config, str(game_path), split,
                       int(runtime.get("seed", 42)), min(120, float(runtime.get("timeout_seconds", 60))))


def _alf_post_observation(observations):
    # Record only the environment's public text, never stringify info/rewards
    # or invent an observation when the native result is absent/malformed.
    if not isinstance(observations, (list, tuple)) or not observations or not isinstance(observations[0], str):
        return {"post_observation": None, "post_observation_status": "missing",
                "post_observation_characters": None, "post_observation_reason": "observation_unavailable"}
    text = observations[0]
    truncated = len(text) > ALFWORLD_TRACE_OBSERVATION_CHARS
    return {"post_observation": text[:ALFWORLD_TRACE_OBSERVATION_CHARS],
            "post_observation_status": "truncated" if truncated else "complete",
            "post_observation_characters": len(text),
            "post_observation_reason": "character_limit" if truncated else "observed"}


def _solve_alfworld(public, skill_text, call, runtime):
    env, calls, usages, trace = None, 0, [], []
    trace_version = runtime.get("alfworld_trace_version")
    if trace_version not in (None, ALFWORLD_TRACE_VERSION):
        return _prediction_unknown("invalid_alfworld_trace_version")
    transitions = trace_version == ALFWORLD_TRACE_VERSION
    # Logging must not change the solver's history keys, content, or window.
    prompt_history = []

    def trace_fields():
        return {"trace": trace, **({"trace_version": trace_version} if transitions else {})}

    max_steps = runtime.get("max_steps", 50)
    if type(max_steps) is not int or not 1 <= max_steps <= 150:
        return _prediction_unknown("invalid_alfworld_step_budget")
    system = "Act in the text environment. Return a single action inside <action>...</action>."
    if skill_text.strip():
        system += "\nFrozen skill guidance (use only where applicable):\n" + skill_text
    try:
        env = _new_alfworld(public, runtime)
        observations, _, infos = env.reset()
        # The task instruction is part of the public reset observation, but
        # subsequent observations generally omit it. Keep it throughout the
        # episode, independently of the bounded recent-action history.
        initial_observation = str(observations[0])
        won = False
        for step in range(max_steps):
            # No expert plan, won/score, extra metadata, task identifiers, or
            # reference trajectory enters the model-visible view.
            view = {"initial_observation": initial_observation,
                    "observation": observations[0], "admissible_commands": infos[0].get("admissible_commands", []),
                    "recent_history": prompt_history[-5:]}
            response, costs, reason = _response(call, system, json.dumps(view, ensure_ascii=False))
            calls += 1
            usages.append(costs.get("usage", {}))
            if response is None:
                return _prediction_unknown(reason, costs={"calls": calls, "usage_by_call": usages}, **trace_fields())
            match = re.search(r"<action>(.*?)</action>", response, re.S)
            action = match.group(1).strip() if match else response.strip()
            if not action or len(action) > 1000:
                return _prediction_unknown("invalid_action_format", costs={"calls": calls, "usage_by_call": usages},
                                           **(trace_fields() if transitions else {}))
            history_item = {"observation": str(observations[0])[:12000], "action": action}
            prompt_history.append(history_item)
            trace.append({**history_item, "post_observation": None, "post_observation_status": "missing",
                          "post_observation_characters": None, "post_observation_reason": "step_result_unavailable"}
                         if transitions else history_item)
            observations, _, _, dones, infos = env.step([action])
            if transitions:
                trace[-1].update(_alf_post_observation(observations))
            if "won" not in infos[0]:
                return _prediction_unknown("native_won_signal_missing", costs={"calls": calls, "usage_by_call": usages},
                                           **(trace_fields() if transitions else {}))
            won = bool(infos[0]["won"])
            if won or dones[0]:
                break
        return {"status": "available", "output": {"won": won, "steps": step + 1},
                "reason": "fresh_native_episode_completed", "costs": {"calls": calls, "usage_by_call": usages},
                **trace_fields()}
    except Exception as exc:
        return _prediction_unknown("alfworld_environment_error:" + type(exc).__name__,
                                   costs={"calls": calls, "usage_by_call": usages}, **trace_fields())
    finally:
        if env is not None:
            try:
                env.close()
            except Exception:
                return _prediction_unknown("alfworld_cleanup_unconfirmed",
                                           costs={"calls": calls, "usage_by_call": usages}, **trace_fields())
