"""Opt-in SpreadsheetBench target/cache compatibility; historical scorer unchanged.

The official evaluator strips single quotes at the edges of sheet/range
segments, including malformed quoting found in Verified400. This adapter tries
the historical strict parser first, then that limited compatibility rule; all
resulting cells and actual reference-sheet names remain mandatory.

OOXML ``t="str"`` + a present empty ``v`` is a cached empty formula string,
unlike an absent value or an empty numeric cache. Reading this evidence neither
recalculates formulas nor establishes freshness. Original workbooks are never
written. See Microsoft CellValue/CellValues and OfficeDev working-with-formulas.
"""
from __future__ import annotations

import importlib.util
import posixpath
import xml.etree.ElementTree as ET
import zipfile
from contextlib import ExitStack
from functools import lru_cache
from pathlib import Path, PurePosixPath

import openpyxl

VERSION = "spreadsheetbench-official-quote-and-string-cache-v1"
MAX_TARGET_CELLS = 200_000
MAX_ZIP_BYTES = 8 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_XML_BYTES = 16 * 1024 * 1024
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


class CompatibilityError(ValueError):
    """Unresolved/unsafe input, not a confirmed model error."""


@lru_cache(maxsize=1)
def _legacy():
    path = Path(__file__).parents[1] / "envs/spreadsheetbench/evaluator.py"
    spec = importlib.util.spec_from_file_location("continual_sheet_legacy_compat", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def resolve_targets(answer_position: str, sheetnames: list[str]) -> dict:
    """Return validated targets and whether official edge-quote parsing was needed.

    No sheet-name guessing, first-sheet fallback for a *named* sheet, fuzzy
    matching, empty union, reversed/out-of-bounds range or arbitrary expression.
    """
    if not sheetnames or len(sheetnames) != len(set(sheetnames)):
        raise CompatibilityError("reference_sheet_identity_unavailable")
    scorer = _legacy()
    mode = "strict"
    try:
        targets = scorer._answer_targets(answer_position, sheetnames[0])
    except (ValueError, TypeError, AttributeError):
        mode = "official_edge_single_quotes"
        if not isinstance(answer_position, str) or not answer_position.strip():
            raise CompatibilityError("empty_target_spec") from None
        targets = []
        # This fallback intentionally follows only the documented official
        # split-and-edge-strip grammar, not general Excel expressions.
        for reference in answer_position.split(","):
            pieces = reference.strip().split("!")
            if len(pieces) == 1:
                sheet, region = sheetnames[0], pieces[0]
            elif len(pieces) == 2:
                sheet, region = pieces
            else:
                raise CompatibilityError("ambiguous_sheet_separator") from None
            sheet, region = sheet.strip().strip("'"), region.strip().strip("'")
            try:
                scorer._range_bounds(region)
            except (ValueError, TypeError, AttributeError) as exc:
                raise CompatibilityError("invalid_target_range") from exc
            targets.append((sheet, region))
    count = 0
    for sheet, region in targets:
        if sheet not in sheetnames:
            raise CompatibilityError("reference_sheet_not_found")
        (sc, sr), (ec, er) = scorer._range_bounds(region)
        count += (ec - sc + 1) * (er - sr + 1)
        if count > MAX_TARGET_CELLS:
            raise CompatibilityError("target_cell_budget_exceeded")
    return {"mode": mode, "targets": targets, "target_cells_including_union_overlap": count}


def _safe_archive(path):
    path = Path(path)
    if not path.is_file() or path.stat().st_size > MAX_ZIP_BYTES:
        raise CompatibilityError("workbook_missing_or_size_limit")
    archive = zipfile.ZipFile(path)
    try:
        entries = archive.infolist()
        if len(entries) > 2048 or len(entries) != len({i.filename for i in entries}):
            raise CompatibilityError("archive_member_count_or_duplicate")
        if sum(i.file_size for i in entries) > MAX_EXPANDED_BYTES:
            raise CompatibilityError("archive_expansion_limit")
        for entry in entries:
            parts = PurePosixPath(entry.filename).parts
            if (entry.flag_bits & 1 or entry.filename.startswith("/") or ".." in parts
                    or "\\" in entry.filename):
                raise CompatibilityError("unsafe_archive_member")
        return archive
    except BaseException:
        archive.close()
        raise


def _xml(archive, member):
    if archive.getinfo(member).file_size > MAX_XML_BYTES:
        raise CompatibilityError("xml_part_size_limit")
    raw = archive.read(member)
    # Restrict this compatibility reader to UTF-8/XML-compatible byte streams;
    # rejecting NUL also prevents UTF-16/32 from bypassing the declaration guard.
    if b"\x00" in raw or b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
        raise CompatibilityError("xml_document_type_or_entity_forbidden")
    return ET.fromstring(raw)


def explicit_empty_formula_strings(path) -> set[tuple[str, str]]:
    """Identify existing empty string caches using ZIP/XML, never infer formulas.

    Shared-formula followers have an ``f`` element even when its text is empty.
    An absent ``v`` or non-string cache remains unresolved. Relationship paths
    are resolved inside the ZIP; no external target is fetched.
    """
    with _safe_archive(path) as archive:
        workbook = _xml(archive, "xl/workbook.xml")
        rels = _xml(archive, "xl/_rels/workbook.xml.rels")
        by_id = {}
        for rel in rels.findall(f"{{{REL_NS}}}Relationship"):
            if rel.get("Id") in by_id:
                raise CompatibilityError("duplicate_relationship")
            by_id[rel.get("Id")] = rel
        result = set()
        sheet_names = set()
        for sheet in workbook.findall(f"{{{MAIN_NS}}}sheets/{{{MAIN_NS}}}sheet"):
            if not sheet.get("name") or sheet.get("name") in sheet_names:
                raise CompatibilityError("duplicate_or_missing_sheet_identity")
            sheet_names.add(sheet.get("name"))
            rel = by_id.get(sheet.get(f"{{{DOC_REL_NS}}}id"))
            if rel is None or rel.get("TargetMode") == "External":
                raise CompatibilityError("worksheet_relationship_unavailable")
            target = rel.get("Target", "")
            # Sheet objects such as chartsheets are not cells. Existing legacy
            # readers determine their usability if explicitly targeted.
            if not rel.get("Type", "").endswith("/worksheet"):
                continue
            if not target or "\\" in target or ":" in target:
                raise CompatibilityError("invalid_worksheet_relationship")
            member = posixpath.normpath(target.lstrip("/") if target.startswith("/")
                                       else posixpath.join("xl", target))
            if not member.startswith("xl/") or member.startswith("../"):
                raise CompatibilityError("worksheet_relationship_outside_package")
            tree = _xml(archive, member)
            seen = set()
            for cell in tree.findall(f".//{{{MAIN_NS}}}c"):
                coordinate = cell.get("r")
                if coordinate in seen:
                    raise CompatibilityError("duplicate_cell_identity")
                seen.add(coordinate)
                value = cell.find(f"{{{MAIN_NS}}}v")
                if (cell.get("t") == "str" and cell.find(f"{{{MAIN_NS}}}f") is not None
                        and value is not None and value.text in (None, "") and len(value) == 0):
                    _legacy()._parse_cell(coordinate)
                    result.add((sheet.get("name"), coordinate.upper()))
        return result


def reference_readiness(path, answer_position) -> dict:
    """Static reference evidence coverage, not solver correctness or freshness."""
    with ExitStack() as stack:
        blanks = explicit_empty_formula_strings(path)
        values = openpyxl.load_workbook(path, data_only=True, keep_links=False)
        stack.callback(values.close)
        formulas = openpyxl.load_workbook(path, data_only=False, keep_links=False)
        stack.callback(formulas.close)
        resolution = resolve_targets(answer_position, values.sheetnames)
        formulas_count = missing = explicit = cells = 0
        for sheet, region in resolution["targets"]:
            for coordinate in _legacy()._iter_cell_names(region):
                cells += 1
                if formulas[sheet][coordinate].data_type != "f":
                    continue
                formulas_count += 1
                if values[sheet][coordinate].value is None:
                    if (sheet, coordinate) in blanks:
                        explicit += 1
                    else:
                        missing += 1
        return {"version": VERSION, "status": "ready" if not missing else "unknown",
                "target_resolution": resolution["mode"], "target_cells": cells,
                "formula_cells": formulas_count, "explicit_empty_string_caches": explicit,
                "unresolved_formula_caches": missing, "cache_freshness_established": False,
                "formula_recalculation_performed": False}


def evaluate(pred_path, gold_path, instruction_type, answer_position) -> dict:
    """Opt-in scorer retaining legacy comparison, unknown and mismatch priority.

    This version must be frozen separately. It does not authorize deployment,
    overwrite historical scores, or repair cached results in either workbook.
    """
    evidence = {"reference_empty_string_caches_used": 0, "prediction_empty_string_caches_used": 0,
                "unavailable_cells": 0, "formula_recalculation_performed": False,
                "cache_freshness_established": False}

    def result(status, reason):
        return {"ok": status == "passed", "status": status, "reason": reason,
                "evaluator_version": VERSION, "instruction_type": instruction_type,
                "evidence": evidence}

    try:
        # Unknown invalid reference must not be converted into a model failure.
        gold_blanks = explicit_empty_formula_strings(gold_path)
        with ExitStack() as stack:
            gold_values = openpyxl.load_workbook(gold_path, data_only=True, keep_links=False)
            stack.callback(gold_values.close)
            gold_formulas = openpyxl.load_workbook(gold_path, data_only=False, keep_links=False)
            stack.callback(gold_formulas.close)
            resolution = resolve_targets(answer_position, gold_values.sheetnames)
            evidence["target_resolution"] = resolution["mode"]
            if not Path(pred_path).is_file():
                return result("missing_output", "missing_output: file not found")
            pred_blanks = explicit_empty_formula_strings(pred_path)
            pred_values = openpyxl.load_workbook(pred_path, data_only=True, keep_links=False)
            stack.callback(pred_values.close)
            pred_formulas = openpyxl.load_workbook(pred_path, data_only=False, keep_links=False)
            stack.callback(pred_formulas.close)
            mismatch = None
            for sheet, region in resolution["targets"]:
                if sheet not in pred_values.sheetnames:
                    return result("missing_output", "missing_output: worksheet not found")
                for coordinate in _legacy()._iter_cell_names(region):
                    available, pair = True, []
                    for label, values, formulas, blanks in (
                        ("reference", gold_values, gold_formulas, gold_blanks),
                        ("prediction", pred_values, pred_formulas, pred_blanks),
                    ):
                        value = values[sheet][coordinate].value
                        if formulas[sheet][coordinate].data_type == "f" and value is None:
                            if (sheet, coordinate) in blanks:
                                value = ""
                                evidence[f"{label}_empty_string_caches_used"] += 1
                            else:
                                available = False
                        pair.append(value)
                    if not available:
                        evidence["unavailable_cells"] += 1
                    elif not _legacy()._compare_cell_value(*pair) and mismatch is None:
                        mismatch = f"value@{sheet}!{coordinate}: mismatch"
            if mismatch is not None:
                return result("failed", mismatch)
            if evidence["unavailable_cells"]:
                return result("unknown", "unavailable: formula cache not established")
            return result("passed", "")
    except CompatibilityError as exc:
        return result("unknown", "unavailable: " + str(exc))
    except (OSError, ValueError, TypeError, KeyError, ET.ParseError, zipfile.BadZipFile) as exc:
        return result("unknown", "unavailable: workbook evidence error: " + type(exc).__name__)
