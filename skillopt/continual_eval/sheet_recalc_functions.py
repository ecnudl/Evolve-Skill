"""Opt-in native IFNA/AGGREGATE execution, never strip namespaces for Calc.

Exact qualified tokens are hidden only in a detached, never-executed screening
view so all other frozen v5 safety checks remain mandatory. Calc receives the
original OOXML bytes: removing _xlfn changes how its importer resolves these
functions. The original input hash, screening-view hash, and actual execution
receipt have distinct roles. Comparison may restore only qualified spelling;
it never edits cached results or original artifacts. The v7 numeric read view,
reference native-cache drift and full original content guard remain required.

Sources: Microsoft MS-XLSX Formulas future-function-list;
help.libreoffice.org/latest/en-GB/text/scalc/01/func_aggregate.html;
help.libreoffice.org/latest/en-US/text/scalc/01/04060104.html (IFNA).
This is a new qualification scope, not inherited v7 authorization.
"""
from __future__ import annotations

import base64
import hashlib
import html
import io
import json
import platform
import shutil
import subprocess
import tempfile
import time
import uuid
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from xml.parsers import expat

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer, TokenizerError

from skillopt.coevolution_v5.core import seal

from . import sheet_recalc as v5
from . import sheet_recalc_v6 as v6
from . import sheet_recalc_v7 as v7
from . import spreadsheet_compat as compat
from .core import require, safe_path

VERSION = "spreadsheet-independent-lo-recalculation-v8b"
ALIAS_VERSION = "native-ifna-aggregate-screen-only-v2"
# Rejected v8a's namespace-stripped inputs exported these lower-case tokens.
# This does NOT establish native function support: #NAME? still fails the guard.
# Normalize only these exact FUNC OPEN tokens, never whole formula text or
# arbitrary function/name spelling. All arguments and quoted content stay exact.
ALIASES = {"_xlfn.IFNA(": "IFNA(", "_xlfn.AGGREGATE(": "AGGREGATE(",
           "ifna(": "IFNA(", "aggregate(": "AGGREGATE("}


def canonical_formula(formula):
    """Preserve every other token, including quoted strings and references."""
    require(isinstance(formula, str) and formula.startswith("="), "Formula text unavailable")
    lexer = Tokenizer(formula)
    require(lexer.render() == formula, "Formula tokenizer roundtrip differs")
    return "=" + "".join(ALIASES.get(t.value, t.value)
                         if t.type == "FUNC" and t.subtype == "OPEN" else t.value for t in lexer.items)


def _formula_map(book):
    result = {}
    for sheet in book:
        for cell in sheet._cells.values():
            if cell.data_type == "f":
                result[(sheet.title, cell.coordinate)] = v5._formula_spec(cell.value)
    return result


def _named_expressions(book):
    return any(name not in {"_xlnm.Print_Area", "_xlnm.Print_Titles"}
               for scope in [book.defined_names, *(s.defined_names for s in book)] for name in scope)


def _patch_formula_xml(raw, edits):
    """Replace only selected f element text; all v elements stay byte-exact."""
    raw.decode("utf-8", errors="strict")
    parser = expat.ParserCreate(namespace_separator="}")
    cell_tag, formula_tag = (compat.MAIN_NS + "}" + name for name in ("c", "f"))
    stack, ranges, seen, selected = [], [], set(), set()
    coordinate, start_text = None, None

    def start(name, attrs):
        nonlocal coordinate, start_text
        if name == cell_tag:
            require(coordinate is None and attrs.get("r") not in seen, "Duplicate or nested cell")
            coordinate = attrs.get("r")
            require(coordinate is not None, "Missing cell coordinate")
            seen.add(coordinate)
        elif name == formula_tag and coordinate in edits:
            require(stack and stack[-1] == cell_tag and start_text is None
                    and coordinate not in selected and attrs.get("t", "normal") in {"normal", "array"},
                    "Unsupported formula text structure")
            selected.add(coordinate)
            offset = parser.CurrentByteIndex
            end = raw.find(b">", offset)
            require(end >= offset and raw[end - 1:end] != b"/", "Missing formula text")
            start_text = end + 1
        elif start_text is not None:
            raise ValueError("Nested formula text")
        stack.append(name)

    def end(name):
        nonlocal coordinate, start_text
        if name == formula_tag and start_text is not None:
            expected, replacement = edits[coordinate]
            offset = parser.CurrentByteIndex
            existing = raw[start_text:offset].decode("utf-8")
            # XML parsing has already rejected DTD/entities. Only built-in XML
            # character escaping is relevant; compare its exact decoded text.
            require(html.unescape(existing) == expected.removeprefix("="), "Formula XML binding differs")
            ranges.append((start_text, offset, html.escape(replacement.removeprefix("="), quote=False).encode()))
            start_text = None
        if name == cell_tag:
            coordinate = None
        require(stack and stack.pop() == name, "Formula XML nesting differs")

    parser.StartElementHandler, parser.EndElementHandler = start, end
    parser.Parse(raw, True)
    require(selected == set(edits) and not stack, "Selected formula text missing")
    chunks, cursor = [], 0
    for start, end, replacement in ranges:
        require(start >= cursor, "Overlapping formula edits")
        chunks.extend((raw[cursor:start], replacement))
        cursor = end
    chunks.append(raw[cursor:])
    return b"".join(chunks)


def _patch_workbook(path, changes):
    with compat._safe_archive(path) as source:
        parts, by_member = v5._sheet_parts(source), {}
        for (sheet, coordinate), texts in changes.items():
            require(sheet in parts, "Formula sheet missing")
            by_member.setdefault(parts[sheet], {})[coordinate] = texts
        target = io.BytesIO()
        with zipfile.ZipFile(target, "w") as output:
            for item in source.infolist():
                raw = source.read(item.filename)
                if item.filename in by_member:
                    raw = _patch_formula_xml(raw, by_member[item.filename])
                output.writestr(item, raw)
        result = target.getvalue()
        require(len(result) <= v5.MAX_BYTES, "Alias workbook size limit")
        return result


def screening_view(path):
    """Return a detached safety-screen view; it must NEVER be executed."""
    path = safe_path(path)
    v6._preflight(path)
    book = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    try:
        require(not _named_expressions(book), "Named expressions not qualified for aliases")
        require(not v6._observes_format(book), "Formula/format observer not qualified for aliases")
        changes, records = {}, []
        for identity, (kind, formula, region) in _formula_map(book).items():
            canonical = canonical_formula(formula)
            if canonical != formula:
                changes[identity] = formula, canonical
                records.append({"sheet": identity[0], "coordinate": identity[1], "kind": kind, "range": region,
                    "original_formula_sha256": hashlib.sha256(formula.encode()).hexdigest(),
                    "screening_formula_sha256": hashlib.sha256(canonical.encode()).hexdigest()})
        require(changes, "No qualified function alias")
        raw = _patch_workbook(path, changes)
        return raw, {"version": ALIAS_VERSION, "mode": "native_original_bytes",
                     "original_input_sha256": v5.sha(path), "execution_input_sha256": v5.sha(path),
                     "screening_view_sha256": hashlib.sha256(raw).hexdigest(), "qualified_formulas": records}
    finally:
        book.close()


def comparison_view(original, exported):
    """Restore only qualified alias spelling, never a changed expression."""
    original, exported = safe_path(original), safe_path(exported)
    v6._preflight(original)
    v6._preflight(exported)
    before = openpyxl.load_workbook(original, data_only=False, keep_links=False)
    after = None
    try:
        after = openpyxl.load_workbook(exported, data_only=False, keep_links=False)
        require(before.sheetnames == after.sheetnames and before.epoch == after.epoch,
                "Alias comparison workbook identity differs")
        require(not _named_expressions(before) and not _named_expressions(after)
                and not v6._observes_format(before) and not v6._observes_format(after),
                "Alias comparison observer or named expression")
        originals, actuals, changes = _formula_map(before), _formula_map(after), {}
        for identity, (kind, formula, region) in originals.items():
            expected = canonical_formula(formula)
            if expected == formula:
                continue
            require(identity in actuals, "Alias formula removed")
            current_kind, current, current_region = actuals[identity]
            require(kind == current_kind and region == current_region
                    and canonical_formula(current) == expected, "Alias formula or array range changed")
            if current != formula:
                changes[identity] = current, formula
        raw = _patch_workbook(exported, changes) if changes else exported.read_bytes()
        return raw, {"comparison_sha256": hashlib.sha256(raw).hexdigest(),
                     "engine_output_sha256": v5.sha(exported), "restored_formula_count": len(changes)}
    finally:
        before.close()
        if after is not None:
            after.close()


class Recalculator(v7.Recalculator):
    @property
    def identity(self):
        identity = super().identity
        identity["version"] = VERSION
        identity["function_alias_view"] = ALIAS_VERSION
        identity["qualified_function_execution"] = "original_bytes_namespace_preserved"
        identity["sources"] = {**identity["sources"], str(Path(__file__)): v5.sha(Path(__file__))}
        return identity

    def _run_native(self, path):
        """Small version-local v5 runner: only its initial screen is specialized.

        No frozen source or global is patched. All container limits, worker,
        numeric manifest/proof checks and cleanup semantics are reused verbatim.
        This method screens even when called directly; no unchecked input hook.
        """
        path = safe_path(path)
        started = time.monotonic()
        base = {"input_sha256": v5.sha(path), "identity": self.identity,
                "cleanup_confirmed": True, "model_api_calls": 0, "container_execution_attempted": False}

        def done(status, reason, **extra):
            return seal({**base, "status": status, "reason": reason,
                         "duration_seconds": time.monotonic() - started, **extra})

        try:
            require(v5.screen(path) == "unsupported_extended_function", "Not qualified extended input")
            screened, proof = screening_view(path)
            with tempfile.TemporaryDirectory(prefix="sheet-function-screen-only-") as directory:
                screen_path = Path(directory).resolve() / "never-executed.xlsx"
                screen_path.write_bytes(screened)
                require(v5.screen(screen_path) is None, "Unqualified extension or input remains")
            base["function_alias_view"] = proof
        except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError,
                expat.ExpatError, zipfile.BadZipFile, TokenizerError):
            return done("unknown", "unsupported_extended_function")
        if platform.system() != "Linux" or shutil.which("docker") is None:
            return done("unknown", "linux_docker_unavailable")
        try:
            manifest = v5._numeric_manifest(path)
            manifest_bytes = json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
            require(len(manifest_bytes) <= v5.MAX_BYTES, "Numeric manifest size limit")
        except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError, zipfile.BadZipFile):
            return done("unknown", "numeric_literal_manifest_unavailable")
        try:
            inspect = v5._bounded_command(["docker", "image", "inspect", "--format", "{{json .}}", self.image], 10)
        except (OSError, subprocess.TimeoutExpired):
            return done("unknown", "pinned_image_or_daemon_unavailable")
        try:
            image = json.loads(inspect.stdout)
            require(inspect.code == 0 and not inspect.timed_out and not inspect.overflow
                    and image["Id"] == self.image and image["Os"] == "linux"
                    and not image["Config"].get("Volumes"), "Unsupported image")
        except (ValueError, KeyError, TypeError):
            return done("unknown", "pinned_image_unavailable_or_unsafe")
        name = "skillopt-sheet-recalc-" + uuid.uuid4().hex
        with tempfile.TemporaryDirectory(prefix="sheet-recalc-native-input-") as directory:
            directory = Path(directory).resolve()
            directory.chmod(0o755)
            frozen = directory / "book.xlsx"
            frozen.write_bytes(path.read_bytes())
            require(v5.sha(frozen) == base["input_sha256"] == proof["execution_input_sha256"],
                    "Original input changed before native execution")
            frozen.chmod(0o444)
            manifest_path = directory / "numeric_literals.json"
            manifest_path.write_bytes(manifest_bytes)
            manifest_path.chmod(0o444)
            worker = safe_path(Path(v5.__file__).with_name("sheet_recalc_worker.py"))
            command = ["docker", "run", "--name", name, "--pull=never", "--network=none", "--read-only",
                       "--user=65534:65534", "--cap-drop=ALL", "--security-opt=no-new-privileges",
                       "--pids-limit=128", "--memory=1024m", "--memory-swap=1024m", "--cpus=1",
                       "--tmpfs=/tmp:rw,nosuid,nodev,noexec,size=256m,mode=1777",
                       "--mount", f"type=bind,src={directory},dst=/input,readonly",
                       "--mount", f"type=bind,src={worker},dst=/worker.py,readonly",
                       "--env=HOME=/tmp", "--env=SAL_USE_VCLPLUGIN=svp",
                       "--env=SKILLOPT_SHEET_RECALC_CONTAINER=1", self.image]
            try:
                base["container_execution_attempted"] = True
                try:
                    command_result = v5._bounded_command(command, self.timeout, 12 * 1024 * 1024)
                except (OSError, subprocess.TimeoutExpired):
                    command_result = v5._Command(None, unavailable=True)
            finally:
                try:
                    cleanup = v5._bounded_command(["docker", "rm", "-f", name], 10)
                except (OSError, subprocess.TimeoutExpired):
                    cleanup = v5._Command(None, unavailable=True)
            base["cleanup_confirmed"] = cleanup.code == 0 and not cleanup.timed_out
            if not base["cleanup_confirmed"]:
                return done("unknown", "container_cleanup_unconfirmed")
            if command_result.timed_out or command_result.overflow or command_result.code != 0:
                return done("unknown", "recalculation_timeout" if command_result.timed_out else "recalculation_execution_error")
            try:
                result = json.loads(command_result.stdout)
                if result["status"] != "available":
                    return done("unknown", result.get("reason", "recalculation_unavailable"),
                                libreoffice_version=result.get("libreoffice_version"))
                raw = base64.b64decode(result["output_base64"], validate=True)
                require(len(raw) <= v5.MAX_BYTES and hashlib.sha256(raw).hexdigest() == result["output_sha256"],
                        "Output bytes mismatch")
                require(result.get("calculate_all") is True and result.get("macros") == "never_execute"
                        and result.get("links") == "no_update", "Worker policy mismatch")
                restored, changed = v5._restore_numeric_export(raw, manifest, result)
                return done("available", "independently_recalculated", output_base64=base64.b64encode(restored).decode(),
                            output_sha256=hashlib.sha256(restored).hexdigest(),
                            raw_export_base64=result["output_base64"], raw_export_sha256=result["output_sha256"],
                            numeric_manifest_sha256=result["numeric_manifest_sha256"],
                            numeric_literals_verified=len(manifest["cells"]), numeric_literal_texts_restored=changed,
                            libreoffice_version=result["libreoffice_version"])
            except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError,
                    expat.ExpatError, zipfile.BadZipFile):
                return done("unknown", "invalid_recalculation_receipt")

    def run(self, path):
        path = safe_path(path)
        original_hash = v5.sha(path)
        if v5.screen(path) != "unsupported_extended_function":
            return super().run(path)
        receipt, proof = None, None
        try:
            receipt = self._run_native(path)
            proof = receipt.get("function_alias_view")
            result = {"input_sha256": original_hash, "identity": self.identity,
                "execution_input_sha256": original_hash, "execution_receipt": receipt,
                "function_alias_view": proof,
                **{k: receipt[k] for k in ("status", "reason", "cleanup_confirmed", "model_api_calls",
                                          "container_execution_attempted", "duration_seconds") if k in receipt}}
            if receipt["status"] != "available":
                return seal(result)
            require(receipt["input_sha256"] == original_hash == proof["execution_input_sha256"]
                    == proof["original_input_sha256"] and receipt["identity"] == self.identity,
                    "Actual native execution binding differs")
            manifest = v5._numeric_manifest(path)
            encoded = json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
            require(type(receipt.get("numeric_literals_verified")) is int
                    and receipt["numeric_literals_verified"] == len(manifest["cells"])
                    and receipt.get("numeric_manifest_sha256") == hashlib.sha256(encoded).hexdigest(),
                    "Missing native numeric computation proof")
            with tempfile.TemporaryDirectory(prefix="sheet-qualified-native-output-") as directory:
                root = Path(directory).resolve()
                exported = root / "engine-view.xlsx"
                raw = base64.b64decode(receipt["output_base64"], validate=True)
                require(hashlib.sha256(raw).hexdigest() == receipt["output_sha256"], "Actual output hash differs")
                exported.write_bytes(raw)
                view, output_proof = comparison_view(path, exported)
                restored_formula = root / "formula-spelling-view.xlsx"
                restored_formula.write_bytes(view)
                view, read_proof = v7.comparison_view(path, restored_formula)
                checked = root / "comparison-view.xlsx"
                checked.write_bytes(view)
                require(v5._recalc_valid(path, checked) is None, "Original content guard rejected alias view")
                require(v5.sha(path) == original_hash, "Original input changed during alias execution")
                return seal({**result, "function_alias_view": {**proof, **output_proof},
                    "comparison_view": read_proof,
                    "output_base64": base64.b64encode(view).decode(), "output_sha256": hashlib.sha256(view).hexdigest()})
        except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError,
                expat.ExpatError, zipfile.BadZipFile, TokenizerError):
            result = {"input_sha256": original_hash, "identity": self.identity, "status": "unknown",
                "reason": "qualified_native_function_view_unverified",
                "cleanup_confirmed": receipt.get("cleanup_confirmed") is True if receipt else True,
                "container_execution_attempted": bool(receipt and receipt.get("container_execution_attempted")),
                "duration_seconds": receipt.get("duration_seconds", 0.) if receipt else 0., "model_api_calls": 0}
            if receipt is not None:
                result.update(execution_receipt=receipt, execution_input_sha256=original_hash, function_alias_view=proof)
            return seal(result)
