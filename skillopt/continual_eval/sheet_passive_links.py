"""Opt-in passive HTTPS worksheet links; original bytes still run offline.

A worksheet hyperlink relationship is not an external-workbook dependency.
Only a one-to-one, single-cell, HTTPS hyperlink shape is qualified here. A
detached screen-only copy removes those links to apply the complete frozen v5
screen to everything else. That copy is never executed. Original bytes enter
the same isolated worker, then both link semantics and v5 content are checked.
No hyperlink is clicked, fetched, removed from the output, or treated as truth.

Shape source: learn.microsoft.com/en-us/dotnet/api/
documentformat.openxml.spreadsheet.hyperlinks?view=openxml-3.0.1
Engineering tests do not authorize this new execution scope on their own.
"""
from __future__ import annotations

import base64
import hashlib
import html
import io
import json
import platform
import re
import shutil
import subprocess
import tempfile
import time
import uuid
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path, PurePosixPath
from urllib.parse import urlsplit
from xml.parsers import expat

import openpyxl

from skillopt.coevolution_v5.core import seal
from skillopt.validator_pilot.api import digest

from . import sheet_recalc as v5
from . import sheet_recalc_functions as v8
from . import sheet_recalc_v6 as v6
from . import sheet_recalc_v7 as v7
from . import spreadsheet_compat as compat
from .core import require, safe_path

VERSION = "spreadsheet-independent-lo-passive-links-v9c"
PROFILE = "single-cell-https-passive-hyperlinks-v3"
TOOLTIP_VIEW = "original-missing-hover-tooltip-read-view-v1"
DISPLAY_VIEW = "original-href-display-unchanged-literal-read-view-v1"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MAIN = "{" + compat.MAIN_NS + "}"
REL = "{" + REL_NS + "}"
RID = "{" + DOC_REL + "}id"
MAX_LINKS = 1000
ERRORS = (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError,
          expat.ExpatError, zipfile.BadZipFile)


def _empty(node):
    return not len(node) and not (node.text or "").strip() and not (node.tail or "").strip()


def _link_parts(archive):
    """Return private semantic records and exact parts for screen-only edits."""
    sheets = v5._sheet_parts(archive)
    relationships, roots, permitted, records = {}, {}, set(), []
    for sheet, member in sheets.items():
        part = PurePosixPath(member)
        relationship_part = str(part.parent / "_rels" / (part.name + ".rels"))
        if relationship_part in archive.namelist():
            root = v5._xml(archive, relationship_part)
            require(root.tag == REL + "Relationships" and not root.attrib
                    and not (root.text or "").strip(), "Invalid worksheet relationships")
            ids = set()
            for item in root:
                require(item.tag == REL + "Relationship" and item.get("Id")
                        and item.get("Id") not in ids and _empty(item), "Ambiguous relationship identity")
                ids.add(item.get("Id"))
                if item.get("TargetMode", "").lower() != "external":
                    continue
                require(set(item.attrib) == {"Id", "Type", "Target", "TargetMode"}
                        and item.get("Type") == DOC_REL + "/hyperlink"
                        and item.get("TargetMode") == "External", "Non-passive external relationship")
                target = item.get("Target", "")
                url = urlsplit(target)
                require(target.startswith("https://") and not any(ord(c) <= 32 or ord(c) == 127 for c in target)
                        and "\\" not in target and url.scheme == "https" and url.hostname
                        and url.username is None and url.password is None
                        and url.port in (None, 443), "Unqualified hyperlink destination")
                relationships[(member, item.get("Id"))] = (relationship_part, item, target)
            roots[relationship_part] = root
        worksheet = v5._xml(archive, member)
        require(worksheet.tag == MAIN + "worksheet", "Invalid worksheet root")
        containers = worksheet.findall(MAIN + "hyperlinks")
        require(len(containers) <= 1, "Duplicate hyperlink containers")
        represented = set()
        for cell in worksheet.iter(MAIN + "c"):
            coordinate = cell.get("r")
            require(coordinate and coordinate not in represented, "Duplicate or missing cell identity")
            represented.add(coordinate)
        seen, used = set(), set()
        for container in containers:
            require(not container.attrib and not (container.text or "").strip()
                    and not (container.tail or "").strip(), "Invalid hyperlinks container")
            for item in container:
                require(item.tag == MAIN + "hyperlink" and _empty(item)
                        and {"ref", RID} <= set(item.attrib)
                        and set(item.attrib) <= {"ref", RID, "display", "tooltip"}, "Unqualified hyperlink shape")
                coordinate, identity = item.get("ref"), item.get(RID)
                column, row = compat._legacy()._parse_cell(coordinate)
                require(1 <= column <= 16384 and 1 <= row <= 1048576, "Hyperlink outside worksheet bounds")
                require(coordinate in represented and coordinate not in seen and identity not in used,
                        "Missing, duplicate, or ranged hyperlink cell")
                seen.add(coordinate)
                used.add(identity)
                key = member, identity
                require(key in relationships, "Unbound hyperlink")
                relpart, relation, target = relationships[key]
                permitted.add((relpart, identity))
                records.append({"sheet": sheet, "coordinate": coordinate, "target": target,
                                "display": item.get("display"), "tooltip": item.get("tooltip")})
            worksheet.remove(container)
        require({key[1] for key in relationships if key[0] == member} == used,
                "Unused external hyperlink relationship")
        roots[member] = worksheet
    require(0 < len(records) <= MAX_LINKS, "No qualified passive hyperlinks or budget exceeded")
    # Every external marker, including ones outside normal .rels locations,
    # must be exactly a whitelisted worksheet relationship. No hidden payload.
    for member in archive.namelist():
        if not member.lower().endswith((".xml", ".rels")):
            continue
        tree = v5._xml(archive, member)
        for node in tree.iter():
            if node.get("TargetMode", "").lower() == "external":
                require(tree.tag == REL + "Relationships" and node in list(tree)
                        and node.tag == REL + "Relationship"
                        and (member, node.get("Id")) in permitted, "Unqualified external marker")
    for relpart, identity in permitted:
        root = roots[relpart]
        root.remove(next(item for item in root if item.get("Id") == identity))
    return sorted(records, key=lambda row: (row["sheet"], row["coordinate"])), roots


def screening_view(path):
    """No output cache changes; this detached XML copy must never be run."""
    path = safe_path(path)
    v6._preflight(path)
    with compat._safe_archive(path) as archive:
        records, edits = _link_parts(archive)
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as target:
            for entry in archive.infolist():
                target.writestr(entry, ET.tostring(edits[entry.filename], encoding="utf-8")
                               if entry.filename in edits else archive.read(entry.filename))
    raw = output.getvalue()
    require(len(raw) <= v5.MAX_BYTES, "Passive link screen size limit")
    return raw, {"version": PROFILE, "mode": "original_bytes_no_link_activation",
                 "original_input_sha256": v5.sha(path), "execution_input_sha256": v5.sha(path),
                 "screening_view_sha256": hashlib.sha256(raw).hexdigest(),
                 "hyperlink_count": len(records), "hyperlink_semantics_sha256": digest(records)}


def _checked_screen(path):
    require(v5.screen(path) == "unsupported_external_relationship", "Not a passive-link candidate")
    raw, proof = screening_view(path)
    with tempfile.TemporaryDirectory(prefix="sheet-passive-screen-only-") as directory:
        screened = Path(directory).resolve() / "never-executed.xlsx"
        screened.write_bytes(raw)
        require(v5.screen(screened) is None, "Another unsupported dependency remains")
    return proof


def verify_preserved(original, exported):
    """Transport relationship IDs may change; all accepted link semantics may not."""
    left, right = [], []
    for path, records in ((original, left), (exported, right)):
        v6._preflight(path)
        with compat._safe_archive(path) as archive:
            values, _ = _link_parts(archive)
            records.extend(values)
    require(left == right, "Passive hyperlink semantics changed")
    return {"hyperlink_count": len(left), "hyperlink_semantics_sha256": digest(left)}


def _literal_link_text(archive, book, sheet, coordinate):
    """Require an explicitly represented string, not a formula or spill result."""
    root = v5._xml(archive, v5._sheet_parts(archive)[sheet])
    cells = [c for c in root.iter(MAIN + "c") if c.get("r") == coordinate]
    require(len(cells) == 1, "Display cell missing or duplicated")
    cell = cells[0]
    require(cell.get("t") in {"s", "inlineStr"} and cell.find(MAIN + "f") is None
            and "cm" not in cell.attrib and "vm" not in cell.attrib,
            "Display cell is not an unannotated literal string")
    payloads = cell.findall(MAIN + ("v" if cell.get("t") == "s" else "is"))
    require(len(payloads) == 1 and (cell.get("t") != "s" or (payloads[0].text or "").isdigit()),
            "Display string payload missing")
    # Neither a cached array follower nor dynamic-array metadata proves a
    # literal, even if openpyxl exposes it as a string.
    require("xl/metadata.xml" not in archive.namelist(), "Display metadata unsupported")
    column, row = compat._legacy()._parse_cell(coordinate)
    for formula in root.iter(MAIN + "f"):
        if formula.get("t") in {"array", "dataTable"}:
            bounds = openpyxl.utils.cell.range_boundaries(formula.get("ref", ""))
            require(all(type(v) is int for v in bounds), "Ambiguous formula coverage")
            left, top, right, bottom = bounds
            require(not (left <= column <= right and top <= row <= bottom), "Display overlaps formula result")
    for table in book[sheet].tables.values():
        left, top, right, bottom = openpyxl.utils.cell.range_boundaries(table.ref)
        require(not (left <= column <= right and top <= row <= bottom), "Display overlaps table result")
    value = book[sheet][coordinate]
    require(value.data_type == "s" and isinstance(value.value, str) and bool(value.value),
            "Display cell string is missing or empty")
    return value.value


def _restore_display_xml(raw, changes):
    """Replace only the known display attribute value, preserving all f/v bytes."""
    raw.decode("utf-8", errors="strict")
    parser = expat.ParserCreate(namespace_separator="}")
    edits, seen = [], set()

    def start(name, attributes):
        if name != compat.MAIN_NS + "}hyperlink" or attributes.get("ref") not in changes:
            return
        coordinate = attributes["ref"]
        old, new = changes[coordinate]
        require(coordinate not in seen and attributes.get("display") == new,
                "Display attribute missing, changed, or duplicated")
        seen.add(coordinate)
        cursor, quote = parser.CurrentByteIndex, None
        while cursor < len(raw):
            char = raw[cursor]
            if quote is not None:
                if char == quote:
                    quote = None
            elif char in (34, 39):
                quote = char
            elif char == 62:
                break
            cursor += 1
        require(cursor < len(raw), "Unclosed hyperlink start tag")
        # Tokenize attributes, not a substring search: a tooltip may itself
        # contain text resembling an XML attribute.
        tag = raw[parser.CurrentByteIndex:cursor]
        matches = list(re.finditer(rb"\s+([^\s=/>]+)\s*=\s*([\"'])(.*?)\2", tag, re.DOTALL))
        display = [m for m in matches if m.group(1) == b"display"]
        require(len(display) == 1, "Display lexical identity ambiguous")
        match = display[0]
        escaped = html.escape(old, quote=True).replace("\t", "&#9;").replace("\n", "&#10;").replace("\r", "&#13;")
        edits.append((parser.CurrentByteIndex + match.start(3), parser.CurrentByteIndex + match.end(3),
                      escaped.encode("utf-8")))

    parser.StartElementHandler = start
    parser.Parse(raw, True)
    require(seen == set(changes), "Display target missing")
    pieces, previous = [], 0
    for start, end, replacement in edits:
        require(start >= previous, "Overlapping display edits")
        pieces.extend((raw[previous:start], replacement))
        previous = end
    pieces.append(raw[previous:])
    return b"".join(pieces)


def display_comparison_view(original, exported):
    """Restore exactly the observed href-to-cell-label export transformation.

    ISO/IEC 29500-1 18.3.1.47 (Microsoft Learn Hyperlink) distinguishes the
    hyperlink object's display property from the string-table cell value.
    Neither is ignored: the detached view restores the original property only
    when it was exactly the href, the new property is exactly the unchanged
    explicit literal string, and all destinations/coordinates stay identical.
    No actual execution input, formula, cached value, or original is modified.
    Source: learn.microsoft.com/en-us/dotnet/api/
    documentformat.openxml.spreadsheet.hyperlink?view=openxml-3.0.1
    """
    original, exported = safe_path(original), safe_path(exported)
    for path in (original, exported):
        v6._preflight(path)
    with compat._safe_archive(original) as source, compat._safe_archive(exported) as output:
        before, _ = _link_parts(source)
        after, _ = _link_parts(output)
        require(len(before) == len(after), "Hyperlink count changed")
        changes = {}
        books = []
        try:
            for old, new in zip(before, after):
                require({k: v for k, v in old.items() if k not in {"display", "tooltip"}}
                        == {k: v for k, v in new.items() if k not in {"display", "tooltip"}},
                        "Hyperlink destination or coordinate changed")
                require(old["tooltip"] == new["tooltip"] or old["tooltip"] is not None and new["tooltip"] is None,
                        "Tooltip changed or newly added")
                if old["display"] == new["display"]:
                    continue
                require(old["display"] == old["target"] and isinstance(new["display"], str),
                        "Display change is not the qualified href-to-literal pattern")
                if not books:
                    books.append(openpyxl.load_workbook(original, data_only=False, keep_links=False))
                    books.append(openpyxl.load_workbook(exported, data_only=False, keep_links=False))
                    require(books[0].sheetnames == books[1].sheetnames and books[0].epoch == books[1].epoch,
                            "Display view workbook identity changed")
                    require(all(not v6._observes_format(b) and not v8._named_expressions(b) for b in books),
                            "Metadata observer or name prevents display restoration")
                left = _literal_link_text(source, books[0], old["sheet"], old["coordinate"])
                right = _literal_link_text(output, books[1], old["sheet"], old["coordinate"])
                require(left == right == new["display"], "Display does not equal the unchanged visible cell string")
                changes.setdefault(old["sheet"], {})[old["coordinate"]] = old["display"], new["display"]
        finally:
            for book in books:
                book.close()
        raw = exported.read_bytes()
        if changes:
            parts = v5._sheet_parts(output)
            edits = {parts[sheet]: entries for sheet, entries in changes.items()}
            target = io.BytesIO()
            with zipfile.ZipFile(target, "w") as dest:
                for entry in output.infolist():
                    content = output.read(entry.filename)
                    if entry.filename in edits:
                        content = _restore_display_xml(content, edits[entry.filename])
                    dest.writestr(entry, content)
            raw = target.getvalue()
            require(len(raw) <= v5.MAX_BYTES, "Display view size limit")
        return raw, {"version": DISPLAY_VIEW, "original_input_sha256": v5.sha(original),
                     "raw_engine_output_sha256": v5.sha(exported), "view_sha256": hashlib.sha256(raw).hexdigest(),
                     "display_attributes_restored": sum(len(values) for values in changes.values()),
                     "formula_or_value_bytes_changed": False}


def _restore_tooltip_xml(raw, tooltips):
    """Insert only a missing hyperlink tooltip attribute; keep every f/v byte."""
    raw.decode("utf-8", errors="strict")
    parser = expat.ParserCreate(namespace_separator="}")
    edits, seen = [], set()

    def start(name, attributes):
        if name != compat.MAIN_NS + "}hyperlink" or attributes.get("ref") not in tooltips:
            return
        coordinate = attributes["ref"]
        require(coordinate not in seen and "tooltip" not in attributes, "Tooltip already present or duplicated")
        seen.add(coordinate)
        cursor, quote = parser.CurrentByteIndex, None
        while cursor < len(raw):
            char = raw[cursor]
            if quote is not None:
                if char == quote:
                    quote = None
            elif char in (34, 39):
                quote = char
            elif char == 62:
                break
            cursor += 1
        require(cursor < len(raw), "Unclosed hyperlink start tag")
        position = cursor - 1 if raw[cursor - 1:cursor] == b"/" else cursor
        escaped = html.escape(tooltips[coordinate], quote=True)
        escaped = escaped.replace("\t", "&#9;").replace("\n", "&#10;").replace("\r", "&#13;")
        edits.append((position, (' tooltip="' + escaped + '"').encode("utf-8")))

    parser.StartElementHandler = start
    parser.Parse(raw, True)
    require(seen == set(tooltips), "Tooltip target missing")
    pieces, previous = [], 0
    for position, insertion in edits:
        require(position >= previous, "Overlapping tooltip edits")
        pieces.extend((raw[previous:position], insertion))
        previous = position
    pieces.append(raw[previous:])
    return b"".join(pieces)


def tooltip_comparison_view(original, exported):
    """Restore lost hover text only, never accept changed link destinations.

    ISO/IEC 29500-1 18.3.1.47, as reproduced in Microsoft Learn Hyperlink,
    defines tooltip as additional user help / hover text, not a cell result.
    Source: learn.microsoft.com/en-us/dotnet/api/
    documentformat.openxml.spreadsheet.hyperlink?view=openxml-3.0.1
    Calc's raw export may lose this presentation attribute. The original input
    and raw engine export stay unchanged; this view copies its original text
    only when all other link fields match and no observer/named expression can
    consume metadata. Arbitrary tooltip changes are not normalized.
    """
    original, exported = safe_path(original), safe_path(exported)
    for path in (original, exported):
        v6._preflight(path)
    with compat._safe_archive(original) as source, compat._safe_archive(exported) as output:
        before, _ = _link_parts(source)
        after, _ = _link_parts(output)
        require(len(before) == len(after), "Hyperlink count changed")
        changes = {}
        for old, new in zip(before, after):
            require({k: v for k, v in old.items() if k != "tooltip"}
                    == {k: v for k, v in new.items() if k != "tooltip"}, "Hyperlink destination or display changed")
            if old["tooltip"] != new["tooltip"]:
                require(old["tooltip"] is not None and new["tooltip"] is None, "Tooltip changed or newly added")
                changes.setdefault(old["sheet"], {})[old["coordinate"]] = old["tooltip"]
        raw = exported.read_bytes()
        if changes:
            books = []
            try:
                books.append(openpyxl.load_workbook(original, data_only=False, keep_links=False))
                books.append(openpyxl.load_workbook(exported, data_only=False, keep_links=False))
                require(books[0].sheetnames == books[1].sheetnames and books[0].epoch == books[1].epoch,
                        "Tooltip view workbook identity changed")
                require(all(not v6._observes_format(b) and not v8._named_expressions(b) for b in books),
                        "Metadata observer or name prevents tooltip restoration")
            finally:
                for book in books:
                    book.close()
            parts = v5._sheet_parts(output)
            edits = {parts[sheet]: entries for sheet, entries in changes.items()}
            target = io.BytesIO()
            with zipfile.ZipFile(target, "w") as dest:
                for entry in output.infolist():
                    content = output.read(entry.filename)
                    if entry.filename in edits:
                        content = _restore_tooltip_xml(content, edits[entry.filename])
                    dest.writestr(entry, content)
            raw = target.getvalue()
            require(len(raw) <= v5.MAX_BYTES, "Tooltip view size limit")
        return raw, {"version": TOOLTIP_VIEW, "original_input_sha256": v5.sha(original),
                     "raw_engine_output_sha256": v5.sha(exported), "view_sha256": hashlib.sha256(raw).hexdigest(),
                     "tooltip_attributes_restored": sum(len(values) for values in changes.values()),
                     "formula_or_value_bytes_changed": False}


class Recalculator(v8.Recalculator):
    @property
    def identity(self):
        result = super().identity
        result.update(version=VERSION, passive_hyperlinks=PROFILE, tooltip_read_view=TOOLTIP_VIEW,
                      display_read_view=DISPLAY_VIEW)
        result["sources"] = {**result["sources"], str(Path(__file__)): v5.sha(Path(__file__))}
        return result

    def _run_passive_native(self, path):
        """Version-local runner; unchanged v5 worker/limits/cleanup, new screen."""
        path = safe_path(path)
        started = time.monotonic()
        base = {"input_sha256": v5.sha(path), "identity": self.identity,
                "cleanup_confirmed": True, "model_api_calls": 0, "container_execution_attempted": False}

        def done(status, reason, **extra):
            return seal({**base, "status": status, "reason": reason,
                         "duration_seconds": time.monotonic() - started, **extra})

        try:
            proof = _checked_screen(path)
            base["passive_hyperlinks"] = proof
        except ERRORS:
            return done("unknown", "unsupported_external_relationship")
        if platform.system() != "Linux" or shutil.which("docker") is None:
            return done("unknown", "linux_docker_unavailable")
        try:
            manifest = v5._numeric_manifest(path)
            manifest_bytes = json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
            require(len(manifest_bytes) <= v5.MAX_BYTES, "Numeric manifest size limit")
        except ERRORS:
            return done("unknown", "numeric_literal_manifest_unavailable")
        try:
            inspect = v5._bounded_command(["docker", "image", "inspect", "--format", "{{json .}}", self.image], 10)
        except (OSError, subprocess.TimeoutExpired):
            return done("unknown", "pinned_image_or_daemon_unavailable")
        try:
            image = json.loads(inspect.stdout)
            require(inspect.code == 0 and not inspect.timed_out and not inspect.overflow
                    and image["Id"] == self.image and image["Os"] == "linux"
                    and not image["Config"].get("Volumes"), "Unsupported image")
        except (ValueError, KeyError, TypeError):
            return done("unknown", "pinned_image_unavailable_or_unsafe")
        name = "skillopt-sheet-recalc-" + uuid.uuid4().hex
        with tempfile.TemporaryDirectory(prefix="sheet-passive-native-") as directory:
            directory = Path(directory).resolve()
            directory.chmod(0o755)
            frozen = directory / "book.xlsx"
            frozen.write_bytes(path.read_bytes())
            require(v5.sha(frozen) == base["input_sha256"] == proof["execution_input_sha256"],
                    "Original passive-link input changed")
            frozen.chmod(0o444)
            literals = directory / "numeric_literals.json"
            literals.write_bytes(manifest_bytes)
            literals.chmod(0o444)
            worker = safe_path(Path(v5.__file__).with_name("sheet_recalc_worker.py"))
            command = ["docker", "run", "--name", name, "--pull=never", "--network=none", "--read-only",
                       "--user=65534:65534", "--cap-drop=ALL", "--security-opt=no-new-privileges",
                       "--pids-limit=128", "--memory=1024m", "--memory-swap=1024m", "--cpus=1",
                       "--tmpfs=/tmp:rw,nosuid,nodev,noexec,size=256m,mode=1777",
                       "--mount", f"type=bind,src={directory},dst=/input,readonly",
                       "--mount", f"type=bind,src={worker},dst=/worker.py,readonly",
                       "--env=HOME=/tmp", "--env=SAL_USE_VCLPLUGIN=svp",
                       "--env=SKILLOPT_SHEET_RECALC_CONTAINER=1", self.image]
            try:
                base["container_execution_attempted"] = True
                try:
                    response = v5._bounded_command(command, self.timeout, 12 * 1024 * 1024)
                except (OSError, subprocess.TimeoutExpired):
                    response = v5._Command(None, unavailable=True)
            finally:
                try:
                    cleanup = v5._bounded_command(["docker", "rm", "-f", name], 10)
                except (OSError, subprocess.TimeoutExpired):
                    cleanup = v5._Command(None, unavailable=True)
            base["cleanup_confirmed"] = cleanup.code == 0 and not cleanup.timed_out
            if not base["cleanup_confirmed"]:
                return done("unknown", "container_cleanup_unconfirmed")
            if response.timed_out or response.overflow or response.code != 0:
                return done("unknown", "recalculation_timeout" if response.timed_out else "recalculation_execution_error")
            try:
                result = json.loads(response.stdout)
                if result["status"] != "available":
                    return done("unknown", result.get("reason", "recalculation_unavailable"),
                                libreoffice_version=result.get("libreoffice_version"))
                raw = base64.b64decode(result["output_base64"], validate=True)
                require(len(raw) <= v5.MAX_BYTES and hashlib.sha256(raw).hexdigest() == result["output_sha256"],
                        "Output bytes mismatch")
                require(result.get("calculate_all") is True and result.get("macros") == "never_execute"
                        and result.get("links") == "no_update", "Worker policy mismatch")
                restored, changed = v5._restore_numeric_export(raw, manifest, result)
                return done("available", "independently_recalculated", output_base64=base64.b64encode(restored).decode(),
                            output_sha256=hashlib.sha256(restored).hexdigest(), raw_export_base64=result["output_base64"],
                            raw_export_sha256=result["output_sha256"], numeric_manifest_sha256=result["numeric_manifest_sha256"],
                            numeric_literals_verified=len(manifest["cells"]), numeric_literal_texts_restored=changed,
                            libreoffice_version=result["libreoffice_version"])
            except ERRORS:
                return done("unknown", "invalid_recalculation_receipt")

    def run(self, path):
        path = safe_path(path)
        if v5.screen(path) != "unsupported_external_relationship":
            return super().run(path)
        original_hash, receipt = v5.sha(path), None
        try:
            receipt = self._run_passive_native(path)
            proof = receipt.get("passive_hyperlinks")
            result = {"input_sha256": original_hash, "execution_input_sha256": original_hash,
                      "identity": self.identity, "execution_receipt": receipt, "passive_hyperlinks": proof,
                      **{key: receipt[key] for key in ("status", "reason", "cleanup_confirmed", "model_api_calls",
                             "container_execution_attempted", "duration_seconds") if key in receipt}}
            if receipt["status"] != "available":
                return seal(result)
            require(receipt["input_sha256"] == original_hash == proof["execution_input_sha256"]
                    == proof["original_input_sha256"] and receipt["identity"] == self.identity,
                    "Native passive-link identity differs")
            manifest = v5._numeric_manifest(path)
            encoded = json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
            require(type(receipt.get("numeric_literals_verified")) is int
                    and receipt["numeric_literals_verified"] == len(manifest["cells"])
                    and receipt.get("numeric_manifest_sha256") == hashlib.sha256(encoded).hexdigest(),
                    "Missing native numeric proof")
            raw = base64.b64decode(receipt["output_base64"], validate=True)
            require(hashlib.sha256(raw).hexdigest() == receipt["output_sha256"], "Export bytes differ")
            with tempfile.TemporaryDirectory(prefix="sheet-passive-read-view-") as directory:
                exported = Path(directory).resolve() / "export.xlsx"
                checked = Path(directory).resolve() / "view.xlsx"
                exported.write_bytes(raw)
                display_raw, display_proof = display_comparison_view(path, exported)
                display_view = Path(directory).resolve() / "display-view.xlsx"
                display_view.write_bytes(display_raw)
                tooltip_raw, tooltip_proof = tooltip_comparison_view(path, display_view)
                tooltip_view = Path(directory).resolve() / "tooltip-view.xlsx"
                tooltip_view.write_bytes(tooltip_raw)
                retained = verify_preserved(path, tooltip_view)
                require(retained["hyperlink_semantics_sha256"] == proof["hyperlink_semantics_sha256"],
                        "Screen/execution link binding differs")
                raw, read_proof = v7.comparison_view(path, tooltip_view)
                checked.write_bytes(raw)
                require(v5._recalc_valid(path, checked) is None, "Original full-content guard rejected output")
                require(v5.sha(path) == original_hash, "Original passive-link source changed")
            return seal({**result, "comparison_view": read_proof, "hyperlinks_preserved": retained,
                         "display_comparison_view": display_proof,
                         "tooltip_comparison_view": tooltip_proof,
                         "output_base64": base64.b64encode(raw).decode(), "output_sha256": hashlib.sha256(raw).hexdigest()})
        except ERRORS:
            return seal({"input_sha256": original_hash, "identity": self.identity, "status": "unknown",
                         "reason": "passive_hyperlink_execution_unverified", "model_api_calls": 0,
                         "cleanup_confirmed": receipt.get("cleanup_confirmed") is True if receipt else True,
                         "container_execution_attempted": bool(receipt and receipt.get("container_execution_attempted")),
                         "duration_seconds": receipt.get("duration_seconds", 0.) if receipt else 0.,
                         "execution_receipt": receipt})
