"""Opt-in scalar-view adapter; does not replace the frozen v5 scorer.

Some numeric literals with reserved format 30 are decoded as numbers by the
input reader, then as dates after Calc exports an explicit date format. This
adapter retains Calc's unmodified output and creates a separate comparison
view. Only proven, unchanged binary64 literals receive an existing General
style in that view. Formulas, caches, values and the original workbook are
never edited. A formatting observer prohibits the adaptation.

Engineering controls are not LibreOffice qualification. An independently
qualified new profile is required before this adapter may score real results.
Empty text lost during import and literal errors converted into formulas are
deliberately NOT repaired by this adapter.
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
import math
import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from xml.parsers import expat

import openpyxl

from skillopt.coevolution_v5.core import seal

from . import sheet_recalc as v5
from . import spreadsheet_compat as compat
from .core import require, safe_path

VERSION = "spreadsheet-independent-lo-recalculation-v6b"
VIEW_VERSION = "format30-exact-numeric-comparison-view-v2"
_OBSERVERS = re.compile(
    r"(?:^|[^A-Z0-9_])(?:CELL|GET\.CELL|GET\.WORKBOOK|GET\.DOCUMENT|"
    r"FORMULATEXT|FORMULA|TEXT|DOLLAR|FIXED)\s*\(", re.I)


def _preflight(path):
    """Bound both archives and reject hostile XML before openpyxl sees it."""
    with compat._safe_archive(path) as archive:
        for name in archive.namelist():
            if name.lower().endswith((".xml", ".rels")):
                v5._xml(archive, name)


def _observes_format(book):
    texts = [v5._formula_spec(cell.value)[1] for sheet in book
             for cell in sheet._cells.values() if cell.data_type == "f"]
    texts.extend(item.attr_text for scope in [book.defined_names,
                 *(sheet.defined_names for sheet in book)] for item in scope.values())
    return any(not isinstance(text, str) or _OBSERVERS.search(text) for text in texts)


def _general_style(archive):
    styles = v5._xml(archive, "xl/styles.xml")
    # A malformed custom definition must not override the General/30 meanings.
    custom = styles.find(f"{{{compat.MAIN_NS}}}numFmts")
    custom_formats = {}
    for item in custom if custom is not None else ():
        identity = int(item.get("numFmtId", "-1"))
        require(item.tag == f"{{{compat.MAIN_NS}}}numFmt" and identity >= 164,
                "Reserved number-format override")
        require(identity not in custom_formats, "Duplicate custom number-format identity")
        custom_formats[identity] = item.get("formatCode")
    # Calc 25.2 may export the existing General style as custom ID 164 instead
    # of builtin ID 0. Only that exact format code qualifies; no date/number
    # grammar normalization, case folding or fallback to an arbitrary style.
    general_ids = {0} | {key for key, code in custom_formats.items() if code == "General"}
    xfs = styles.find(f"{{{compat.MAIN_NS}}}cellXfs")
    require(xfs is not None and len(xfs) <= v5.MAX_NUMERIC_LITERALS,
            "Missing or oversized cell styles")
    for index, style in enumerate(xfs):
        if (style.tag == f"{{{compat.MAIN_NS}}}xf"
                and int(style.get("numFmtId", "0")) in general_ids):
            return index
    raise ValueError("No existing General style")


def _numeric_cells(archive):
    """No date coercion: read exact raw numeric literals, never formula caches."""
    cells = {}
    for sheet, member in v5._sheet_parts(archive).items():
        seen = set()
        for cell in v5._xml(archive, member).iter(f"{{{compat.MAIN_NS}}}c"):
            coordinate = cell.get("r")
            require(coordinate and coordinate not in seen, "Duplicate or missing cell")
            seen.add(coordinate)
            values = cell.findall(f"{{{compat.MAIN_NS}}}v")
            if (cell.get("t", "n") == "n" and cell.find(f"{{{compat.MAIN_NS}}}f") is None
                    and len(values) == 1 and values[0].text
                    and v5.NUMERIC_TEXT.fullmatch(values[0].text)):
                value = float(values[0].text)
                if math.isfinite(value):
                    cells[(sheet, coordinate)] = value.hex()
    return cells


def _replace_styles(raw, coordinates, style):
    """Modify only existing numeric cell style attribute bytes; keep all else."""
    # Calc emits UTF-8 worksheet parts. Other encodings remain unsupported for
    # this optional view, although v5 can screen certain UTF-16 metadata parts.
    raw.decode("utf-8", errors="strict")
    parser = expat.ParserCreate(namespace_separator="}")
    edits, seen = [], set()

    def start(name, attrs):
        if name != compat.MAIN_NS + "}c" or attrs.get("r") not in coordinates:
            return
        coordinate = attrs["r"]
        require(coordinate not in seen and attrs.get("t", "n") == "n", "Changed cell identity")
        seen.add(coordinate)
        begin = parser.CurrentByteIndex
        end = raw.find(b">", begin)
        require(end >= begin, "Missing cell opening tag")
        opening = raw[begin:end + 1]
        # Restrict to ordinary unescaped decimal style IDs in the opening tag.
        matches = list(re.finditer(rb"\s+s\s*=\s*(['\"])([0-9]+)\1", opening))
        require(len(matches) == 1 and int(matches[0][2]) == int(attrs["s"]),
                "Unsupported cell style spelling")
        found = matches[0]
        edits.append((begin + found.start(2), begin + found.end(2), str(style).encode("ascii")))

    parser.StartElementHandler = start
    parser.Parse(raw, True)
    require(seen == set(coordinates), "Missing comparison-view cell")
    chunks, cursor = [], 0
    for start, end, replacement in edits:
        require(start >= cursor, "Overlapping cell style edits")
        chunks.extend((raw[cursor:start], replacement))
        cursor = end
    chunks.append(raw[cursor:])
    return b"".join(chunks)


def comparison_view(original, exported):
    """Return view bytes plus an audit description, without running any engine.

    The caller must additionally bind the unchanged-computation UNO proof.
    This function alone is not evidence of successful recalculation.
    """
    original, exported = safe_path(original), safe_path(exported)
    _preflight(original)
    _preflight(exported)
    raw = exported.read_bytes()
    require(len(raw) <= v5.MAX_BYTES, "Comparison view size limit")
    books = []
    try:
        before = openpyxl.load_workbook(original, data_only=False, keep_links=False)
        books.append(before)
        after = openpyxl.load_workbook(exported, data_only=False, keep_links=False)
        books.append(after)
        proof = {"version": VIEW_VERSION, "source_sha256": v5.sha(original),
                 "export_sha256": v5.sha(exported), "adapted_cells": []}
        if before.sheetnames != after.sheetnames or before.epoch != after.epoch:
            return raw, proof
        derived = v5._array_outputs(before)
        eligible = []
        for sheet in before:
            for cell in sheet._cells.values():
                changed = after[sheet.title][cell.coordinate]
                if (type(cell.value) in (int, float) and not cell.is_date
                        and cell._style is not None and cell._style.numFmtId == 30
                        and (sheet.title, cell.coordinate) not in derived
                        and changed.data_type == "d" and changed.is_date):
                    eligible.append((sheet.title, cell.coordinate))
        if not eligible or _observes_format(before) or _observes_format(after):
            return raw, proof
        with compat._safe_archive(original) as left, compat._safe_archive(exported) as right:
            _general_style(left)  # Also reject reserved custom-format overrides in input.
            general = _general_style(right)
            original_numbers, exported_numbers = _numeric_cells(left), _numeric_cells(right)
            parts, changes = v5._sheet_parts(right), {}
            for identity in eligible:
                require(identity in original_numbers and identity in exported_numbers
                        and original_numbers[identity] == exported_numbers[identity],
                        "Format30 literal binary64 changed")
                sheet, coordinate = identity
                require(float(before[sheet][coordinate].value).hex() == original_numbers[identity],
                        "Original scalar does not match raw binary64")
                changes.setdefault(parts[sheet], set()).add(coordinate)
                proof["adapted_cells"].append({"sheet": sheet, "coordinate": coordinate,
                    "binary64": original_numbers[identity], "original_numFmtId": 30,
                    "comparison_style_index": general})
            target = io.BytesIO()
            with zipfile.ZipFile(target, "w") as dest:
                for info in right.infolist():
                    content = right.read(info.filename)
                    if info.filename in changes:
                        content = _replace_styles(content, changes[info.filename], general)
                    dest.writestr(info, content)
            view = target.getvalue()
            require(len(view) <= v5.MAX_BYTES, "Comparison view size limit")
            proof["view_sha256"] = hashlib.sha256(view).hexdigest()
            return view, proof
    finally:
        for book in books:
            book.close()


class Recalculator(v5.Recalculator):
    """Explicit new source identity; all container isolation stays in v5."""

    @property
    def identity(self):
        identity = super().identity
        identity["version"] = VERSION
        identity["comparison_view"] = VIEW_VERSION
        identity["sources"] = {**identity["sources"], str(Path(__file__)): v5.sha(Path(__file__))}
        return identity

    def run(self, path):
        receipt = super().run(path)
        if receipt["status"] != "available":
            return receipt
        # Do not silently trust a missing/stale numeric computation proof.
        try:
            manifest = v5._numeric_manifest(path)
            encoded = json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
            require(receipt["input_sha256"] == v5.sha(path)
                    and type(receipt.get("numeric_literals_verified")) is int
                    and receipt["numeric_literals_verified"] == len(manifest["cells"])
                    and receipt.get("numeric_manifest_sha256") == hashlib.sha256(encoded).hexdigest(),
                    "Missing numeric computation proof")
            raw = base64.b64decode(receipt["output_base64"], validate=True)
            require(hashlib.sha256(raw).hexdigest() == receipt["output_sha256"], "Export hash mismatch")
            with tempfile.TemporaryDirectory(prefix="sheet-recalc-v6-view-") as directory:
                exported = Path(directory).resolve() / "unchanged-export.xlsx"
                exported.write_bytes(raw)
                view, proof = comparison_view(path, exported)
            result = {key: value for key, value in receipt.items() if key != "record_hash"}
            result.update(engine_export_base64=receipt["output_base64"],
                          engine_export_sha256=receipt["output_sha256"], comparison_view=proof,
                          output_base64=base64.b64encode(view).decode(),
                          output_sha256=hashlib.sha256(view).hexdigest())
            return seal(result)
        except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError,
                expat.ExpatError, zipfile.BadZipFile):
            return seal({**{key: value for key, value in receipt.items() if key != "record_hash"},
                         "status": "unknown", "reason": "format30_comparison_view_unverified"})
