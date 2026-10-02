"""Opt-in frozen benchmark-cache H, candidate-only qualified recalculation.

This changes the reference preprocessing policy, not the task or candidate.
H is read only by the host; it is never mounted in the candidate calculator.
Its cache is a dataset label, not independently established fresh computation.
Known reference disagreements remain quarantined. No learning authorization.
"""
from __future__ import annotations

import base64
import datetime
import hashlib
import io
import math
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from contextlib import ExitStack
from pathlib import Path

import openpyxl

from skillopt.coevolution_v5.core import seal
from skillopt.validator_pilot.api import digest

from . import sheet_recalc as recalc
from . import sheet_recalc_v6 as views
from . import spreadsheet_compat as compat
from .core import read_json, require, write_json

VERSION = "spreadsheet-frozen-dataset-cache-h-v3"
METADATA_PROFILE = "xldapr-single-cell-saved-cache-v1"
METADATA_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"
METADATA_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata"
METADATA_EXT = "{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}"
DYNAMIC_NS = "http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray"
UNRESOLVED_ERRORS = {"#GETTING_DATA", "#SPILL!", "#NAME?"}
SCORE_EXPECTED = {
    "numeric_correct": "pass", "numeric_wrong": "fail", "stale_correct_cache": "fail",
    "empty_string": "pass", "missing_formula_cache": "unknown", "missing_array_follower": "unknown",
    "array_zero_false": "pass", "reference_unsupported_but_cached": "pass",
    "known_reference_dispute": "unknown", "reference_error_literal": "pass", "true_date": "pass",
    "candidate_unsupported": "unknown", "missing_target_sheet": "unknown",
    "candidate_boolean_literal_limit": "unknown",
    "metadata_single_cached": "pass", "metadata_missing_cache": "unknown",
    "metadata_multicell": "unknown", "metadata_unknown_type": "unknown",
}
SCORE_CONTROLS = frozenset(SCORE_EXPECTED)


def _typed(value):
    if isinstance(value, (datetime.date, datetime.time, datetime.datetime)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    require(value is None or isinstance(value, (str, bool, int, float)), "Unsupported cached value")
    if isinstance(value, float):
        require(math.isfinite(value), "Nonfinite reference value")
    return {"type": type(value).__name__, "value": value}


def _metadata_profile(archive):
    """Qualify one observed XLDAPR shape; never interpret arbitrary metadata.

    MS-XLSX CT_DynamicArrayProperties describes fDynamic as a dynamic-array
    marker, not evidence of a current calculation. MS-OE376 2.1.624 specifies
    Office's one-based cm index (unlike generic SDK zero-based prose). This
    intentionally supports only the exact single-record cm=1 representation.
    Sources:
    learn.microsoft.com/en-us/openspecs/office_standards/ms-xlsx/d1ef676d-d970-4ddd-b642-62f5e0c291e0
    learn.microsoft.com/en-us/openspecs/office_standards/ms-oe376/b3e61d30-f15f-4712-9252-50883cbf0f8c
    learn.microsoft.com/en-us/office/open-xml/spreadsheet/working-with-formulas
    """
    relationships = []
    for member in archive.namelist():
        if member.endswith(".rels"):
            relationship_root = recalc._xml(archive, member)
            for rel in relationship_root:
                if rel.get("Type", "").endswith("/sheetMetadata"):
                    require(relationship_root.tag == "{http://schemas.openxmlformats.org/package/2006/relationships}Relationships"
                            and len({r.get("Id") for r in relationship_root}) == len(relationship_root),
                            "Metadata relationship identity ambiguous")
                    relationships.append((member, rel))
    types = recalc._xml(archive, "[Content_Types].xml")
    require(types.tag == "{http://schemas.openxmlformats.org/package/2006/content-types}Types",
            "Content type root changed")
    declarations = [node for node in types if node.get("ContentType") == METADATA_MIME
                    or node.get("PartName") == "/xl/metadata.xml"]
    if "xl/metadata.xml" not in archive.namelist() and not relationships and not declarations:
        return None
    require("xl/metadata.xml" in archive.namelist() and len(relationships) == len(declarations) == 1,
            "Unbound or duplicate metadata part")
    member, relation = relationships[0]
    require(member == "xl/_rels/workbook.xml.rels"
            and relation.tag == "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
            and relation.get("Type") == METADATA_REL and relation.get("Id")
            and relation.get("Target") in {"metadata.xml", "/xl/metadata.xml"}
            and set(relation.attrib) <= {"Id", "Type", "Target", "TargetMode"}
            and relation.get("TargetMode", "Internal") == "Internal" and not len(relation)
            and not (relation.text or "").strip() and not (relation.tail or "").strip(),
            "Unqualified metadata relationship")
    require(declarations[0].tag == "{http://schemas.openxmlformats.org/package/2006/content-types}Override"
            and declarations[0].attrib == {"PartName": "/xl/metadata.xml", "ContentType": METADATA_MIME}
            and not len(declarations[0]) and not (declarations[0].text or "").strip()
            and not (declarations[0].tail or "").strip(),
            "Unqualified metadata content type")
    root = recalc._xml(archive, "xl/metadata.xml")
    ns = "{" + compat.MAIN_NS + "}"

    def shape(node, tag, attributes, children=()):
        require(node.tag == tag and node.attrib == attributes
                and [child.tag for child in node] == list(children)
                and not (node.text or "").strip() and not (node.tail or "").strip(),
                "Unqualified XLDAPR metadata shape")

    shape(root, ns + "metadata", {}, [ns + "metadataTypes", ns + "futureMetadata", ns + "cellMetadata"])
    definitions, future, cells = root
    shape(definitions, ns + "metadataTypes", {"count": "1"}, [ns + "metadataType"])
    attributes = {name: "1" for name in ("copy", "pasteAll", "pasteValues", "merge", "splitFirst",
                  "rowColShift", "clearFormats", "clearComments", "assign", "coerce", "cellMeta")}
    shape(definitions[0], ns + "metadataType", {**attributes, "name": "XLDAPR", "minSupportedVersion": "120000"})
    shape(future, ns + "futureMetadata", {"name": "XLDAPR", "count": "1"}, [ns + "bk"])
    shape(future[0], ns + "bk", {}, [ns + "extLst"])
    shape(future[0][0], ns + "extLst", {}, [ns + "ext"])
    shape(future[0][0][0], ns + "ext", {"uri": METADATA_EXT}, ["{" + DYNAMIC_NS + "}dynamicArrayProperties"])
    shape(future[0][0][0][0], "{" + DYNAMIC_NS + "}dynamicArrayProperties", {"fDynamic": "1", "fCollapsed": "0"})
    shape(cells, ns + "cellMetadata", {"count": "1"}, [ns + "bk"])
    shape(cells[0], ns + "bk", {}, [ns + "rc"])
    shape(cells[0][0], ns + "rc", {"t": "1", "v": "0"})
    return {"profile": METADATA_PROFILE, "part_sha256": hashlib.sha256(archive.read("xl/metadata.xml")).hexdigest(),
            "dynamic_semantics_executed": False, "saved_target_cache_only": True}


def _single_cell_metadata_cache(cell, formula, coordinate):
    """No spill inference: explicit one-cell shape and an explicitly saved v."""
    require(cell.get("cm") == "1" and formula is not None and formula.get("t") == "array"
            and formula.text and formula.text.strip() and not len(formula), "Unqualified dynamic array anchor")
    left, right = compat._legacy()._range_bounds(formula.get("ref", ""))
    require(left == right == compat._legacy()._parse_cell(coordinate), "Dynamic saved range is not a single cell")
    values = cell.findall(f"{{{compat.MAIN_NS}}}v")
    require(len(values) == 1 and not len(values[0]), "Dynamic saved value missing or malformed")
    text, kind = values[0].text, cell.get("t", "n")
    require(kind in {"n", "b", "str", "e"}, "Unqualified dynamic cache type")
    if kind == "n":
        require(text and recalc.NUMERIC_TEXT.fullmatch(text) and math.isfinite(float(text)), "Dynamic numeric cache missing")
    elif kind == "b":
        require(text in {"0", "1"}, "Dynamic boolean cache missing")
    elif kind == "e":
        require(text and text.strip(), "Dynamic error cache missing")
    # t=str with an actual empty v is an explicit saved empty string, not a
    # missing result. Its freshness still is not established by this reader.


def freeze_reference(path, expected_sha256, answer_position, *, provenance, dispute=False):
    """Bind static target-cache completeness to a verified source declaration.

    The caller verifies provenance against the original dataset inventory.
    This function does not authenticate a self-authored provenance declaration.
    No target values are returned, only their digest and availability counts.
    """
    require(isinstance(expected_sha256, str) and len(expected_sha256) == 64
            and recalc.sha(path) == expected_sha256, "Reference bytes changed")
    require(type(provenance) is dict and provenance.get("source_protocol_hash")
            and provenance.get("task_id") and provenance.get("evidence_kind") in
            {"engineering_fixture", "frozen_real_model_replay"}, "Reference provenance required")
    require(type(dispute) is bool, "Explicit reference dispute flag required")
    base = {"version": VERSION, "reference_sha256": expected_sha256,
            "answer_position": answer_position, "provenance": provenance,
            "known_reference_dispute": dispute, "host_only": True,
            "reference_recalculated": False, "cache_freshness_established": False,
            "truth_basis": "frozen_benchmark_label_not_independent_oracle",
            "model_api_calls": 0}

    def done(status, reason, **kwargs):
        return seal({**base, "status": status, "reason": reason, **kwargs})

    try:
        views._preflight(path)
        raw_cells = {}
        with compat._safe_archive(path) as archive:
            metadata = _metadata_profile(archive)
            for sheet, member in recalc._sheet_parts(archive).items():
                cells, shared = {}, {}
                for cell in recalc._xml(archive, member).iter(f"{{{compat.MAIN_NS}}}c"):
                    coordinate = cell.get("r")
                    require(coordinate and coordinate.upper() not in cells, "Duplicate or absent cell identity")
                    compat._legacy()._parse_cell(coordinate)
                    if "vm" in cell.attrib or ("cm" in cell.attrib and metadata is None):
                        return done("unknown", "reference_cell_metadata_unqualified")
                    require(len(cell.findall(f"{{{compat.MAIN_NS}}}f")) <= 1
                            and len(cell.findall(f"{{{compat.MAIN_NS}}}v")) <= 1,
                            "Duplicate formula or value")
                    formula = cell.find(f"{{{compat.MAIN_NS}}}f")
                    if "cm" in cell.attrib:
                        _single_cell_metadata_cache(cell, formula, coordinate)
                    if formula is not None and formula.get("t") == "dataTable":
                        return done("unknown", "reference_data_table_unqualified")
                    if formula is not None:
                        kind = formula.get("t", "normal")
                        require(kind in {"normal", "array", "shared"}, "Unknown formula type")
                        if kind == "shared":
                            identity = formula.get("si", "")
                            require(identity.isdecimal() and int(identity) >= 0, "Invalid shared formula identity")
                            shared.setdefault(identity, []).append((coordinate.upper(), formula))
                        else:
                            require(bool(formula.text and formula.text.strip()), "Missing formula expression")
                    cells[coordinate.upper()] = cell
                for members in shared.values():
                    anchors = [(coordinate, formula) for coordinate, formula in members
                               if formula.text and formula.text.strip()]
                    require(len(anchors) == 1, "Shared formula needs one anchor")
                    anchor, formula = anchors[0]
                    (sc, sr), (ec, er) = compat._legacy()._range_bounds(formula.get("ref", ""))
                    require((ec - sc + 1) * (er - sr + 1) <= compat.MAX_TARGET_CELLS,
                            "Shared formula range budget")
                    for coordinate, member_formula in members:
                        col, row = compat._legacy()._parse_cell(coordinate)
                        require(sc <= col <= ec and sr <= row <= er, "Shared formula outside anchor range")
                        require(coordinate == anchor or "ref" not in member_formula.attrib,
                                "Follower has a conflicting shared range")
                raw_cells[sheet] = cells
        with ExitStack() as stack:
            values = openpyxl.load_workbook(path, data_only=True, keep_links=False)
            stack.callback(values.close)
            formulas = openpyxl.load_workbook(path, data_only=False, keep_links=False)
            stack.callback(formulas.close)
            resolution = compat.resolve_targets(answer_position, values.sheetnames)
            require(values.sheetnames == formulas.sheetnames, "Reference sheet identity mismatch")
            derived = recalc._array_outputs(formulas)
            counts, target_values, missing = Counter(), [], []
            for sheet, region in resolution["targets"]:
                for coordinate in compat._legacy()._iter_cell_names(region):
                    cell = raw_cells.get(sheet, {}).get(coordinate)
                    raw_f = None if cell is None else cell.find(f"{{{compat.MAIN_NS}}}f")
                    raw_v = None if cell is None else cell.find(f"{{{compat.MAIN_NS}}}v")
                    calculated = (raw_f is not None or (sheet, coordinate) in derived
                                  or formulas[sheet][coordinate].data_type == "f")
                    value = values[sheet][coordinate].value
                    empty_string = (calculated and cell is not None and cell.get("t") == "str"
                                    and raw_v is not None and len(raw_v) == 0 and raw_v.text in (None, ""))
                    if metadata and (cell is None or (value is None and not empty_string)):
                        # Even an unmarked target cannot be inferred blank in
                        # a dynamic workbook from an absent/empty numeric node.
                        counts["missing_explicit_metadata_target"] += 1
                        missing.append([sheet, coordinate])
                    if calculated and (raw_v is None or (value is None and not empty_string)):
                        counts["missing_calculated_cache"] += 1
                        missing.append([sheet, coordinate])
                    elif empty_string:
                        value = ""
                        counts["explicit_empty_calculated_string"] += 1
                    elif calculated:
                        counts["available_calculated_cache"] += 1
                    else:
                        counts["ordinary_blank" if value is None else "literal"] += 1
                    if values[sheet][coordinate].data_type == "e" and value in UNRESOLVED_ERRORS:
                        counts["unresolved_error"] += 1
                    target_values.append([sheet, coordinate, _typed(value)])
            information = {"target_resolution": resolution["mode"], "target_cells": len(target_values),
                           "cache_counts": dict(counts), "target_values_sha256": digest(target_values),
                           "missing_cache_locations_sha256": digest(missing)}
            if metadata:
                information["metadata_snapshot"] = metadata
            if dispute:
                return done("unknown", "known_reference_dispute", **information)
            if missing or counts["unresolved_error"]:
                return done("unknown", "reference_target_cache_incomplete", **information)
            return done("available", "frozen_label_cache_complete", **information)
    except (OSError, ValueError, TypeError, KeyError, IndexError, AttributeError,
            ET.ParseError, zipfile.BadZipFile) as exc:
        return done("unknown", "reference_static_evidence_error:" + type(exc).__name__)


def evaluate(prediction, reference, answer_position, engine, manifest):
    """Use a frozen host-only H and a qualified candidate calculator.

    Receipt identity, bytes, cleanup and original content checks are mandatory.
    A complete H is required even when another target would already mismatch.
    The caller qualifies the calculator and freezes this policy before use.
    """
    require(type(manifest) is dict and manifest == seal({k: v for k, v in manifest.items()
                                                       if k != "record_hash"}), "Reference manifest seal mismatch")
    verified = freeze_reference(reference, manifest["reference_sha256"], answer_position,
                                provenance=manifest["provenance"], dispute=manifest["known_reference_dispute"])
    require(manifest == verified, "Reference policy, target or evidence changed")
    result = {"reference_manifest_hash": manifest["record_hash"], "receipts": {},
              "version": VERSION, "reference_executed": False, "model_api_calls": 0,
              "feedback_allowed": False}
    if manifest["status"] != "available":
        return {**result, "status": "unknown", "reason": manifest["reason"]}
    input_hash = recalc.sha(prediction)
    receipt = engine.run(prediction)
    require(receipt == seal({k: v for k, v in receipt.items() if k != "record_hash"})
            and receipt["input_sha256"] == input_hash and receipt["identity"] == engine.identity,
            "Candidate receipt identity mismatch")
    require(recalc.sha(prediction) == input_hash, "Candidate changed during recalculation")
    result["receipts"]["prediction"] = receipt["record_hash"]
    if receipt.get("cleanup_confirmed") is not True:
        return {**result, "status": "unknown", "reason": "candidate_cleanup_unconfirmed"}
    if receipt["status"] != "available":
        return {**result, "status": "unknown", "reason": "prediction:" + receipt["reason"]}
    raw = base64.b64decode(receipt["output_base64"], validate=True)
    require(len(raw) <= recalc.MAX_BYTES and hashlib.sha256(raw).hexdigest() == receipt["output_sha256"],
            "Candidate receipt output mismatch")
    with tempfile.TemporaryDirectory(prefix="candidate-only-score-") as temp:
        fresh = Path(temp).resolve() / "candidate.xlsx"
        fresh.write_bytes(raw)
        reason = recalc._recalc_valid(prediction, fresh)
        if reason:
            return {**result, "status": "unknown", "reason": "prediction:" + reason}
        scored = compat.evaluate(fresh, reference, "", answer_position)
        return {**result, "status": {"passed": "pass", "failed": "fail", "missing_output": "fail"}
                .get(scored["status"], "unknown"), "reason": scored["reason"], "evidence": scored["evidence"]}


def _fixture_book(path, cells, caches=None):
    """Author labelled engineering fixtures only; never used on real artifacts."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "S"
    for coordinate, value in cells.items():
        sheet[coordinate] = value
    buffer = io.BytesIO()
    book.save(buffer)
    book.close()
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive, zipfile.ZipFile(output, "w") as target:
        for item in archive.infolist():
            raw = archive.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml" and caches:
                tree = ET.fromstring(raw)
                for cell in tree.iter(f"{{{compat.MAIN_NS}}}c"):
                    if cell.get("r") not in caches:
                        continue
                    kind, value = caches[cell.get("r")]
                    cell.set("t", kind)
                    old = cell.find(f"{{{compat.MAIN_NS}}}v")
                    if old is not None:
                        cell.remove(old)
                    ET.SubElement(cell, f"{{{compat.MAIN_NS}}}v").text = value
                raw = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            target.writestr(item, raw)
    Path(path).write_bytes(output.getvalue())


def _fixture_xldapr(path, *, unknown_type=False):
    """Add exact metadata only to authored qualification fixtures, never H."""
    ns, relationships, types = ("{" + name + "}" for name in (
        compat.MAIN_NS, "http://schemas.openxmlformats.org/package/2006/relationships",
        "http://schemas.openxmlformats.org/package/2006/content-types"))
    metadata = ET.Element(ns + "metadata")
    definitions = ET.SubElement(metadata, ns + "metadataTypes", count="1")
    attrs = {name: "1" for name in ("copy", "pasteAll", "pasteValues", "merge", "splitFirst",
             "rowColShift", "clearFormats", "clearComments", "assign", "coerce", "cellMeta")}
    ET.SubElement(definitions, ns + "metadataType", {**attrs,
        "name": "UNQUALIFIED" if unknown_type else "XLDAPR", "minSupportedVersion": "120000"})
    future = ET.SubElement(metadata, ns + "futureMetadata", name="XLDAPR", count="1")
    extension = ET.SubElement(ET.SubElement(ET.SubElement(future, ns + "bk"), ns + "extLst"), ns + "ext", uri=METADATA_EXT)
    ET.SubElement(extension, "{" + DYNAMIC_NS + "}dynamicArrayProperties", fDynamic="1", fCollapsed="0")
    ET.SubElement(ET.SubElement(ET.SubElement(metadata, ns + "cellMetadata", count="1"), ns + "bk"),
                  ns + "rc", t="1", v="0")
    output = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(output, "w") as target:
        require("xl/metadata.xml" not in source.namelist(), "Fixture metadata already authored")
        for item in source.infolist():
            raw = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                tree = ET.fromstring(raw)
                cells = [c for c in tree.iter(ns + "c") if c.get("r") == "B1"]
                require(len(cells) == 1, "Fixture dynamic anchor missing")
                cells[0].set("cm", "1")
                raw = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            elif item.filename == "xl/_rels/workbook.xml.rels":
                tree = ET.fromstring(raw)
                ET.SubElement(tree, relationships + "Relationship", Id="rIdXldaprFixture",
                              Type=METADATA_REL, Target="metadata.xml")
                raw = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            elif item.filename == "[Content_Types].xml":
                tree = ET.fromstring(raw)
                ET.SubElement(tree, types + "Override", PartName="/xl/metadata.xml", ContentType=METADATA_MIME)
                raw = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            target.writestr(item, raw)
        target.writestr("xl/metadata.xml", ET.tostring(metadata, encoding="utf-8", xml_declaration=True))
    Path(path).write_bytes(output.getvalue())


def qualify(output, engine):
    """Eighteen predeclared score controls using the actual candidate engine.

    The surrounding CLI binds source/engine before this call. Fixture authoring
    caches are labelled controls, never real benchmark answer modifications.
    DurableEngine ensures replay reuses each closed candidate execution.
    """
    root = Path(output).resolve()
    root.mkdir(parents=True, exist_ok=True)
    manifest_path = root / "fixtures.json"
    if not manifest_path.exists():
        require(not list(root.iterdir()), "Incomplete score fixture setup; preserve and use a new directory")
        from openpyxl.worksheet.formula import ArrayFormula
        controls = []
        for name, expected in SCORE_EXPECTED.items():
            gold_cells, pred_cells = {"B1": 2}, {"B1": 2}
            gold_caches, pred_caches = None, None
            target = "S!B1"
            if name == "numeric_wrong":
                pred_cells = {"B1": 3}
            elif name == "stale_correct_cache":
                pred_cells, pred_caches = {"A1": 3, "B1": "=A1+1"}, {"B1": ("n", "2")}
            elif name == "empty_string":
                gold_cells = pred_cells = {"B1": '=IF(TRUE(),"",2)'}
                gold_caches = {"B1": ("str", "")}
            elif name == "missing_formula_cache":
                gold_cells = {"B1": "=1+1", "B2": 999}
                pred_cells, target = {"B1": 2, "B2": 0}, "S!B1:B2"
            elif name in {"missing_array_follower", "array_zero_false"}:
                gold_cells = {"A1": 0, "A2": 1, "B1": ArrayFormula(ref="B1:B2", text="=IF(A1:A2=0,0,FALSE())"), "B2": 0}
                gold_caches = {"B1": ("n", "0"), "B2": ("n", "") if name == "missing_array_follower" else ("b", "0")}
                # H's False/zero completeness is independent of whether the
                # calculator preserves a candidate boolean *literal*. Calc
                # exports that literal as a formula, which v8b rejects. Keep
                # this limitation in a separate explicit rejection control.
                pred_cells, target = {"B1": "=0", "B2": "=FALSE()"}, "S!B1:B2"
            elif name == "reference_unsupported_but_cached":
                gold_cells, gold_caches = {"A1": 2, "B1": "=_xlfn.UNIQUE(A1)"}, {"B1": ("n", "2")}
            elif name == "reference_error_literal":
                gold_cells, pred_cells = {"B1": "#N/A"}, {"B1": "=NA()"}
            elif name == "true_date":
                gold_cells = pred_cells = {"B1": datetime.datetime(2020, 1, 2)}
            elif name == "candidate_unsupported":
                pred_cells = {"A1": 2, "B1": "=_xlfn.UNIQUE(A1)"}
            elif name == "missing_target_sheet":
                target = "Absent!B1"
            elif name == "candidate_boolean_literal_limit":
                gold_cells = pred_cells = {"B1": 0, "B2": False}
                target = "S!B1:B2"
            elif name.startswith("metadata_"):
                gold_cells = {"B1": ArrayFormula(ref="B1:B2" if name == "metadata_multicell" else "B1", text="=1+1")}
                gold_caches = None if name == "metadata_missing_cache" else {"B1": ("n", "2")}
                if name == "metadata_multicell":
                    gold_cells["B2"] = 2
                    gold_caches["B2"] = ("n", "2")
                    pred_cells["B2"] = 2
                    target = "S!B1:B2"
            gold, pred = root / (name + "-gold.xlsx"), root / (name + "-candidate.xlsx")
            _fixture_book(gold, gold_cells, gold_caches)
            _fixture_book(pred, pred_cells, pred_caches)
            if name.startswith("metadata_"):
                _fixture_xldapr(gold, unknown_type=name == "metadata_unknown_type")
            controls.append({"name": name, "expected": expected, "reference": str(gold), "prediction": str(pred),
                             "reference_sha256": recalc.sha(gold), "prediction_sha256": recalc.sha(pred),
                             "answer_position": target, "dispute": name == "known_reference_dispute"})
        write_json(manifest_path, seal({"version": VERSION, "controls": controls,
                                      "evidence_kind": "engineering_fixture"}))
    fixtures = read_json(manifest_path, sealed=True)
    require(fixtures["version"] == VERSION and len(fixtures["controls"]) == len(SCORE_CONTROLS)
            and {c["name"]: c["expected"] for c in fixtures["controls"]} == SCORE_EXPECTED,
            "Score control roster or expectations changed")
    rows = []
    for control in fixtures["controls"]:
        require(recalc.sha(control["reference"]) == control["reference_sha256"]
                and recalc.sha(control["prediction"]) == control["prediction_sha256"], "Score fixture bytes changed")
        h = freeze_reference(control["reference"], control["reference_sha256"], control["answer_position"],
                             provenance={"source_protocol_hash": fixtures["record_hash"],
                                         "task_id": control["name"], "evidence_kind": "engineering_fixture"},
                             dispute=control["dispute"])
        observed = evaluate(control["prediction"], control["reference"], control["answer_position"], engine, h)
        rows.append({"name": control["name"], "expected": control["expected"], "result": observed,
                     "reference_manifest": h, "qualified": observed["status"] == control["expected"]})
        if hasattr(engine, "unsafe_cleanup") and engine.unsafe_cleanup():
            break
    return seal({"version": VERSION, "status": "qualified" if len(rows) == len(SCORE_CONTROLS)
                 and all(c["qualified"] for c in rows) else "rejected", "engine": engine.identity,
                 "module_sha256": recalc.sha(__file__), "fixtures_hash": fixtures["record_hash"], "controls": rows,
                 "evidence_kind": "engineering_fixture", "model_api_calls": 0, "feedback_allowed": False})
