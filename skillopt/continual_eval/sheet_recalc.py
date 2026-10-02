"""Zero-model-call, opt-in LibreOffice rescore of immutable delivered workbooks.

This is a new scoring protocol, not an update to historical scores or an Excel
equivalence claim. Reference/prediction copies are recalculated independently;
reference cache drift and unsupported semantics remain unknown. Originals and
the complete position denominator are retained. No model Python is reexecuted.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import math
import platform
import posixpath
import re
import shutil
import subprocess
import tempfile
import time
import uuid
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path
from xml.parsers import expat

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.worksheet.formula import ArrayFormula

from skillopt.coevolution_v5.core import seal
from skillopt.skill_validation.sandbox import _bounded_command, _Command
from skillopt.validator_pilot.api import digest

from . import spreadsheet_compat as compat
from .core import output_lock, panel_tasks, read_json, require, safe_path, write_json

VERSION = "spreadsheet-independent-lo-recalculation-v5"
MAX_BYTES = 8 * 1024 * 1024
MAX_NUMERIC_LITERALS = 100_000
NUMERIC_TEXT = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?\Z")
IMAGE = re.compile(r"sha256:[0-9a-f]{64}\Z")
FORBIDDEN_FUNCTIONS = re.compile(
    r"(?:^|[^A-Z0-9_.])(?:WEBSERVICE|HYPERLINK|RTD|DDE|CALL|REGISTER\.ID|INDIRECT|"
    r"NOW|TODAY|RAND|RANDBETWEEN|RANDARRAY)\s*\(", re.I)


def sha(path):
    return hashlib.sha256(safe_path(path).read_bytes()).hexdigest()


def _sources():
    paths = [Path(__file__), Path(__file__).with_name("sheet_recalc_worker.py"),
             Path(compat.__file__), Path(compat._legacy().__file__),
             Path(__file__).parents[1] / "skill_validation/sandbox.py"]
    return {str(path): sha(path) for path in paths}


def _xml(archive, member):
    """Version-local XML reader: explicit UTF-16 BOM, never rewrite the member.

    Preserve the original UTF-8 reader and all budgets. Check declarations on
    decoded Unicode, so UTF-16 cannot hide DTD/entities or external links from
    the same downstream semantic screening. Unsupported encodings stay unknown.
    """
    if archive.getinfo(member).file_size > compat.MAX_XML_BYTES:
        raise compat.CompatibilityError("xml_part_size_limit")
    raw = archive.read(member)
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        raise compat.CompatibilityError("xml_utf32_unsupported")
    if not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return compat._xml(archive, member)
    try:
        text = raw.decode("utf-16", errors="strict")
    except UnicodeError as exc:
        raise compat.CompatibilityError("xml_invalid_utf16") from exc
    if text.startswith("\ufeff"):
        raise compat.CompatibilityError("xml_duplicate_byte_order_mark")
    if len(text.encode("utf-8")) > compat.MAX_XML_BYTES:
        raise compat.CompatibilityError("xml_decoded_size_limit")
    if "\x00" in text or "<!DOCTYPE" in text.upper() or "<!ENTITY" in text.upper():
        raise compat.CompatibilityError("xml_document_type_or_entity_forbidden")
    declaration = re.match(r"\A<\?xml\s+.*?\?>", text, flags=re.S)
    if declaration:
        encoding = re.search(r"\bencoding\s*=\s*(['\"])([^'\"]+)\1", declaration.group())
        if encoding:
            # Only XML/IANA labels, not arbitrary Python codec aliases.
            allowed = {"utf-16", "utf-16le" if raw.startswith(b"\xff\xfe") else "utf-16be"}
            if encoding.group(2).lower() not in allowed:
                raise compat.CompatibilityError("xml_encoding_bom_mismatch")
    return ET.fromstring(text)


def _formula_spec(value):
    if isinstance(value, str):
        return ("scalar", value, None)
    if isinstance(value, ArrayFormula) and isinstance(value.text, str) and isinstance(value.ref, str):
        (sc, sr), (ec, er) = compat._legacy()._range_bounds(value.ref)
        require((ec - sc + 1) * (er - sr + 1) <= compat.MAX_TARGET_CELLS, "Array formula budget")
        # Excel/LibreOffice may spell the same fixed single-cell result range
        # B1 or B1:B1. Compare its parsed coordinates, never normalize formula
        # text or permit a changed extent. This does not claim general formula
        # equivalence and must not change v1's frozen rejection.
        canonical_range = (f"{compat._legacy()._col_num2name(sc)}{sr}:"
                           f"{compat._legacy()._col_num2name(ec)}{er}")
        return ("array", value.text, canonical_range)
    raise ValueError("Unsupported formula representation")


def _boolean_call_normal_form(formula):
    """One qualified spelling equivalence, not general formula normalization.

    Excel documents TRUE/FALSE literals and their zero-argument functions as
    the same logical values (support.microsoft.com/en-us/excel/functions/true-function
    and /false-function). Tokenization preserves quoted text and references.
    Only unqualified uppercase empty calls in operand positions are reduced.
    Array constants, name/reference positions and other edits remain exact.
    """
    try:
        lexer = Tokenizer(formula)
        if not formula.startswith("=") or lexer.render() != formula:
            return None
        items, stack, result, index = lexer.items, [], [], 0
        infix = {"+", "-", "*", "/", "^", "&", "=", "<", ">", "<=", ">=", "<>"}

        def adjacent(start, step):
            while 0 <= start < len(items) and items[start].type == "WHITE-SPACE":
                start += step
            return items[start] if 0 <= start < len(items) else None

        def operand_boundary(token, before):
            if token is None:
                return True
            if token.type == "OPERATOR-INFIX" and token.value in infix:
                return True
            if token.type == "SEP" and token.subtype == "ARG":
                return bool(stack) and stack[-1] == "FUNC"
            if before:
                return ((token.type in {"FUNC", "PAREN"} and token.subtype == "OPEN")
                        or token.type == "OPERATOR-PREFIX" and token.value in {"+", "-"})
            return ((token.type in {"FUNC", "PAREN"} and token.subtype == "CLOSE")
                    or token.type == "OPERATOR-POSTFIX" and token.value == "%")

        while index < len(items):
            token = items[index]
            if (token.type == "FUNC" and token.subtype == "OPEN" and token.value in {"TRUE(", "FALSE("}
                    and index + 1 < len(items) and items[index + 1].type == "FUNC"
                    and items[index + 1].subtype == "CLOSE" and items[index + 1].value == ")"
                    and "ARRAY" not in stack
                    and operand_boundary(adjacent(index - 1, -1), True)
                    and operand_boundary(adjacent(index + 2, 1), False)):
                result.append(token.value[:-1])
                index += 2
                continue
            if token.type in {"FUNC", "PAREN", "ARRAY"}:
                if token.subtype == "OPEN":
                    stack.append(token.type)
                elif token.subtype == "CLOSE":
                    if not stack or stack.pop() != token.type:
                        return None
            result.append(token.value)
            index += 1
        return "=" + "".join(result) if not stack else None
    except (TokenizerError, IndexError, ValueError):
        return None


def _same_formula(before, after):
    left, right = _formula_spec(before), _formula_spec(after)
    if left == right:
        return True
    if left[0] != right[0] or left[2] != right[2]:
        return False
    normalized = _boolean_call_normal_form(left[1])
    return normalized is not None and normalized == _boolean_call_normal_form(right[1])


def _array_outputs(book):
    derived = set()
    for sheet in book:
        # Iterate represented cells, not the rectangular extent of a sparse sheet.
        for cell in sheet._cells.values():
            if cell.data_type == "f" and isinstance(cell.value, ArrayFormula):
                _, _, region = _formula_spec(cell.value)
                (sc, sr), _ = compat._legacy()._range_bounds(region)
                require((cell.column, cell.row) == (sc, sr), "Array anchor does not match its range")
                for coordinate in compat._legacy()._iter_cell_names(region):
                    identity = (sheet.title, coordinate)
                    require(identity not in derived, "Overlapping array regions")
                    derived.add(identity)
                    require(len(derived) <= compat.MAX_TARGET_CELLS, "Array output budget")
    return derived


def screen(path):
    """Fail closed on macros, external links/data, volatile or unsupported math.

    This conservative scope can reject otherwise harmless hyperlinks. It is
    intentionally reported as unknown, not a model error or task filter.
    """
    try:
        with compat._safe_archive(path) as archive:
            names = archive.namelist()
            for name in names:
                lower = name.lower()
                if any(part in lower for part in ("vba", "macrosheet", "dialogsheet", "externallink",
                                                  "connections.xml", "querytable", "embeddings/", "activex/")):
                    return "unsupported_active_or_external_workbook"
                if lower.endswith((".xml", ".rels")):
                    tree = _xml(archive, name)
                    if any(node.get("TargetMode", "").lower() == "external" for node in tree.iter()):
                        return "unsupported_external_relationship"
                    for node in tree.iter():
                        if node.tag.rsplit("}", 1)[-1] == "definedName" and node.text:
                            if "[" in node.text or "|" in node.text or FORBIDDEN_FUNCTIONS.search(node.text):
                                return "unsupported_external_or_volatile_defined_name"
                    if lower == "[content_types].xml" and any(
                        "macroenabled" in str(node.attrib).lower() for node in tree.iter()
                    ):
                        return "unsupported_macro_content_type"
        book = openpyxl.load_workbook(path, data_only=False, keep_links=False)
        try:
            _array_outputs(book)
            for sheet in book:
                if sheet.max_row * sheet.max_column > 2_000_000:
                    return "unsupported_worksheet_iteration_budget"
                for row in sheet:
                    for cell in row:
                        if cell.data_type != "f":
                            continue
                        try:
                            _, formula, _ = _formula_spec(cell.value)
                        except ValueError:
                            return "unsupported_formula_representation_or_array_range"
                        if "[" in formula or "|" in formula or FORBIDDEN_FUNCTIONS.search(formula):
                            return "unsupported_external_structured_or_volatile_formula"
                        if re.search(r"(?:_xlfn\.|_xlws\.|COM\.MICROSOFT\.|ORG\.OPENOFFICE\.)", formula, re.I):
                            return "unsupported_extended_function"
        finally:
            book.close()
        return None
    except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError, zipfile.BadZipFile):
        return "unsupported_workbook_structure"


def _sheet_parts(archive):
    rels = _xml(archive, "xl/_rels/workbook.xml.rels")
    by_id = {rel.get("Id"): rel for rel in rels}
    require(len(by_id) == len(rels), "Duplicate worksheet relationship")
    result = {}
    for sheet in _xml(archive, "xl/workbook.xml").findall(
            f"{{{compat.MAIN_NS}}}sheets/{{{compat.MAIN_NS}}}sheet"):
        name = sheet.get("name")
        rel = by_id.get(sheet.get(f"{{{compat.DOC_REL_NS}}}id"))
        require(name and name not in result and rel is not None, "Worksheet identity unavailable")
        require(rel.get("TargetMode") != "External" and rel.get("Type", "").endswith("/worksheet"),
                "Unsupported worksheet relationship")
        target = rel.get("Target", "")
        require(target and "\\" not in target and ":" not in target, "Invalid worksheet relationship")
        member = posixpath.normpath(target.lstrip("/") if target.startswith("/")
                                   else posixpath.join("xl", target))
        require(member.startswith("xl/") and member not in result.values(), "Unsafe worksheet member")
        result[name] = member
    return result


def _numeric_manifest(path):
    """Literal values only; never a formula cache, array result or other workbook.

    Exact binary64 is the computation contract. No epsilon or native answer
    scoring tolerance may be used to repair an intermediate input.
    """
    book = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    cells = []
    try:
        derived = _array_outputs(book)
        with compat._safe_archive(path) as archive:
            for sheet, member in _sheet_parts(archive).items():
                seen = set()
                for cell in _xml(archive, member).iter(f"{{{compat.MAIN_NS}}}c"):
                    coordinate = cell.get("r")
                    require(coordinate and coordinate not in seen, "Missing or duplicate cell identity")
                    seen.add(coordinate)
                    value = book[sheet][coordinate]
                    if (value.data_type == "f" or (sheet, coordinate) in derived
                            or type(value.value) not in (int, float)):
                        continue
                    nodes = cell.findall(f"{{{compat.MAIN_NS}}}v")
                    require(cell.get("t", "n") == "n" and cell.find(f"{{{compat.MAIN_NS}}}f") is None
                            and len(nodes) == 1 and nodes[0].text
                            and NUMERIC_TEXT.fullmatch(nodes[0].text), "Numeric literal unavailable")
                    number = float(nodes[0].text)
                    require(math.isfinite(number) and number == value.value,
                            "Numeric literal not exactly binary64 representable")
                    cells.append({"sheet": sheet, "coordinate": coordinate,
                                  "row": value.row - 1, "column": value.column - 1,
                                  "binary64": number.hex(), "xml_value": nodes[0].text})
                    require(len(cells) <= MAX_NUMERIC_LITERALS, "Numeric literal budget")
        return {"schema": "numeric-literal-export-proof-v1", "input_sha256": sha(path), "cells": cells}
    finally:
        book.close()


def _restore_numeric_xml(raw, entries):
    """Byte-local replacement of existing literal v text, never XML reserialization.

    Expat supplies namespace-aware element locations. Every unchanged byte,
    including namespace declarations, formulas and caches, is retained.
    """
    wanted = {item["coordinate"]: item for item in entries}
    require(len(wanted) == len(entries), "Duplicate numeric restoration identity")
    parser = expat.ParserCreate(namespace_separator="}")
    stack, edits, seen = [], [], set()
    cell_coordinate, start_value = None, None
    cell_tag, value_tag, formula_tag = [compat.MAIN_NS + "}" + tag for tag in ("c", "v", "f")]

    def start(name, attrs):
        nonlocal cell_coordinate, start_value
        stack.append(name)
        if name == cell_tag:
            require(cell_coordinate is None, "Nested cell")
            cell_coordinate = attrs.get("r")
            if cell_coordinate in wanted:
                require(cell_coordinate not in seen and attrs.get("t", "n") == "n", "Changed numeric cell type")
                seen.add(cell_coordinate)
        elif cell_coordinate in wanted:
            require(name != formula_tag, "Literal became formula")
            if name == value_tag:
                require(len(stack) >= 2 and stack[-2] == cell_tag and start_value is None,
                        "Invalid numeric value nesting")
                offset = parser.CurrentByteIndex
                end = raw.find(b">", offset)
                require(end >= offset and raw[end - 1:end] != b"/", "Missing numeric value")
                start_value = end + 1

    def end(name):
        nonlocal cell_coordinate, start_value
        if name == value_tag and cell_coordinate in wanted:
            require(start_value is not None, "Missing numeric text")
            stop = parser.CurrentByteIndex
            text = raw[start_value:stop].decode("ascii")
            require(NUMERIC_TEXT.fullmatch(text) and math.isfinite(float(text)), "Invalid exported numeric text")
            edits.append((start_value, stop, wanted[cell_coordinate]["xml_value"].encode("ascii")))
            start_value = None
        if name == cell_tag:
            cell_coordinate = None
        stack.pop()

    parser.StartElementHandler, parser.EndElementHandler = start, end
    parser.Parse(raw, True)
    require(seen == set(wanted) and len(edits) == len(wanted), "Missing or duplicate numeric value")
    parts, cursor, changed = [], 0, 0
    for start, stop, replacement in edits:
        require(start >= cursor, "Overlapping numeric edits")
        changed += raw[start:stop] != replacement
        parts.extend((raw[cursor:start], replacement))
        cursor = stop
    parts.append(raw[cursor:])
    return b"".join(parts), changed


def _restore_numeric_export(raw, manifest, proof):
    encoded = json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
    count = len(manifest["cells"])
    require(proof.get("numeric_manifest_sha256") == hashlib.sha256(encoded).hexdigest()
            and type(proof.get("numeric_verified_before")) is int
            and type(proof.get("numeric_verified_after")) is int
            and proof.get("numeric_verified_before") == count
            and proof.get("numeric_verified_after") == count,
            "Missing exact computation input proof")
    with tempfile.TemporaryDirectory(prefix="sheet-recalc-export-") as directory:
        path = Path(directory) / "raw.xlsx"
        path.write_bytes(raw)
        with compat._safe_archive(path) as archive:
            sheets, by_member = _sheet_parts(archive), {}
            for item in manifest["cells"]:
                require(item["sheet"] in sheets, "Numeric sheet missing after export")
                by_member.setdefault(sheets[item["sheet"]], []).append(item)
            output, changed = io.BytesIO(), 0
            with zipfile.ZipFile(output, "w") as dest:
                for entry in archive.infolist():
                    data = archive.read(entry.filename)
                    if entry.filename in by_member:
                        _xml(archive, entry.filename)
                        data, n = _restore_numeric_xml(data, by_member[entry.filename])
                        changed += n
                    dest.writestr(entry, data)
            require(len(output.getvalue()) <= MAX_BYTES, "Restored workbook size limit")
            return output.getvalue(), changed


class Recalculator:
    """Each workbook receives a fresh networkless, read-only-root container."""

    def __init__(self, image, timeout=120):
        require(isinstance(image, str) and IMAGE.fullmatch(image), "Pinned local image ID required")
        require(type(timeout) is int and 20 <= timeout <= 300, "Bounded recalculation timeout required")
        self.image, self.timeout = image, timeout

    @property
    def identity(self):
        return {"version": VERSION, "image_id": self.image, "timeout_seconds": self.timeout,
                "memory_mb": 1024, "cpus": 1, "pids_limit": 128, "network": "none",
                "root_read_only": True, "source_read_only": True, "user": "65534:65534",
                "macros": "never_execute", "links": "no_update", "sources": _sources()}

    def run(self, path):
        path = safe_path(path)
        started = time.monotonic()
        base = {"input_sha256": sha(path), "identity": self.identity,
                "cleanup_confirmed": True, "model_api_calls": 0, "container_execution_attempted": False}

        def done(status, reason, **extra):
            return seal({**base, "status": status, "reason": reason,
                         "duration_seconds": time.monotonic() - started, **extra})

        reason = screen(path)
        if reason:
            return done("unknown", reason)
        if platform.system() != "Linux" or shutil.which("docker") is None:
            return done("unknown", "linux_docker_unavailable")
        try:
            manifest = _numeric_manifest(path)
            manifest_bytes = json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
            require(len(manifest_bytes) <= MAX_BYTES, "Numeric manifest size limit")
        except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError, zipfile.BadZipFile):
            return done("unknown", "numeric_literal_manifest_unavailable")
        try:
            inspect = _bounded_command(["docker", "image", "inspect", "--format",
                                        "{{json .}}", self.image], 10)
        except (OSError, subprocess.TimeoutExpired):
            return done("unknown", "pinned_image_or_daemon_unavailable")
        try:
            image = json.loads(inspect.stdout)
            require(inspect.code == 0 and not inspect.timed_out and not inspect.overflow
                    and image["Id"] == self.image and image["Os"] == "linux"
                    and not image["Config"].get("Volumes"), "Unsupported image")
        except (ValueError, KeyError, TypeError):
            return done("unknown", "pinned_image_unavailable_or_unsafe")
        name = "skillopt-sheet-recalc-" + uuid.uuid4().hex
        with tempfile.TemporaryDirectory(prefix="sheet-recalc-input-") as directory:
            directory = Path(directory).resolve()
            directory.chmod(0o755)
            frozen = directory / "book.xlsx"
            frozen.write_bytes(path.read_bytes())
            frozen.chmod(0o444)
            manifest_path = directory / "numeric_literals.json"
            manifest_path.write_bytes(manifest_bytes)
            manifest_path.chmod(0o444)
            worker = safe_path(Path(__file__).with_name("sheet_recalc_worker.py"))
            command = ["docker", "run", "--name", name, "--pull=never", "--network=none", "--read-only",
                       "--user=65534:65534", "--cap-drop=ALL", "--security-opt=no-new-privileges",
                       "--pids-limit=128", "--memory=1024m", "--memory-swap=1024m", "--cpus=1",
                       "--tmpfs=/tmp:rw,nosuid,nodev,noexec,size=256m,mode=1777",
                       "--mount", f"type=bind,src={directory},dst=/input,readonly",
                       "--mount", f"type=bind,src={worker},dst=/worker.py,readonly",
                       "--env=HOME=/tmp", "--env=SAL_USE_VCLPLUGIN=svp",
                       "--env=SKILLOPT_SHEET_RECALC_CONTAINER=1", self.image]
            try:
                base["container_execution_attempted"] = True
                try:
                    command_result = _bounded_command(command, self.timeout, 12 * 1024 * 1024)
                except (OSError, subprocess.TimeoutExpired):
                    command_result = _Command(None, unavailable=True)
            finally:
                try:
                    cleanup = _bounded_command(["docker", "rm", "-f", name], 10)
                except (OSError, subprocess.TimeoutExpired):
                    cleanup = _Command(None, unavailable=True)
            base["cleanup_confirmed"] = cleanup.code == 0 and not cleanup.timed_out
            if not base["cleanup_confirmed"]:
                return done("unknown", "container_cleanup_unconfirmed")
            if command_result.timed_out or command_result.overflow or command_result.code != 0:
                return done("unknown", "recalculation_timeout" if command_result.timed_out else "recalculation_execution_error")
            try:
                result = json.loads(command_result.stdout)
                if result["status"] != "available":
                    return done("unknown", result.get("reason", "recalculation_unavailable"),
                                libreoffice_version=result.get("libreoffice_version"))
                raw = base64.b64decode(result["output_base64"], validate=True)
                require(len(raw) <= MAX_BYTES and hashlib.sha256(raw).hexdigest() == result["output_sha256"],
                        "Output bytes mismatch")
                require(result.get("calculate_all") is True and result.get("macros") == "never_execute"
                        and result.get("links") == "no_update", "Worker policy mismatch")
                restored, changed = _restore_numeric_export(raw, manifest, result)
                return done("available", "independently_recalculated", output_base64=base64.b64encode(restored).decode(),
                            output_sha256=hashlib.sha256(restored).hexdigest(),
                            raw_export_base64=result["output_base64"], raw_export_sha256=result["output_sha256"],
                            numeric_manifest_sha256=result["numeric_manifest_sha256"],
                            numeric_literals_verified=len(manifest["cells"]), numeric_literal_texts_restored=changed,
                            libreoffice_version=result["libreoffice_version"])
            except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError,
                    expat.ExpatError, zipfile.BadZipFile):
                return done("unknown", "invalid_recalculation_receipt")


def _recalc_valid(before, after):
    try:
        return _recalc_content_guard(before, after)
    except (OSError, ValueError, KeyError, IndexError, TypeError, ET.ParseError, zipfile.BadZipFile):
        return "recalculation_invalid_workbook_structure"


def _recalc_content_guard(before, after):
    """Require formula retention and no unsupported-function errors after save."""
    with compat._safe_archive(after) as archive:
        for name in archive.namelist():
            if name.lower().endswith((".xml", ".rels")):
                _xml(archive, name)
    originals = openpyxl.load_workbook(before, data_only=False, keep_links=False)
    formulas = openpyxl.load_workbook(after, data_only=False, keep_links=False)
    values = openpyxl.load_workbook(after, data_only=True, keep_links=False)
    try:
        if originals.sheetnames != formulas.sheetnames:
            return "recalculation_sheet_identity_changed"

        def names(book):
            return {(scope, name): value.attr_text for scope, values in
                    [("workbook", book.defined_names), *((sheet.title, sheet.defined_names) for sheet in book)]
                    for name, value in values.items() if name not in {"_xlnm.Print_Area", "_xlnm.Print_Titles"}}

        if names(originals) != names(formulas):
            return "recalculation_changed_defined_names"
        derived = _array_outputs(originals)
        if derived != _array_outputs(formulas):
            return "recalculation_changed_array_range"
        boolean_spelling_changed = False
        for sheet in originals:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        if formulas[sheet.title][cell.coordinate].data_type != "f":
                            return "recalculation_removed_formula"
                        if not _same_formula(cell.value, formulas[sheet.title][cell.coordinate].value):
                            return "recalculation_changed_formula"
                        boolean_spelling_changed |= (_formula_spec(cell.value)
                            != _formula_spec(formulas[sheet.title][cell.coordinate].value))
                        value = values[sheet.title][cell.coordinate]
                        if value.data_type == "e" and value.value in {"#NAME?", "#GETTING_DATA", "#SPILL!"}:
                            return "unsupported_formula_result"
                    elif (formulas[sheet.title][cell.coordinate].data_type == "f"
                          or ((sheet.title, cell.coordinate) not in derived
                              and cell.value != values[sheet.title][cell.coordinate].value)):
                        return "recalculation_changed_literal_cell"
            for cell in formulas[sheet.title]._cells.values():
                if ((cell.row, cell.column) not in sheet._cells and cell.value is not None
                        and ((sheet.title, cell.coordinate) not in derived or cell.data_type == "f")):
                    return "recalculation_added_cell_content"
        if boolean_spelling_changed:
            # Logical-value equivalence does not make formula text equivalent:
            # introspection elsewhere can observe FALSE versus FALSE(). Do not
            # authorize that broader claim or patch any dependent result cache.
            observers = re.compile(r"(?:^|[^A-Z0-9_])(?:FORMULATEXT|FORMULA|GET\.CELL)\s*\(", re.I)
            texts = [_formula_spec(cell.value)[1] for sheet in originals
                     for cell in sheet._cells.values() if cell.data_type == "f"]
            texts.extend(names(originals).values())
            if any(observers.search(text) for text in texts):
                return "recalculation_formula_text_dependency_unverified"
        empty_formula_caches = compat.explicit_empty_formula_strings(after)
        for sheet, coordinate in derived:
            value = values[sheet][coordinate]
            if value.data_type == "e" and value.value in {"#NAME?", "#GETTING_DATA", "#SPILL!"}:
                return "unsupported_formula_result"
            # Followers without a calculable result are not ordinary blank cells.
            # A real empty-string cache is accepted only where OOXML establishes
            # it explicitly; absent follower results stay conservatively unknown.
            if value.value is None and (sheet, coordinate) not in empty_formula_caches:
                return "recalculation_array_result_unavailable"
        return None
    finally:
        originals.close()
        formulas.close()
        values.close()


def evaluate_pair(prediction, reference, answer_position, engine, cache):
    """Both inputs are detached. Original reference caches are a drift guard only."""
    receipts, paths = {}, {}
    with tempfile.TemporaryDirectory(prefix="sheet-recalc-compare-") as directory:
        for role, original in (("reference", reference), ("prediction", prediction)):
            raw_hash = sha(original)
            key = digest({"input": raw_hash, "engine": engine.identity})
            receipt_path = safe_path(cache / f"{key}.json")
            if receipt_path.is_file():
                receipt = read_json(receipt_path, sealed=True)
                require(receipt["input_sha256"] == raw_hash and receipt["identity"] == engine.identity,
                        "Recalculation cache binding mismatch")
            else:
                receipt = engine.run(original)
                write_json(receipt_path, receipt)
            receipts[role] = receipt["record_hash"]
            require(receipt["input_sha256"] == raw_hash and receipt["identity"] == engine.identity,
                    "Recalculation receipt binding mismatch")
            if receipt["status"] != "available":
                return {"status": "unknown", "reason": role + ":" + receipt["reason"], "receipts": receipts}
            path = Path(directory) / f"{role}.xlsx"
            raw = base64.b64decode(receipt["output_base64"], validate=True)
            require(len(raw) <= MAX_BYTES and receipt["cleanup_confirmed"] is True
                    and hashlib.sha256(raw).hexdigest() == receipt["output_sha256"], "Recalculation cache bytes mismatch")
            path.write_bytes(raw)
            paths[role] = path
            reason = _recalc_valid(original, path)
            if reason:
                return {"status": "unknown", "reason": role + ":" + reason, "receipts": receipts}
        # Native cached reference and this engine must agree on the target.
        # Do not turn reference corruption/Excel-vs-Calc drift into model error.
        drift = compat.evaluate(paths["reference"], reference, "", answer_position)
        if drift["status"] != "passed":
            return {"status": "unknown", "reason": "reference_native_cache_drift_or_unavailable", "receipts": receipts}
        scored = compat.evaluate(paths["prediction"], paths["reference"], "", answer_position)
        return {"status": {"passed": "pass", "failed": "fail", "missing_output": "fail"}.get(scored["status"], "unknown"),
                "reason": scored["reason"], "evidence": scored["evidence"], "receipts": receipts}


def prepare(source, output, qualification):
    source, output = safe_path(source), safe_path(output)
    require(not output.is_relative_to(source) and not source.is_relative_to(output), "Separate immutable output required")
    qualification_path = safe_path(qualification)
    qualified = read_json(qualification_path, sealed=True)
    require(qualified["version"] == VERSION and qualified["status"] == "qualified", "Successful qualification required")
    require(qualified["engine"]["sources"] == _sources(), "Qualification source mismatch")
    inventory = {str(qualification_path): sha(qualification_path)}

    def read(path, sealed=False):
        value = read_json(path, sealed=sealed)
        inventory[str(safe_path(path))] = sha(path)
        return value

    plan = read(source / "plan.json", sealed=True)
    require(plan["config"]["partition"] == "development", "Only authorized development replay")
    panel_path = plan["panels"]["spreadsheetbench"]["path"]
    read(panel_path)
    tasks = {digest(task): task for task in panel_tasks(plan, "spreadsheetbench")}
    scores = {}
    for path in sorted((source / "host_only/scores").glob("*.json")):
        record = read(path, sealed=True)
        if record["benchmark"] == "spreadsheetbench":
            require(record["prediction_hash"] not in scores, "Duplicate native score")
            scores[record["prediction_hash"]] = record
    slots, identities = [], set()
    for path in sorted((source / "predictions").glob("*/prediction.json")):
        record = read(path, sealed=True)
        request = record["request"]
        if request["benchmark"] != "spreadsheetbench":
            continue
        require(request["plan_hash"] == plan["record_hash"] and path.parent.name == digest(request), "Source position mismatch")
        require(read(path.parent / "intent.json", sealed=True) == seal(request), "Position intent mismatch")
        task = tasks[request["task_hash"]]
        identity = (request["checkpoint_hash"], task["task_id"], request["repeat"])
        require(identity not in identities and 0 <= request["repeat"] < plan["repeats"], "Duplicate or invalid task repeat")
        identities.add(identity)
        cp_path = source / "checkpoints/no_skill/h0/s0.json"
        cp = read(cp_path, sealed=True)
        require(cp["record_hash"] == request["checkpoint_hash"] and cp["method"] == "no_skill"
                and cp["skill_text"] == "", "This protocol replays the completed no-skill panel only")
        old = scores.pop(record["record_hash"])
        require(old["task_id"] == task["task_id"] and old["repeat"] == request["repeat"]
                and old["checkpoint_hash"] == request["checkpoint_hash"] and old["plan_hash"] == plan["record_hash"],
                "Native score identity mismatch")
        for gold in task["private"]["test_files"]:
            digest_gold = sha(gold)
            require(task["private"]["asset_sha256"].get(gold) == digest_gold, "Source reference changed")
            inventory[str(safe_path(gold))] = digest_gold
        slots.append({"id": path.parent.name, "task": task, "request": request,
                      "prediction_path": str(path), "prediction_hash": record["record_hash"],
                      "old_score": old["status"], "old_score_hash": old["record_hash"]})
    require(not scores and len(slots) == len(tasks) * plan["repeats"], "Complete panel with all repeats required")
    protocol = seal({"version": VERSION, "source": str(source), "plan_hash": plan["record_hash"],
                     "slots": slots, "inventory": inventory, "engine": qualified["engine"],
                     "qualification_hash": qualified["record_hash"], "model_api_calls": 0,
                     "historical_scores_replaced": False, "feedback_allowed": False,
                     "scope": "uniform_all_delivered_workbooks_full_denominator"})
    with output_lock(output):
        write_json(output / "protocol.json", protocol)
    return {"status": "prepared", "positions": len(slots), "protocol_hash": protocol["record_hash"]}


def run(output):
    output = safe_path(output)
    with output_lock(output):
        protocol = read_json(output / "protocol.json", sealed=True)
        require(protocol["engine"]["sources"] == _sources(), "Frozen sources changed")
        for path, expected in protocol["inventory"].items():
            require(sha(path) == expected, "Frozen source evidence changed")
        engine = Recalculator(protocol["engine"]["image_id"], protocol["engine"]["timeout_seconds"])
        require(engine.identity == protocol["engine"], "Execution policy changed")
        for slot in protocol["slots"]:
            destination = output / "positions" / (slot["id"] + ".json")
            if destination.is_file():
                old = read_json(destination, sealed=True)
                require(old["protocol_hash"] == protocol["record_hash"] and old["slot_id"] == slot["id"], "Replay result mismatch")
                continue
            if (output / "PAUSE").exists():
                break
            record = read_json(slot["prediction_path"], sealed=True)
            require(record["record_hash"] == slot["prediction_hash"], "Prediction changed")
            prediction, task = record["prediction"], slot["task"]
            cases = (prediction.get("output") or {}).get("cases", [])
            results = []
            if not cases:
                results = [_undelivered(prediction["reason"])]
            elif len(cases) != len(task["private"]["test_files"]):
                results = [{"status": "unknown", "reason": "original_case_count_mismatch"}]
            else:
                with tempfile.TemporaryDirectory(prefix="sheet-recalc-prediction-") as temporary:
                    for index, (case, reference) in enumerate(zip(cases, task["private"]["test_files"])):
                        if case["status"] != "available":
                            results.append(_undelivered(case["reason"]))
                            continue
                        raw = base64.b64decode(case["output_base64"], validate=True)
                        require(len(raw) <= MAX_BYTES, "Prediction exceeds workbook limit")
                        path = Path(temporary).resolve() / f"{index}.xlsx"
                        path.write_bytes(raw)
                        results.append(evaluate_pair(path, reference, task["private"]["answer_position"], engine, output / "recalculations"))
            counts = Counter(result["status"] for result in results)
            status = "fail" if counts["fail"] else "unknown" if counts["unknown"] else "pass"
            write_json(destination, seal({"protocol_hash": protocol["record_hash"], "slot_id": slot["id"],
                "task_id": task["task_id"], "repeat": slot["request"]["repeat"], "old_status": slot["old_score"],
                "status": status, "cases": results, "model_api_calls": 0}))
    return report(output)


def _undelivered(reason):
    category = ("generation_truncation" if reason == "model_response_truncated" else
                "generated_code_execution_exception" if reason.startswith("native_exception:") else
                "missing_deliverable" if reason == "output.xlsx_not_produced" else "unresolved_execution")
    return {"status": "unknown", "reason": "original_no_delivered_workbook:" + reason,
            "failure_category": category, "new_execution_performed": False}


def report(output):
    output = safe_path(output)
    protocol = read_json(output / "protocol.json", sealed=True)
    records = []
    for slot in protocol["slots"]:
        path = output / "positions" / (slot["id"] + ".json")
        if path.is_file():
            record = read_json(path, sealed=True)
            require(record["protocol_hash"] == protocol["record_hash"] and record["slot_id"] == slot["id"], "Result identity mismatch")
            records.append(record)
    receipts = [read_json(path, sealed=True) for path in (output / "recalculations").glob("*.json")]
    return seal({"version": VERSION, "protocol_hash": protocol["record_hash"],
        "status": "complete" if len(records) == len(protocol["slots"]) else "pending",
        "positions": len(protocol["slots"]), "completed": len(records),
        "old_counts": dict(Counter(slot["old_score"] for slot in protocol["slots"])),
        "counts": dict(Counter(record["status"] for record in records)),
        "transitions": dict(Counter(record["old_status"] + "->" + record["status"] for record in records)),
        "unknown_reasons": dict(Counter(case["reason"] for record in records for case in record["cases"] if case["status"] == "unknown")),
        "original_undelivered_categories": dict(Counter(case["failure_category"] for record in records
                                                         for case in record["cases"] if "failure_category" in case)),
        "recalculation_costs": {"unique_workbook_receipts": len(receipts),
            "containers_attempted": sum(receipt.get("container_execution_attempted", False) for receipt in receipts),
            "duration_seconds_sum": sum(receipt.get("duration_seconds", 0) for receipt in receipts),
            "cleanup_unconfirmed": sum(not receipt["cleanup_confirmed"] for receipt in receipts),
            "libreoffice_versions": sorted({receipt["libreoffice_version"] for receipt in receipts if receipt.get("libreoffice_version")})},
        "model_api_calls": 0, "historical_scores_replaced": False, "feedback_allowed": False,
        "excel_equivalence_proven": False, "is_method_effect": False})


def qualify(output, image):
    """Engine controls are fixtures, not evidence of model or Skill benefit."""
    output = safe_path(output)
    engine = Recalculator(image)
    with output_lock(output):
        fixtures = {
            "correct": ("=SUM(A1:A2)", 5), "wrong": ("=SUM(A1:A2)+1", 6),
            "empty_string": ('=IF(A1=2,"",1)', None), "retained_dependency": ("=C1*2", 10),
            "unsupported": ("=SKILLOPT_NOT_A_FUNCTION(A1)", "#NAME?"),
            "array_correct": (ArrayFormula(ref="B1:B2", text="=A1:A2*2"), {"B1": 4, "B2": 6}),
            "array_wrong": (ArrayFormula(ref="B1", text="=SUM(A1:A2*2)+1"), 11),
            "numeric_roundtrip": ("=A1*10", 1.234567890123456),
            "numeric_wrong": ("=A1*10+1", 2.234567890123456),
            "boolean_false": ("=IF(FALSE,999,5)", 5),
            "boolean_true": ("=IF(TRUE,5,999)", 5),
            "boolean_wrong": ("=IF(FALSE,5,6)", 6),
            "boolean_string_wrong": ('="FALSE()"', "FALSE()"),
            "utf16le_metadata": ("=SUM(A1:A2)", 5),
            "utf16be_metadata": ("=SUM(A1:A2)", 5),
            "utf16_external": ("=SUM(A1:A2)", None),
            "utf16_doctype": ("=SUM(A1:A2)", None),
        }
        results = []
        for name, (formula, expected) in fixtures.items():
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "S"
            sheet["A1"], sheet["A2"], sheet["C1"], sheet["B1"] = 2, 3, "=SUM(A1:A2)", formula
            if name.startswith("numeric_"):
                sheet["A1"] = 0.1234567890123456
            raw = io.BytesIO()
            workbook.save(raw)
            workbook.close()
            if name.startswith("utf16"):
                endian = "be" if name == "utf16be_metadata" else "le"
                text = '<?xml version="1.0" encoding="UTF-16"?><fixture><note>\u5b89\u5168</note></fixture>'
                if name == "utf16_external":
                    text = ('<?xml version="1.0" encoding="UTF-16"?>'
                            '<Relationships><Relationship TargetMode="External" '
                            'Target="https://example.invalid/never-fetch"/></Relationships>')
                elif name == "utf16_doctype":
                    text = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE fixture><fixture/>'
                encoded = (b"\xff\xfe" if endian == "le" else b"\xfe\xff") + text.encode("utf-16-" + endian)
                enriched = io.BytesIO()
                with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as source, zipfile.ZipFile(enriched, "w") as dest:
                    for item in source.infolist():
                        dest.writestr(item, source.read(item.filename))
                    dest.writestr("customXml/skillopt-qualification.xml", encoded)
                raw = enriched
            path = output / "fixtures" / (name + ".xlsx")
            path.parent.mkdir(parents=True, exist_ok=True)
            require(not path.exists() or path.read_bytes() == raw.getvalue(), "Fixture directory already used")
            path.write_bytes(raw.getvalue())
            receipt = engine.run(path)
            write_json(output / "receipts" / (name + ".json"), receipt)
            qualified = False
            reason = receipt["reason"]
            actual = actual_cells = None
            if name in {"utf16_external", "utf16_doctype"}:
                expected_reason = ("unsupported_external_relationship" if name == "utf16_external"
                                   else "unsupported_workbook_structure")
                qualified = (receipt["status"] == "unknown" and reason == expected_reason
                             and receipt["container_execution_attempted"] is False)
            if receipt["status"] == "available":
                after = output / "fixtures" / (name + "-recalculated.xlsx")
                after.write_bytes(base64.b64decode(receipt["output_base64"], validate=True))
                values = openpyxl.load_workbook(after, data_only=True)
                actual = values["S"]["B1"].value
                actual_cells = ({coordinate: values["S"][coordinate].value for coordinate in expected}
                                if isinstance(expected, dict) else None)
                values.close()
                reason = _recalc_valid(path, after)
                if name == "unsupported":
                    qualified = actual == "#NAME?" and reason in {
                        "unsupported_formula_result", "recalculation_changed_formula"}
                else:
                    matches = (all(compat._legacy()._compare_cell_value(actual_cells[cell], value)
                                   for cell, value in expected.items()) if isinstance(expected, dict)
                               else compat._legacy()._compare_cell_value(actual, expected))
                    qualified = reason is None and matches
                    if name == "empty_string":
                        qualified = qualified and ("S", "B1") in compat.explicit_empty_formula_strings(after)
                if name == "wrong":
                    qualified = qualified and not compat._legacy()._compare_cell_value(actual, 5)
                if name == "array_wrong":
                    qualified = qualified and not compat._legacy()._compare_cell_value(actual, 10)
                if name == "numeric_wrong":
                    qualified = qualified and not compat._legacy()._compare_cell_value(actual, 1.234567890123456)
                if name == "boolean_wrong":
                    qualified = qualified and not compat._legacy()._compare_cell_value(actual, 5)
                if name == "boolean_string_wrong":
                    qualified = qualified and not compat._legacy()._compare_cell_value(actual, "FALSE")
            results.append({"name": name, "qualified": qualified, "reason": reason,
                            "receipt_hash": receipt["record_hash"], "cleanup_confirmed": receipt["cleanup_confirmed"],
                            "libreoffice_version": receipt.get("libreoffice_version"),
                            "numeric_literals_verified": receipt.get("numeric_literals_verified"),
                            "numeric_literal_texts_restored": receipt.get("numeric_literal_texts_restored"),
                            "expected": expected, "actual": actual_cells if isinstance(expected, dict) else actual})
        report = seal({"version": VERSION, "status": "qualified" if all(item["qualified"] for item in results) else "rejected",
                       "engine": engine.identity, "controls": results, "model_api_calls": 0,
                       "evidence_kind": "engineering_fixture", "excel_equivalence_proven": False})
        write_json(output / "qualification.json", report)
        return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    qualification = sub.add_parser("qualify")
    qualification.add_argument("--output", required=True)
    qualification.add_argument("--image", required=True)
    preparation = sub.add_parser("prepare")
    preparation.add_argument("--source", required=True)
    preparation.add_argument("--output", required=True)
    preparation.add_argument("--qualification", required=True)
    for name in ("run", "report"):
        sub.add_parser(name).add_argument("--output", required=True)
    args = vars(parser.parse_args())
    command = args.pop("command")
    print(json.dumps(globals()[command](**args), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
