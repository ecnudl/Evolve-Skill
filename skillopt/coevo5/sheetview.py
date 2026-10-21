"""v3 (registration delta §6): the answer-range cells of a SpreadsheetBench answer as the scorer compared them, for the
diagnosis, the rubric reviser and the analyst (never the judge, which sees no reference).

For the first case of a cell's final scoring: the scorer's own recalculated reference and prediction workbooks, found
by the receipt hashes its result names for that case in the cell's host-only scoring directory (each receipt's input
hash checked against the reference file and the written workbook), and the workbook the program wrote (a value or a
formula per cell). Up to ``SHEET_CELLS`` cells in answer-range order; the first cell the scorer named as mismatching is
always among them. Values stay JSON-native (numbers are numbers), text and dates clipped to ``SHEET_VALUE_CHARS``. A
recalculated reference that cannot be found falls back to the reference file's own cached values and says so; anything
else that cannot be found or read leaves its column null. Zero calls; reads only the cell's own records.
"""
from __future__ import annotations

import base64
import binascii
import hashlib
import io
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from skillopt.continual_eval.core import read_json, safe_path
from skillopt.validator_pilot.api import digest

from .protocol import SHEET_CELLS, SHEET_VALUE_CHARS

MISMATCH = re.compile(r"value@(.+)!([A-Z]{1,3}[0-9]{1,7}): mismatch")
CLIP_MARK = " [...]"
ERRORS = (OSError, ValueError, KeyError, TypeError, IndexError, zipfile.BadZipFile, binascii.Error)


def _value(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (datetime, date, time)):
        value = value.isoformat()
    text = str(value)
    return text if len(text) <= SHEET_VALUE_CHARS else text[:SHEET_VALUE_CHARS - len(CLIP_MARK)] + CLIP_MARK


def _workbook(raw, *, data_only):
    import openpyxl

    return openpyxl.load_workbook(io.BytesIO(raw), data_only=data_only, keep_links=False)


def _read(workbook, sheet, coordinate):
    try:
        return _value(workbook[sheet][coordinate].value) if sheet in workbook.sheetnames else None
    except ERRORS:
        return None


def _scored(store, record):
    """(the final score artifact, its scoring directory) of a cell, or (None, None)."""
    scored = [a for a in record["attempts"] if a.get("score")]
    if not scored or not record["attempts"][-1].get("score"):
        return None, None
    request = {"benchmark": "spreadsheetbench", "unit": record["cell"]["unit"], "cell": record["cell"]["id"],
               "scoring": len(scored) - 1}
    score = read_json(store / "artifacts" / (record["attempts"][-1]["score"]["sha256"] + ".json"))
    return score, store / "host_only" / "scorer_artifacts" / digest(request)


def _recalculated(directory, wanted, source_sha):
    """The recalculated workbook bytes of the receipt whose record hash is ``wanted`` and whose input is
    ``source_sha``, or None."""
    if directory is None or not wanted or not (directory / "recalculations").is_dir():
        return None
    for path in sorted((directory / "recalculations").glob("*.json")):
        receipt = read_json(path, sealed=True)
        if receipt.get("record_hash") == wanted:
            if receipt.get("input_sha256") != source_sha or receipt.get("status") != "available":
                return None
            raw = base64.b64decode(receipt["output_base64"], validate=True)
            return raw if hashlib.sha256(raw).hexdigest() == receipt.get("output_sha256") else None
    return None


def cells(store, record, prediction, private):
    """The compared answer-range cells of the cell ``record`` (its ``prediction`` artifact; the task's ``private``
    part), or None when the answer range or the reference cannot be read."""
    from skillopt.continual_eval.spreadsheet_compat import _legacy, resolve_targets

    store = safe_path(store)
    try:
        reference_path = Path(private["test_files"][0])
        reference_raw = reference_path.read_bytes()
        output = prediction.get("output") if prediction.get("status") == "available" else None
        case = (output or {}).get("cases", [None])[0] if isinstance(output, dict) else None
        written_raw = (base64.b64decode(case["output_base64"], validate=True)
                       if isinstance(case, dict) and case.get("status") == "available" else None)
        score, directory = _scored(store, record)
        result = ((score or {}).get("metrics") or {}).get("cases") or [{}]
        receipts = result[0].get("receipts") or {}
        recalculated_reference = _recalculated(directory, receipts.get("reference"),
                                               hashlib.sha256(reference_raw).hexdigest())
        recalculated = (_recalculated(directory, receipts.get("prediction"), hashlib.sha256(written_raw).hexdigest())
                        if written_raw is not None else None)
        expected = _workbook(recalculated_reference or reference_raw, data_only=True)
        targets = resolve_targets(private["answer_position"], expected.sheetnames)["targets"]
        coordinates = [(sheet, c) for sheet, region in targets for c in _legacy()._iter_cell_names(region)]
    except ERRORS:
        return None
    reason = str(result[0].get("reason") or "")
    found = MISMATCH.fullmatch(reason)
    first = (found.group(1), found.group(2)) if found else None
    shown = coordinates[:SHEET_CELLS]
    if first is not None and first in coordinates and first not in shown:
        shown = shown[:SHEET_CELLS - 1] + [first]
    try:
        values = _workbook(recalculated, data_only=True) if recalculated is not None else None
    except ERRORS:
        values = None
    try:
        formulas = _workbook(written_raw, data_only=False) if written_raw is not None else None
    except ERRORS:
        formulas = None
    return {"case": 0, "cells_in_range": len(coordinates),
            "first_mismatch": "%s!%s" % first if first is not None else None,
            "expected_source": "recalculated_reference" if recalculated_reference is not None else "reference_cache",
            "recalculated": "available" if values is not None else "unavailable",
            "written": "available" if formulas is not None else "unavailable",
            "shown": [{"cell": "%s!%s" % (sheet, c), "expected": _read(expected, sheet, c),
                       "recalculated": _read(values, sheet, c) if values is not None else None,
                       "written": _read(formulas, sheet, c) if formulas is not None else None}
                      for sheet, c in shown]}
